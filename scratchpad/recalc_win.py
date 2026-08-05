"""
recalc_win.py
==============
Recalcula un .xlsx de verdad usando LibreOffice en modo headless (Windows),
y luego escanea todas las celdas en busca de errores de fórmula
(#DIV/0!, #REF!, #VALUE!, #NAME?, #N/A, #NULL!, #NUM!).

Por qué existe: LibreOffice, al abrir un .xlsx con --convert-to, NO
recalcula por defecto las fórmulas que ya traen un valor en caché de Excel
(para no ralentizar la carga) — así que simplemente "leer con openpyxl
data_only=True" solo te devuelve el último valor guardado por Excel, no
detecta fórmulas rotas nuevas. Para forzar el recálculo real hay que decirle
a LibreOffice explícitamente "recalcula siempre las fórmulas de Excel al
cargar" (OOXMLRecalcMode=0), lo cual se hace con un perfil de usuario
temporal que trae ese ajuste ya puesto (la skill oficial de recálculo usa
sockets Unix para hablar con soffice, que no existen en Windows; aquí se
usa en su lugar el propio --convert-to con un perfil aislado).

Uso:
    python scratchpad/recalc_win.py "ruta\\al\\archivo.xlsx" [timeout_seg]
"""
import sys
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"

REGISTRYMODS = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
 <item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
 <item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""

ERRORES_EXCEL = ("#DIV/0!", "#REF!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!")


def _preparar_perfil() -> str:
    perfil = Path(tempfile.mkdtemp(prefix="lo_profile_"))
    (perfil / "user" / "registrymodifications.xml").parent.mkdir(parents=True, exist_ok=True)
    (perfil / "user" / "registrymodifications.xml").write_text(REGISTRYMODS, encoding="utf-8")
    return str(perfil)


def recalcular(ruta_xlsx: str, timeout: int = 120) -> str:
    """
    Recalcula ruta_xlsx con LibreOffice y devuelve la ruta al .xlsx
    recalculado (en una carpeta temporal, no toca el original).
    """
    ruta_xlsx = str(Path(ruta_xlsx).resolve())
    perfil = _preparar_perfil()
    outdir = tempfile.mkdtemp(prefix="lo_out_")

    cmd = [
        SOFFICE,
        f"-env:UserInstallation=file:///{Path(perfil).as_posix()}",
        "--headless", "--norestore", "--nolockcheck",
        "--convert-to", "xlsx",
        "--outdir", outdir,
        ruta_xlsx,
    ]
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
    if proc.returncode != 0:
        raise RuntimeError(f"soffice falló ({proc.returncode}): {proc.stdout}\n{proc.stderr}")

    nombre = Path(ruta_xlsx).stem + ".xlsx"
    salida = Path(outdir) / nombre
    if not salida.exists():
        candidatos = list(Path(outdir).glob("*.xlsx"))
        if not candidatos:
            raise RuntimeError(f"soffice no generó salida en {outdir}. stdout={proc.stdout}")
        salida = candidatos[0]
    return str(salida)


def buscar_errores(ruta_xlsx: str) -> dict:
    """Abre el xlsx recalculado y devuelve {hoja: [(celda, valor_error), ...]}."""
    import openpyxl
    wb = openpyxl.load_workbook(ruta_xlsx, data_only=True)
    errores = {}
    for hoja in wb.sheetnames:
        ws = wb[hoja]
        encontrados = []
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.strip() in ERRORES_EXCEL:
                    encontrados.append((cell.coordinate, v))
        if encontrados:
            errores[hoja] = encontrados
    return errores


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) < 2:
        print("Uso: python recalc_win.py <archivo.xlsx> [timeout_seg]")
        sys.exit(1)
    ruta = sys.argv[1]
    timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 120

    print(f"Recalculando {ruta} con LibreOffice (timeout={timeout}s)...")
    salida = recalcular(ruta, timeout=timeout)
    print(f"Recalculado en: {salida}")

    errores = buscar_errores(salida)
    if not errores:
        print("\n✔ 0 errores de fórmula en todo el libro.")
    else:
        total = sum(len(v) for v in errores.values())
        print(f"\n✘ {total} errores de fórmula encontrados:")
        for hoja, celdas in errores.items():
            print(f"  [{hoja}]")
            for coord, val in celdas:
                print(f"    {coord}: {val}")
