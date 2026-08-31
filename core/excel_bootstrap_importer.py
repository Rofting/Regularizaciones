"""Importación auditable de libros maestros y fuentes de reparto.

El módulo convierte fuentes históricas en datos normalizados sin usar el libro
como base de cálculo futura. Cada valor conserva la hoja y celda de origen y
las fuentes se archivan antes de persistir su contenido.
"""

from __future__ import annotations

import csv
import hashlib
import os
import re
import shutil
import sqlite3
import tempfile
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

import xlrd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from document_review import create_review_issue
from excel_profiles import ExcelProfile
from expedient_service import (
    get_case,
    link_case_to_period,
    register_source_document,
)
from importar_excel_referencia import ReferenceValidationError, parse_reference_workbook


@dataclass(frozen=True)
class BootstrapImportResult:
    id_batch: int
    id_periodo: int
    imported_invoice_count: int
    imported_owner_count: int
    imported_reading_count: int
    open_issue_count: int
    installed_template_path: Path | None = None
    installed_template_sha256: str | None = None


@dataclass(frozen=True)
class CompanionImportResult:
    id_batches: tuple[int, ...]
    id_periodo: int
    imported_owner_count: int
    imported_reading_count: int
    open_issue_count: int


class TemplateInstallationError(ValueError):
    """No se puede instalar una plantilla privada de forma segura."""


class TemplateInstallationConflictError(TemplateInstallationError):
    """La comunidad ya tiene una plantilla diferente para ese perfil."""


@dataclass(frozen=True)
class _Source:
    entity_type: str
    entity_key: str
    field_name: str
    sheet_name: str
    cell_address: str
    raw_value: Any
    normalized_value: Any


_INVOICE_SCHEMAS = {
    "GAS": {
        "service": "GAS",
        "date": ("FECHA FRA", "FECHA FACTURA"),
        "start": ("FECHA INI", "FECHA INICIAL"),
        "end": ("FECHA FIN", "FECHA FINAL"),
        "consumption": ("KWH", "M3"),
        "fixed": ("EUR FIJO", "TERMINO FIJO", "CUOTA FIJA"),
        "variable": ("EUR VAR", "TERMINO VARIABLE", "CUOTA VARIABLE"),
        "total": ("TOTAL", "TOTAL IVA"),
        "notes": ("NOTA", "PROVEEDOR"),
    },
    "ELECTRICIDAD": {
        "service": "ELECTRICIDAD",
        "date": ("FECHA FRA", "FECHA FACTURA"),
        "start": ("FECHA INI", "FECHA INICIAL"),
        "end": ("FECHA FIN", "FECHA FINAL"),
        "consumption": ("KWH",),
        "fixed": ("EUR FIJO", "TERMINO FIJO", "CUOTA FIJA"),
        "variable": ("EUR VAR", "TERMINO VARIABLE", "CUOTA VARIABLE"),
        "total": ("TOTAL", "TOTAL IVA"),
        "notes": ("NOTA", "PROVEEDOR"),
    },
    "AGUA": {
        "service": "AGUA",
        "date": ("FECHA FRA", "FECHA FACTURA"),
        "start": ("FECHA INI", "FECHA INICIAL"),
        "end": ("FECHA FIN", "FECHA FINAL"),
        "consumption": ("M3",),
        "fixed": ("C FIJA", "CUOTA FIJA", "EUR FIJO"),
        "variable": ("C VARIABLE", "CUOTA VARIABLE", "EUR VAR"),
        "total": ("TOTAL IVA", "TOTAL"),
        "notes": ("NOTA", "PROVEEDOR"),
    },
}


def _header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").upper())
    text = "".join(character for character in text if not unicodedata.combining(character))
    text = text.replace("€", " EUR ").replace("³", "3")
    return " ".join(re.sub(r"[^A-Z0-9]+", " ", text).split())


def _clean(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


def _number(value: Any) -> float | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _clean(value).replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean(value)
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text[:10], pattern).date()
        except ValueError:
            continue
    return None


def _normalized(value: Any) -> str | None:
    if value is None:
        return None
    parsed_date = _date(value)
    if parsed_date is not None and not isinstance(value, (int, float)):
        return parsed_date.isoformat()
    if isinstance(value, float):
        return format(value, ".12g")
    return _clean(value) or None


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while chunk := source.read(64 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _archive_root(connection: sqlite3.Connection, source_path: Path) -> Path:
    database = next(
        (row[2] for row in connection.execute("PRAGMA database_list") if row[1] == "main"),
        "",
    )
    if database and database != ":memory:":
        return Path(database).resolve().parent / "expedientes"
    return source_path.resolve().parent / ".archivo_fuentes"


def _case_context(
    connection: sqlite3.Connection,
    id_case: int,
    profile: ExcelProfile,
) -> tuple[Any, int]:
    case = get_case(connection, id_case)
    community = connection.execute(
        "SELECT codigo FROM comunidades WHERE id_comunidad=?", (case.community_id,)
    ).fetchone()
    if community is None or str(community["codigo"]) != profile.community_code:
        raise ValueError("El perfil Excel no corresponde a la comunidad del expediente")
    return case, link_case_to_period(connection, id_case)


def _register_source(
    connection: sqlite3.Connection,
    *,
    id_case: int,
    source_path: Path,
    document_kind: str,
):
    if connection.in_transaction:
        connection.commit()
    return register_source_document(
        connection,
        id_case,
        source_path=source_path,
        archive_root=_archive_root(connection, source_path),
        document_kind=document_kind,
    )[0]


def _verified_archived_path(document) -> Path:
    archived_path = Path(document.archived_path).resolve()
    if not archived_path.is_file() or _sha256(archived_path) != document.sha256:
        raise RuntimeError("La copia archivada no supera la verificación de hash")
    return archived_path


def _private_template_path(project_root: Path, profile: ExcelProfile) -> Path:
    root = Path(project_root).resolve()
    destination = (root / profile.template_relative_path).resolve()
    try:
        destination.relative_to(root)
    except ValueError:
        raise TemplateInstallationError("La plantilla queda fuera del proyecto") from None
    return destination


def _profile_sha256(project_root: Path, profile: ExcelProfile) -> str:
    config_path = Path(project_root).resolve() / "config" / "excel_profiles" / f"{profile.key}.json"
    if not config_path.is_file():
        raise TemplateInstallationError(
            f"No existe el perfil configurado para instalar la plantilla: {config_path}"
        )
    return _sha256(config_path)


def _record_template_conflict(
    connection: sqlite3.Connection,
    *,
    id_case: int,
    id_document: int,
    profile: ExcelProfile,
    message: str,
) -> None:
    _issue(
        connection,
        id_case=id_case,
        id_document=id_document,
        code="TEMPLATE_VERSION_CONFLICT",
        field_name=f"template.{profile.key}",
        message=message,
    )


def _install_private_template(
    connection: sqlite3.Connection,
    *,
    id_case: int,
    document,
    profile: ExcelProfile,
    project_root: Path,
) -> tuple[Path, str]:
    """Publica una copia verificable del archivo ya archivado, sin sobrescribir.

    ``os.link`` reserva el nombre de destino en la misma carpeta de forma
    atómica. Si otra ejecución ya publicó el archivo, sólo se acepta cuando
    su huella es idéntica; de otro modo se registra una incidencia y se exige
    una versión de perfil nueva.
    """
    archive = _verified_archived_path(document)
    expected_hash = document.sha256
    if _sha256(archive) != expected_hash:
        raise TemplateInstallationError("La plantilla archivada no supera la verificación de hash")
    destination = _private_template_path(project_root, profile)
    _profile_sha256(project_root, profile)
    if destination.exists():
        if not destination.is_file():
            raise TemplateInstallationError("La ruta de plantilla instalada no es un archivo")
        installed_hash = _sha256(destination)
        if installed_hash == expected_hash:
            return destination, installed_hash
        message = (
            "La plantilla instalada no coincide con el maestro importado; "
            "debe registrarse una nueva versión de perfil antes de reemplazarla"
        )
        _record_template_conflict(
            connection, id_case=id_case, id_document=document.id_document,
            profile=profile, message=message,
        )
        raise TemplateInstallationConflictError(message)

    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary_name: str | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=destination.parent,
            prefix=f".{destination.name}.",
            suffix=".tmp",
            delete=False,
        ) as temporary:
            temporary_name = temporary.name
        temporary_path = Path(temporary_name)
        shutil.copy2(archive, temporary_path)
        if _sha256(temporary_path) != expected_hash:
            raise TemplateInstallationError("La copia temporal de plantilla no supera la verificación de hash")
        try:
            os.link(temporary_path, destination)
        except FileExistsError:
            installed_hash = _sha256(destination) if destination.is_file() else None
            if installed_hash == expected_hash:
                return destination, installed_hash
            message = (
                "Otra instalación publicó una plantilla distinta; debe registrarse "
                "una nueva versión de perfil antes de reemplazarla"
            )
            _record_template_conflict(
                connection, id_case=id_case, id_document=document.id_document,
                profile=profile, message=message,
            )
            raise TemplateInstallationConflictError(message)
        installed_hash = _sha256(destination)
        if installed_hash != expected_hash:
            try:
                destination.unlink(missing_ok=True)
            except OSError:
                pass
            raise TemplateInstallationError("La plantilla publicada no supera la verificación de hash")
        return destination, installed_hash
    finally:
        if temporary_name is not None:
            try:
                Path(temporary_name).unlink(missing_ok=True)
            except OSError:
                pass


def _register_installed_template(
    connection: sqlite3.Connection,
    *,
    id_case: int,
    document,
    community_id: int,
    profile: ExcelProfile,
    project_root: Path,
    installed_hash: str,
) -> None:
    """Registra sólo una plantilla que ya existe y cuya huella fue validada."""
    destination = _private_template_path(project_root, profile)
    if not destination.is_file() or _sha256(destination) != installed_hash:
        raise TemplateInstallationError("La plantilla instalada dejó de coincidir antes de registrarla")
    profile_hash = _profile_sha256(project_root, profile)
    active = connection.execute(
        """SELECT profile_key,profile_version,template_relative_path,template_sha256
           FROM excel_template_profiles WHERE id_comunidad=? AND status='active'
           ORDER BY id_template_profile""",
        (community_id,),
    ).fetchall()
    if active:
        if len(active) != 1:
            message = "La comunidad tiene varios perfiles activos; debe revisar su configuración"
            _record_template_conflict(
                connection, id_case=id_case, id_document=document.id_document,
                profile=profile, message=message,
            )
            raise TemplateInstallationConflictError(message)
        existing = active[0]
        if (
            existing["profile_key"] == profile.key
            and existing["profile_version"] == profile.version
            and existing["template_relative_path"] == profile.template_relative_path
            and existing["template_sha256"] == installed_hash
        ):
            return
        message = (
            "El registro de plantilla no coincide; debe registrarse una nueva versión "
            "de perfil antes de reemplazarla"
        )
        _record_template_conflict(
            connection, id_case=id_case, id_document=document.id_document,
            profile=profile, message=message,
        )
        raise TemplateInstallationConflictError(message)
    connection.execute(
        """INSERT INTO excel_template_profiles
           (id_comunidad,profile_key,profile_version,template_relative_path,
            template_sha256,profile_sha256,status)
           VALUES (?,?,?,?,?,?,'active')""",
        (
            community_id, profile.key, profile.version, profile.template_relative_path,
            installed_hash, profile_hash,
        ),
    )
    connection.commit()


def _existing_batch(
    connection: sqlite3.Connection,
    community_id: int,
    source_hash: str,
    source_kind: str,
) -> sqlite3.Row | None:
    return connection.execute(
        """SELECT id_batch,id_periodo FROM import_batches
           WHERE id_comunidad=? AND source_sha256=? AND source_kind=?""",
        (community_id, source_hash, source_kind),
    ).fetchone()


def _start_batch(
    connection: sqlite3.Connection,
    *,
    community_id: int,
    period_id: int,
    source_kind: str,
    archived_path: str,
    source_hash: str,
) -> int:
    cursor = connection.execute(
        """INSERT INTO import_batches
           (id_comunidad,id_periodo,source_kind,source_path,source_sha256,status)
           VALUES (?,?,?,?,?,'processing')""",
        (community_id, period_id, source_kind, str(archived_path), source_hash),
    )
    return int(cursor.lastrowid)


def _link_batch_to_case(
    connection: sqlite3.Connection,
    *,
    id_case: int,
    id_batch: int,
    id_periodo: int,
    source_kind: str,
    source_hash: str,
) -> None:
    connection.execute(
        """INSERT OR IGNORE INTO case_import_batches
           (id_case,id_batch,id_periodo,source_kind,source_sha256)
           VALUES (?,?,?,?,?)""",
        (id_case, id_batch, id_periodo, source_kind, source_hash),
    )


def _batch_is_linked_to_case(
    connection: sqlite3.Connection,
    *,
    id_case: int,
    id_batch: int,
    id_periodo: int,
    source_kind: str,
) -> bool:
    return connection.execute(
        """SELECT 1 FROM case_import_batches
           WHERE id_case=? AND id_batch=? AND id_periodo=? AND source_kind=?""",
        (id_case, id_batch, id_periodo, source_kind),
    ).fetchone() is not None


def _store_sources(
    connection: sqlite3.Connection, id_batch: int, values: Iterable[_Source]
) -> None:
    connection.executemany(
        """INSERT OR IGNORE INTO source_values
           (id_batch,entity_type,entity_key,field_name,sheet_name,cell_address,
            raw_value,normalized_value,validation_status)
           VALUES (?,?,?,?,?,?,?,?,?)""",
        [
            (
                id_batch,
                value.entity_type,
                value.entity_key,
                value.field_name,
                value.sheet_name,
                value.cell_address,
                None if value.raw_value is None else str(value.raw_value),
                _normalized(value.normalized_value),
                "validated" if _normalized(value.normalized_value) is not None else "candidate",
            )
            for value in values
        ],
    )


def _open_issue_count(connection: sqlite3.Connection, id_case: int) -> int:
    return int(connection.execute(
        "SELECT COUNT(*) FROM review_issues WHERE id_case=? AND status='open'",
        (id_case,),
    ).fetchone()[0])


def _validate_document_if_clean(connection: sqlite3.Connection, document_id: int) -> None:
    open_count = connection.execute(
        """SELECT COUNT(*) FROM review_issues
           WHERE id_document=? AND status='open'""",
        (document_id,),
    ).fetchone()[0]
    if not open_count:
        connection.execute(
            "UPDATE source_documents SET status='validated' WHERE id_document=?",
            (document_id,),
        )


def _issue(
    connection: sqlite3.Connection,
    *,
    id_case: int,
    id_document: int,
    code: str,
    field_name: str,
    message: str,
    detected_value: Any = None,
) -> None:
    create_review_issue(
        connection,
        id_case,
        id_document,
        code=code,
        field_name=field_name,
        message=message,
        detected_value=None if detected_value is None else str(detected_value),
    )


def _column_for(headers: dict[int, str], aliases: Iterable[str]) -> int | None:
    wanted = {_header(alias) for alias in aliases}
    for column, label in headers.items():
        if label in wanted:
            return column
    return None


def _columns_for(headers: dict[int, str], aliases: Iterable[str]) -> tuple[int, ...]:
    """Devuelve columnas en el orden de preferencia declarado por el perfil local."""
    result = []
    for alias in aliases:
        wanted = _header(alias)
        column = next((key for key, label in headers.items() if label == wanted), None)
        if column is not None and column not in result:
            result.append(column)
    return tuple(result)


def _find_invoice_header(sheet, schema: dict[str, Any]) -> tuple[int, dict[str, int | None]] | None:
    for row in range(1, min(sheet.max_row, 60) + 1):
        headers = {cell.column: _header(cell.value) for cell in sheet[row] if cell.value is not None}
        columns: dict[str, Any] = {
            key: _column_for(headers, schema[key])
            for key in ("date", "start", "end", "fixed", "variable", "total", "notes")
        }
        columns["consumption"] = _columns_for(headers, schema["consumption"])
        if columns["date"] is not None and columns["total"] is not None:
            return row, columns
    return None


def _invoice_identity(service: str, invoice_date: date, start: date | None,
                      end: date | None, provider: str) -> str:
    raw = "|".join((
        service,
        invoice_date.isoformat(),
        start.isoformat() if start else "",
        end.isoformat() if end else "",
        _header(provider),
    ))
    return f"BOOT-{service}-{hashlib.sha256(raw.encode('utf-8')).hexdigest()[:20]}"


def _cell_has_value(cell) -> bool:
    return cell is not None and _clean(cell.value) != ""


def _is_invoice_summary(value: Any) -> bool:
    label = _header(value)
    return label.startswith("SUMA") or label.startswith("TOTAL") or label.startswith("SUBTOTAL")


def _parse_invoice_sheets(
    connection: sqlite3.Connection,
    *,
    workbook,
    id_case: int,
    id_document: int,
    id_batch: int,
    community_id: int,
    period_id: int,
    case,
    archived_path: str,
) -> int:
    inserted = 0
    for sheet_name, schema in _INVOICE_SCHEMAS.items():
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        found = _find_invoice_header(sheet, schema)
        if found is None:
            _issue(
                connection, id_case=id_case, id_document=id_document,
                code="UNRECOGNIZED_HEADER", field_name=f"{sheet_name}.headers",
                message=f"No se reconocen las cabeceras de {sheet_name}",
            )
            continue
        header_row, columns = found
        for row in range(header_row + 1, sheet.max_row + 1):
            date_cell = sheet.cell(row, int(columns["date"]))
            start_cell = sheet.cell(row, int(columns["start"])) if columns["start"] else None
            end_cell = sheet.cell(row, int(columns["end"])) if columns["end"] else None
            total_cell = sheet.cell(row, int(columns["total"]))
            notes_cell = sheet.cell(row, int(columns["notes"])) if columns["notes"] else None
            relevant_cells = [date_cell, start_cell, end_cell, total_cell, notes_cell]
            relevant_cells.extend(
                sheet.cell(row, int(column)) for column in columns["consumption"]
            )
            for key in ("fixed", "variable"):
                if columns[key]:
                    relevant_cells.append(sheet.cell(row, int(columns[key])))
            if not any(_cell_has_value(cell) for cell in relevant_cells):
                continue
            if _is_invoice_summary(date_cell.value):
                continue
            invoice_date = _date(date_cell.value)
            if invoice_date is None:
                missing = not _cell_has_value(date_cell)
                _issue(
                    connection, id_case=id_case, id_document=id_document,
                    code="MISSING_REQUIRED_FIELD" if missing else "INCOMPATIBLE_DATE",
                    field_name=f"{sheet_name}.{date_cell.coordinate}.invoice_date",
                    message=(
                        "Falta la fecha requerida de la factura"
                        if missing else "La fecha de la factura no es válida"
                    ),
                    detected_value=date_cell.value,
                )
                continue
            if not (case.start_date <= invoice_date <= case.end_date):
                _issue(
                    connection, id_case=id_case, id_document=id_document,
                    code="INCOMPATIBLE_DATE",
                    field_name=f"{sheet_name}.{date_cell.coordinate}.invoice_date",
                    message="La fecha de la factura queda fuera del expediente",
                    detected_value=date_cell.value,
                )
                continue
            invalid_endpoint = False
            parsed_endpoints = {}
            for field_name, cell in (("start_date", start_cell), ("end_date", end_cell)):
                parsed = _date(cell.value) if cell is not None else None
                if parsed is None:
                    missing = not _cell_has_value(cell)
                    address = cell.coordinate if cell is not None else "?"
                    _issue(
                        connection, id_case=id_case, id_document=id_document,
                        code="MISSING_REQUIRED_FIELD" if missing else "INCOMPATIBLE_DATE",
                        field_name=f"{sheet_name}.{address}.{field_name}",
                        message=(
                            f"Falta el extremo requerido {field_name}"
                            if missing else f"El extremo {field_name} no es una fecha válida"
                        ),
                        detected_value=cell.value if cell is not None else None,
                    )
                    invalid_endpoint = True
                parsed_endpoints[field_name] = parsed
            if invalid_endpoint:
                continue
            start_date = parsed_endpoints["start_date"]
            end_date = parsed_endpoints["end_date"]
            if end_date < start_date:
                _issue(
                    connection, id_case=id_case, id_document=id_document,
                    code="INCOMPATIBLE_DATE",
                    field_name=f"{sheet_name}.{start_cell.coordinate}:{end_cell.coordinate}.supply_range",
                    message="El fin de suministro es anterior al inicio",
                    detected_value=f"{start_date} - {end_date}",
                )
                continue
            total = _number(total_cell.value)
            if total is None:
                _issue(
                    connection, id_case=id_case, id_document=id_document,
                    code="MISSING_REQUIRED_FIELD",
                    field_name=f"{sheet_name}.{total_cell.coordinate}.total",
                    message="Falta el total requerido de la factura",
                )
                continue
            provider = _clean(notes_cell.value if notes_cell else None) or schema["service"]
            invoice_number = _invoice_identity(
                schema["service"], invoice_date, start_date, end_date, provider
            )
            consumption_cell = next(
                (
                    sheet.cell(row, int(column))
                    for column in columns["consumption"]
                    if _number(sheet.cell(row, int(column)).value) is not None
                ),
                None,
            )
            fixed_cell = sheet.cell(row, int(columns["fixed"])) if columns["fixed"] else None
            variable_cell = (
                sheet.cell(row, int(columns["variable"])) if columns["variable"] else None
            )
            consumption = _number(consumption_cell.value) if consumption_cell else None
            fixed = _number(fixed_cell.value) if fixed_cell else None
            variable = _number(variable_cell.value) if variable_cell else None
            missing_required = []
            for key, cell, amount in (
                ("consumption", consumption_cell, consumption),
                ("fixed", fixed_cell, fixed),
                ("variable", variable_cell, variable),
            ):
                if cell is None or amount is None:
                    address = cell.coordinate if cell is not None else "?"
                    missing_required.append(key)
                    _issue(
                        connection, id_case=id_case, id_document=id_document,
                        code="MISSING_REQUIRED_FIELD",
                        field_name=f"{sheet_name}.{address}.{key}",
                        message=f"Falta el valor obligatorio {key} de la factura",
                    )
            if missing_required:
                continue
            if fixed is not None and variable is not None and abs((fixed + variable) - total) > 0.02:
                _issue(
                    connection, id_case=id_case, id_document=id_document,
                    code="TOTAL_MISMATCH", field_name=f"{sheet_name}.{total_cell.coordinate}.total",
                    message="El total no coincide con la suma de los componentes",
                    detected_value=f"{fixed} + {variable} != {total}",
                )
            cursor = connection.execute(
                """INSERT OR IGNORE INTO facturas
                   (id_comunidad,id_periodo,tipo_suministro,proveedor,cups_o_referencia,
                    num_factura,fecha_factura,fecha_inicio,fecha_fin,consumo_total,
                    unidad_consumo,termino_fijo,termino_variable,importe_total,
                    archivo_origen,notas)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    community_id, period_id, schema["service"], provider,
                    f"BOOTSTRAP:{schema['service']}", invoice_number,
                    invoice_date.isoformat(),
                    start_date.isoformat() if start_date else None,
                    end_date.isoformat() if end_date else None,
                    consumption or 0.0,
                    (
                        "m3"
                        if consumption_cell is not None
                        and _header(sheet.cell(header_row, consumption_cell.column).value) == "M3"
                        else "kWh"
                    ),
                    fixed or 0.0, variable or 0.0, total,
                    str(archived_path), provider,
                ),
            )
            if cursor.rowcount:
                inserted += 1
            invoice = connection.execute(
                """SELECT id_factura FROM facturas
                   WHERE id_comunidad=? AND num_factura=? AND cups_o_referencia=?""",
                (community_id, invoice_number, f"BOOTSTRAP:{schema['service']}"),
            ).fetchone()
            components = (
                ("fixed", fixed, fixed_cell),
                ("variable", variable, variable_cell),
                ("total", total, total_cell),
            )
            for component_key, amount, cell in components:
                if amount is None or cell is None:
                    continue
                connection.execute(
                    """INSERT OR IGNORE INTO invoice_components
                       (id_factura,component_key,amount,unit,source_sheet,source_cell)
                       VALUES (?,?,?,?,?,?)""",
                    (invoice[0], component_key, amount, "EUR", sheet_name, cell.coordinate),
                )
            values = [
                _Source("invoice", invoice_number, "invoice_date", sheet_name,
                        date_cell.coordinate, date_cell.value, invoice_date),
                _Source("invoice", invoice_number, "total", sheet_name,
                        total_cell.coordinate, total_cell.value, total),
            ]
            for field, cell, value in (
                ("start_date", start_cell, start_date),
                ("end_date", end_cell, end_date),
                ("consumption", consumption_cell, consumption),
                ("fixed", fixed_cell, fixed),
                ("variable", variable_cell, variable),
                ("provider", notes_cell, provider),
            ):
                if cell is not None:
                    values.append(_Source(
                        "invoice", invoice_number, field, sheet_name,
                        cell.coordinate, cell.value, value,
                    ))
            _store_sources(connection, id_batch, values)
    return inserted


def _store_parameter(
    connection: sqlite3.Connection,
    *,
    community_id: int,
    period_id: int,
    key: str,
    value: float,
    unit: str,
    source_sheet: str,
    source_cell: str,
) -> None:
    existing = connection.execute(
        """SELECT numeric_value FROM period_parameters
           WHERE id_comunidad=? AND id_periodo=? AND parameter_key=?""",
        (community_id, period_id, key),
    ).fetchone()
    if existing is None:
        connection.execute(
            """INSERT INTO period_parameters
               (id_comunidad,id_periodo,parameter_key,numeric_value,unit,source_sheet,source_cell)
               VALUES (?,?,?,?,?,?,?)""",
            (community_id, period_id, key, value, unit, source_sheet, source_cell),
        )


def _parse_reference_parameters(
    connection: sqlite3.Connection,
    *,
    path: Path,
    id_case: int,
    id_document: int,
    id_batch: int,
    community_id: int,
    period_id: int,
    case,
) -> None:
    try:
        reference = parse_reference_workbook(path)
    except (ReferenceValidationError, ValueError, TypeError) as error:
        _issue(
            connection, id_case=id_case, id_document=id_document,
            code="UNRECOGNIZED_HEADER", field_name="ANALISIS.reference",
            message="No se puede interpretar el bloque económico del análisis",
            detected_value=error,
        )
        return
    if reference.start_date != case.start_date or reference.end_date != case.end_date:
        _issue(
            connection, id_case=id_case, id_document=id_document,
            code="INCOMPATIBLE_DATE", field_name="DATOS.period",
            message="El periodo del libro no coincide con el expediente",
            detected_value=f"{reference.start_date} - {reference.end_date}",
        )
    sources = []
    for concept in reference.concepts:
        for suffix, cents in (("actual", concept.actual_cents), ("billed", concept.billed_cents)):
            key = f"{concept.concept_key}_{suffix}"
            matching = next(
                (source for source in concept.source_values if source.field_name == f"{suffix}_cents"),
                concept.source_values[0],
            )
            _store_parameter(
                connection, community_id=community_id, period_id=period_id,
                key=key, value=cents / 100, unit="EUR",
                source_sheet=matching.sheet_name, source_cell=matching.cell_address,
            )
        sources.extend(
            _Source(
                source.entity_type, source.entity_key, source.field_name,
                source.sheet_name, source.cell_address,
                source.raw_value, source.normalized_value,
            )
            for source in concept.source_values
        )
    sources.extend(
        _Source(
            source.entity_type, source.entity_key, source.field_name,
            source.sheet_name, source.cell_address,
            source.raw_value, source.normalized_value,
        )
        for source in reference.source_values
    )
    _store_sources(connection, id_batch, sources)


def _validate_reused_master_period(
    connection: sqlite3.Connection,
    *,
    archived_path: Path,
    id_case: int,
    id_document: int,
    case,
) -> None:
    """Revalida solo el contrato temporal al reutilizar un lote físico."""
    try:
        reference = parse_reference_workbook(archived_path)
    except (ReferenceValidationError, ValueError, TypeError) as error:
        _issue(
            connection, id_case=id_case, id_document=id_document,
            code="UNRECOGNIZED_HEADER", field_name="ANALISIS.reference",
            message="No se puede reinterpretar el período del libro maestro",
            detected_value=error,
        )
        return
    if reference.start_date != case.start_date or reference.end_date != case.end_date:
        _issue(
            connection, id_case=id_case, id_document=id_document,
            code="INCOMPATIBLE_DATE", field_name="DATOS.period",
            message="El periodo del libro no coincide con el expediente actual",
            detected_value=f"{reference.start_date} - {reference.end_date}",
        )


def _parse_other_expenses(
    connection: sqlite3.Connection,
    *,
    workbook,
    id_batch: int,
    community_id: int,
    period_id: int,
    case_name: str,
) -> None:
    if "OTROS GASTOS" not in workbook.sheetnames:
        return
    sheet = workbook["OTROS GASTOS"]
    total = 0.0
    sources = []
    active = False
    for row in range(1, sheet.max_row + 1):
        labels = [_header(cell.value) for cell in sheet[row]]
        if any(label == _header(f"EJERCICIO {case_name}") for label in labels):
            active = True
            continue
        if active and any(label.startswith("EJERCICIO ") for label in labels):
            break
        if not active:
            continue
        comment = _clean(sheet.cell(row, 5).value)
        amount_cell = sheet.cell(row, 6)
        amount = _number(amount_cell.value)
        if not comment or comment.upper().startswith("TOTAL") or amount is None:
            continue
        total += amount
        sources.append(_Source(
            "period_parameter", "extraordinary_expense_actual", "amount",
            sheet.title, amount_cell.coordinate, amount_cell.value, amount,
        ))
    if sources:
        _store_parameter(
            connection, community_id=community_id, period_id=period_id,
            key="extraordinary_expense_actual", value=total, unit="EUR",
            source_sheet=sources[0].sheet_name, source_cell=sources[0].cell_address,
        )
        _store_sources(connection, id_batch, sources)


def _trace_aggregate_readings(connection: sqlite3.Connection, *, workbook,
                              id_batch: int, case_name: str) -> None:
    if "LECTURAS ACS M3" not in workbook.sheetnames:
        return
    sheet = workbook["LECTURAS ACS M3"]
    sources = []
    for row in range(1, sheet.max_row + 1):
        if _clean(sheet.cell(row, 1).value) != case_name:
            continue
        for column, field in ((4, "initial"), (6, "final"), (7, "consumption")):
            cell = sheet.cell(row, column)
            sources.append(_Source(
                "aggregate_reading", f"ACS:{case_name}:{row}", field,
                sheet.title, cell.coordinate, cell.value, cell.value,
            ))
    _store_sources(connection, id_batch, sources)


def import_master_excel(
    connection: sqlite3.Connection,
    *,
    id_case: int,
    workbook_path: Path,
    profile: ExcelProfile,
    actor: str,
    project_root: Path | None = None,
) -> BootstrapImportResult:
    """Importa un modelo maestro sin usarlo como fuente viva de cálculo."""
    path = Path(workbook_path).resolve()
    if not path.is_file():
        raise FileNotFoundError(f"No existe el libro maestro: {path}")
    if not actor.strip():
        raise ValueError("El responsable de la importación es obligatorio")
    case, period_id = _case_context(connection, id_case, profile)
    document = _register_source(
        connection, id_case=id_case, source_path=path,
        document_kind="excel_master_bootstrap",
    )
    archived_path = _verified_archived_path(document)
    installed_template_path: Path | None = None
    installed_template_sha256: str | None = None
    if project_root is not None:
        installed_template_path, installed_template_sha256 = _install_private_template(
            connection,
            id_case=id_case,
            document=document,
            profile=profile,
            project_root=Path(project_root),
        )
        _register_installed_template(
            connection,
            id_case=id_case,
            document=document,
            community_id=case.community_id,
            profile=profile,
            project_root=Path(project_root),
            installed_hash=installed_template_sha256,
        )
    source_kind = "excel_master_bootstrap"
    existing = _existing_batch(
        connection, case.community_id, document.sha256, source_kind
    )
    if existing is not None:
        existing_batch_id = int(existing["id_batch"])
        if _batch_is_linked_to_case(
            connection, id_case=id_case, id_batch=existing_batch_id,
            id_periodo=period_id, source_kind=source_kind,
        ):
            return BootstrapImportResult(
                existing_batch_id, period_id, 0, 0, 0,
                _open_issue_count(connection, id_case), installed_template_path,
                installed_template_sha256,
            )
        connection.execute("BEGIN IMMEDIATE")
        try:
            _link_batch_to_case(
                connection, id_case=id_case, id_batch=existing_batch_id,
                id_periodo=period_id, source_kind=source_kind,
                source_hash=document.sha256,
            )
            _validate_reused_master_period(
                connection, archived_path=archived_path, id_case=id_case,
                id_document=document.id_document, case=case,
            )
            _validate_document_if_clean(connection, document.id_document)
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        return BootstrapImportResult(
            existing_batch_id, period_id, 0, 0, 0,
            _open_issue_count(connection, id_case), installed_template_path,
            installed_template_sha256,
        )

    archived_bytes = archived_path.read_bytes()
    formula_book = load_workbook(BytesIO(archived_bytes), data_only=False)
    values_book = load_workbook(BytesIO(archived_bytes), data_only=True)
    connection.execute("BEGIN IMMEDIATE")
    try:
        batch_id = _start_batch(
            connection, community_id=case.community_id, period_id=period_id,
            source_kind=source_kind,
            archived_path=document.archived_path, source_hash=document.sha256,
        )
        _link_batch_to_case(
            connection, id_case=id_case, id_batch=batch_id,
            id_periodo=period_id, source_kind=source_kind,
            source_hash=document.sha256,
        )
        for sheet_name in profile.required_sheets:
            if sheet_name not in formula_book.sheetnames:
                _issue(
                    connection, id_case=id_case, id_document=document.id_document,
                    code="MISSING_SHEET", field_name=sheet_name,
                    message=f"Falta la hoja obligatoria {sheet_name}",
                )
        imported = _parse_invoice_sheets(
            connection, workbook=values_book, id_case=id_case,
            id_document=document.id_document, id_batch=batch_id,
            community_id=case.community_id, period_id=period_id, case=case,
            archived_path=document.archived_path,
        )
        _parse_reference_parameters(
            connection, path=archived_path, id_case=id_case,
            id_document=document.id_document, id_batch=batch_id,
            community_id=case.community_id, period_id=period_id, case=case,
        )
        _parse_other_expenses(
            connection, workbook=values_book, id_batch=batch_id,
            community_id=case.community_id, period_id=period_id,
            case_name=case.name,
        )
        _trace_aggregate_readings(
            connection, workbook=values_book, id_batch=batch_id, case_name=case.name
        )
        connection.execute(
            """UPDATE import_batches SET status='validated',validated_at=datetime('now')
               WHERE id_batch=?""",
            (batch_id,),
        )
        _validate_document_if_clean(connection, document.id_document)
        connection.commit()
    except Exception as error:
        connection.rollback()
        raise error
    finally:
        formula_book.close()
        values_book.close()
    return BootstrapImportResult(
        batch_id, period_id, imported, 0, 0,
        _open_issue_count(connection, id_case), installed_template_path,
        installed_template_sha256,
    )


def _import_owner_source(
    connection: sqlite3.Connection,
    *,
    id_case: int,
    case,
    period_id: int,
    path: Path,
) -> tuple[int, int]:
    document = _register_source(
        connection, id_case=id_case, source_path=path,
        document_kind="owner_list",
    )
    archived_path = _verified_archived_path(document)
    source_kind = "owner_list"
    existing = _existing_batch(
        connection, case.community_id, document.sha256, source_kind
    )
    existing_batch_id = int(existing["id_batch"]) if existing is not None else None
    if existing_batch_id is not None and _batch_is_linked_to_case(
        connection, id_case=id_case, id_batch=existing_batch_id,
        id_periodo=period_id, source_kind=source_kind,
    ):
        return existing_batch_id, 0
    with archived_path.open(
        "r", encoding="utf-8-sig", errors="replace", newline=""
    ) as handle:
        rows = list(csv.reader(handle, delimiter=";"))
    connection.execute("BEGIN IMMEDIATE")
    try:
        batch_id = existing_batch_id or _start_batch(
            connection, community_id=case.community_id, period_id=period_id,
            source_kind=source_kind, archived_path=document.archived_path,
            source_hash=document.sha256,
        )
        _link_batch_to_case(
            connection, id_case=id_case, id_batch=batch_id,
            id_periodo=period_id, source_kind=source_kind,
            source_hash=document.sha256,
        )
        inserted = 0
        headers = {_header(value): index for index, value in enumerate(rows[0] if rows else [])}
        property_column = headers.get("FDENOMINACION")
        name_column = headers.get("NOMBRE")
        coefficient_column = headers.get("COEFICIENTE")
        email_column = headers.get("EMAIL")
        if property_column is None or name_column is None:
            _issue(
                connection, id_case=id_case, id_document=document.id_document,
                code="UNRECOGNIZED_HEADER", field_name="owner_list.headers",
                message="No se reconocen las cabeceras del listado de propietarios",
            )
        else:
            for row_number, row in enumerate(rows[1:], start=2):
                def field(column: int | None):
                    return row[column] if column is not None and column < len(row) else None

                property_code = _clean(field(property_column))
                name = _clean(field(name_column))
                if not property_code and not name:
                    continue
                if not property_code or not name:
                    _issue(
                        connection, id_case=id_case, id_document=document.id_document,
                        code="MISSING_REQUIRED_FIELD",
                        field_name=f"owner_list.row_{row_number}",
                        message="Falta código de vivienda o nombre de propietario",
                    )
                    continue
                coefficient = _number(field(coefficient_column))
                if coefficient_column is None or coefficient is None:
                    _issue(
                        connection, id_case=id_case, id_document=document.id_document,
                        code="MISSING_REQUIRED_FIELD",
                        field_name=f"owner_list.row_{row_number}.coefficient",
                        message="Falta el coeficiente requerido del propietario",
                    )
                    continue
                email = _clean(field(email_column)) or None
                cursor = connection.execute(
                    """INSERT OR IGNORE INTO propietarios
                       (id_comunidad,codigo_vivienda,nombre_propietario,coeficiente,email)
                       VALUES (?,?,?,?,?)""",
                    (case.community_id, property_code, name, coefficient, email),
                )
                inserted += int(bool(cursor.rowcount))
                sources = []
                for key, column, value in (
                    ("property_code", property_column, property_code),
                    ("owner_name", name_column, name),
                    ("coefficient", coefficient_column, coefficient),
                    ("email", email_column, email),
                ):
                    if column is not None:
                        sources.append(_Source(
                            "owner", property_code, key, "CSV",
                            f"{get_column_letter(column + 1)}{row_number}",
                            field(column), value,
                        ))
                _store_sources(connection, batch_id, sources)
        if existing_batch_id is None:
            connection.execute(
                """UPDATE import_batches SET status='validated',validated_at=datetime('now')
                   WHERE id_batch=?""",
                (batch_id,),
            )
        _validate_document_if_clean(connection, document.id_document)
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    return batch_id, inserted


def _reading_grid(path: Path) -> tuple[str, list[list[Any]]]:
    if path.suffix.lower() == ".xls":
        book = xlrd.open_workbook(path)
        sheet = book.sheet_by_index(0)
        return sheet.name, [
            [sheet.cell_value(row, column) for column in range(sheet.ncols)]
            for row in range(sheet.nrows)
        ]
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        return sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


_MONTH_NUMBERS = {
    "ene": 1, "enero": 1,
    "feb": 2, "febrero": 2,
    "mar": 3, "marzo": 3,
    "abr": 4, "abril": 4,
    "may": 5, "mayo": 5,
    "jun": 6, "junio": 6,
    "jul": 7, "julio": 7,
    "ago": 8, "agosto": 8,
    "sep": 9, "sept": 9, "septiembre": 9,
    "oct": 10, "octubre": 10,
    "nov": 11, "noviembre": 11,
    "dic": 12, "diciembre": 12,
}


def _period_header_date(value: Any) -> tuple[int, int] | None:
    text = _clean(value).lower()
    numeric = re.fullmatch(r"(\d{1,2})/(\d{2,4})", text)
    if numeric:
        month, year = int(numeric[1]), int(numeric[2])
        if year < 100:
            year += 2000
        return (year, month) if 1 <= month <= 12 else None
    named = re.fullmatch(r"([a-záéíóú]{3,10})-(\d{2,4})", text)
    if named:
        month_name = unicodedata.normalize("NFKD", named[1])
        month_name = "".join(
            character for character in month_name if not unicodedata.combining(character)
        )
        month = _MONTH_NUMBERS.get(month_name)
        year = int(named[2])
        if year < 100:
            year += 2000
        return (year, month) if month is not None else None
    return None


def _period_header(value: Any) -> bool:
    return _period_header_date(value) is not None


def _owner_key(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", _header(value))


def _import_reading_source(
    connection: sqlite3.Connection,
    *,
    id_case: int,
    case,
    period_id: int,
    path: Path,
    service: str,
) -> tuple[int, int]:
    document = _register_source(
        connection, id_case=id_case, source_path=path,
        document_kind="meter_readings",
    )
    archived_path = _verified_archived_path(document)
    source_kind = "meter_readings"
    existing = _existing_batch(
        connection, case.community_id, document.sha256, source_kind
    )
    existing_batch_id = int(existing["id_batch"]) if existing is not None else None
    if existing_batch_id is not None and _batch_is_linked_to_case(
        connection, id_case=id_case, id_batch=existing_batch_id,
        id_periodo=period_id, source_kind=source_kind,
    ):
        return existing_batch_id, 0
    sheet_name, rows = _reading_grid(archived_path)
    connection.execute("BEGIN IMMEDIATE")
    try:
        batch_id = existing_batch_id or _start_batch(
            connection, community_id=case.community_id, period_id=period_id,
            source_kind=source_kind, archived_path=document.archived_path,
            source_hash=document.sha256,
        )
        _link_batch_to_case(
            connection, id_case=id_case, id_batch=batch_id,
            id_periodo=period_id, source_kind=source_kind,
            source_hash=document.sha256,
        )
        header_row = None
        property_column = None
        period_columns: list[int] = []
        for index, row in enumerate(rows):
            labels = [_header(value) for value in row]
            if "PROPIEDAD" in labels and any(label in {"COD", "CODIGO"} for label in labels):
                header_row = index
                property_column = labels.index("PROPIEDAD")
                dated_columns = [
                    (i, _period_header_date(value))
                    for i, value in enumerate(row)
                    if _period_header(value)
                ]
                period_columns = [i for i, _ in dated_columns]
                break
        if header_row is None or property_column is None or len(period_columns) < 2:
            _issue(
                connection, id_case=id_case, id_document=document.id_document,
                code="UNRECOGNIZED_HEADER", field_name="readings.headers",
                message="No se reconocen vivienda y dos columnas de lectura",
            )
            period_columns = []
        owners = {
            _owner_key(row["codigo_vivienda"]): row
            for row in connection.execute(
                """SELECT id_propietario,codigo_vivienda FROM propietarios
                   WHERE id_comunidad=? AND activo=1""",
                (case.community_id,),
            )
        }
        inserted = 0
        if period_columns:
            initial_column = next(
                (
                    column for column, header_date in dated_columns
                    if header_date == (case.start_date.year, case.start_date.month)
                ),
                None,
            )
            final_column = next(
                (
                    column for column, header_date in dated_columns
                    if header_date == (case.end_date.year, case.end_date.month)
                ),
                None,
            )
            if initial_column is None or final_column is None:
                code = (
                    "INCOMPATIBLE_DATE"
                    if initial_column is None and final_column is None
                    else "MISSING_READING_RANGE"
                )
                _issue(
                    connection, id_case=id_case, id_document=document.id_document,
                    code=code, field_name="readings.period_columns",
                    message="Las columnas de lectura no corresponden al rango del expediente",
                    detected_value=", ".join(str(row[column]) for column in period_columns),
                )
        if period_columns and initial_column is not None and final_column is not None:
            for row_number, row in enumerate(rows[header_row + 1:], start=header_row + 2):
                property_code = _clean(row[property_column] if property_column < len(row) else None)
                if not property_code:
                    continue
                initial_raw = row[initial_column] if initial_column < len(row) else None
                final_raw = row[final_column] if final_column < len(row) else None
                initial = _number(initial_raw)
                final = _number(final_raw)
                owner = owners.get(_owner_key(property_code))
                if owner is None:
                    if initial is not None or final is not None:
                        _issue(
                            connection, id_case=id_case, id_document=document.id_document,
                            code="UNMATCHED_OWNER", field_name=f"reading.{property_code}",
                            message="La vivienda de la lectura no existe en propietarios",
                            detected_value=property_code,
                        )
                    continue
                _store_sources(connection, batch_id, (
                    _Source("reading", f"{property_code}:{service}", "initial", sheet_name,
                            f"{get_column_letter(initial_column + 1)}{row_number}", initial_raw, initial),
                    _Source("reading", f"{property_code}:{service}", "final", sheet_name,
                            f"{get_column_letter(final_column + 1)}{row_number}", final_raw, final),
                ))
                if initial is None or final is None:
                    _issue(
                        connection, id_case=id_case, id_document=document.id_document,
                        code="MISSING_REQUIRED_FIELD",
                        field_name=f"reading.{property_code}.{service}",
                        message="Falta una lectura inicial o final",
                    )
                    continue
                reset = final < initial
                for reading_date, value, state in (
                    (case.start_date, initial, "real"),
                    (case.end_date, final, "contador_averiado" if reset else "real"),
                ):
                    cursor = connection.execute(
                        """INSERT OR IGNORE INTO lecturas_vecino
                           (id_propietario,id_periodo,tipo,fecha_lectura,valor_acumulado,
                            estado,fuente,notas)
                           VALUES (?,?,?,?,?,?,?,?)""",
                        (
                            owner["id_propietario"], period_id, service,
                            reading_date.isoformat(), value, state,
                            str(document.archived_path),
                            "lectura inferior a la anterior; pendiente de estimación aprobada"
                            if reset and reading_date == case.end_date else None,
                        ),
                    )
                    inserted += int(bool(cursor.rowcount))
                if reset:
                    counter_field = "estado_contador_acs" if service == "ACS" else "estado_contador_cal"
                    connection.execute(
                        f"UPDATE propietarios SET {counter_field}='averiado' WHERE id_propietario=?",
                        (owner["id_propietario"],),
                    )
                    _issue(
                        connection, id_case=id_case, id_document=document.id_document,
                        code="COUNTER_RESET",
                        field_name=f"reading.{property_code}.{service}",
                        message="El contador disminuye y requiere una estimación aprobada",
                        detected_value=f"{initial} -> {final}",
                    )
        for owner in owners.values():
            count = connection.execute(
                """SELECT COUNT(*) FROM lecturas_vecino
                   WHERE id_propietario=? AND id_periodo=? AND tipo=?
                     AND fecha_lectura IN (?,?)""",
                (
                    owner["id_propietario"], period_id, service,
                    case.start_date.isoformat(), case.end_date.isoformat(),
                ),
            ).fetchone()[0]
            if count < 2:
                _issue(
                    connection, id_case=id_case, id_document=document.id_document,
                    code="MISSING_READING_RANGE",
                    field_name=f"reading.{owner['codigo_vivienda']}.{service}",
                    message="Faltan las lecturas inicial y final del propietario activo",
                )
        if existing_batch_id is None:
            connection.execute(
                """UPDATE import_batches SET status='validated',validated_at=datetime('now')
                   WHERE id_batch=?""",
                (batch_id,),
            )
        _validate_document_if_clean(connection, document.id_document)
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    return batch_id, inserted


def import_companion_sources(
    connection: sqlite3.Connection,
    *,
    id_case: int,
    owner_list_path: str | Path,
    readings_path: str | Path,
    profile: ExcelProfile,
    actor: str,
    service: str = "ACS",
) -> CompanionImportResult:
    """Importa propietarios y lecturas preservando toda revisión previa."""
    if not actor.strip():
        raise ValueError("El responsable de la importación es obligatorio")
    owner_path = Path(owner_list_path).resolve()
    reading_path = Path(readings_path).resolve()
    for path in (owner_path, reading_path):
        if not path.is_file():
            raise FileNotFoundError(f"No existe la fuente complementaria: {path}")
    case, period_id = _case_context(connection, id_case, profile)
    owner_batch, owner_count = _import_owner_source(
        connection, id_case=id_case, case=case, period_id=period_id,
        path=owner_path,
    )
    reading_batch, reading_count = _import_reading_source(
        connection, id_case=id_case, case=case, period_id=period_id,
        path=reading_path, service=service,
    )
    return CompanionImportResult(
        (owner_batch, reading_batch), period_id, owner_count, reading_count,
        _open_issue_count(connection, id_case),
    )
