"""Análisis no destructivo de fuentes para el alta guiada de comunidades."""

from __future__ import annotations

import hashlib
import json
import os
import re
import sqlite3
import tempfile
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Mapping

from openpyxl import load_workbook

import excel_generator
import expedient_service
import gestor_bd
from excel_profiles import (
    MODULE_REQUIRED_SHEETS,
    canonical_concepts_for_module,
    runtime_profile_path,
    validate_profile_payload,
)
from lector_pdf import extraer_texto
from meter_reading_sources import ReadingObservation, excel_observations


_EXCEL_SUFFIXES = frozenset({".xls", ".xlsx"})
_METER_MODULES = ("ACS", "CALEFACCION")
_SERVICE_LABELS = {
    "ACS": ("acs", "agua caliente"),
    "CALEFACCION": ("calefaccion", "calefacción"),
}
_SERVICE_DECISIONS = {
    "ACS": ("ACS",),
    "CALEFACCION": ("CALEFACCION",),
    "ACS+CALEFACCION": ("ACS", "CALEFACCION"),
    "NO_APLICA": (),
}
_WINDOWS_RESERVED_NAMES = frozenset({
    "CON", "PRN", "AUX", "NUL",
    *(f"COM{number}" for number in range(1, 10)),
    *(f"LPT{number}" for number in range(1, 10)),
})
_WINDOWS_RESERVED_CHARACTERS = frozenset('<>:"/\\|?*')


@dataclass(frozen=True)
class SourceCandidate:
    path: Path
    kind: str
    sha256: str
    headers: tuple[str, ...] = ()
    text: str = ""


@dataclass(frozen=True)
class OnboardingQuestion:
    key: str
    prompt: str
    candidates: tuple[str, ...]
    required: bool


@dataclass(frozen=True)
class ReadingBinding:
    module: str
    column: str
    source_sha256s: tuple[str, ...]
    meter: str = ""
    date: str = ""
    value: str = ""


@dataclass(frozen=True)
class InvoiceEvidence:
    source_sha256: str
    providers: tuple[str, ...]
    periods: tuple[str, ...]
    amounts: tuple[str, ...]
    concepts: tuple[str, ...]


@dataclass(frozen=True)
class ReadingEvidence:
    source_sha256: str
    meters: tuple[str, ...]
    columns: tuple[str, ...]
    dates: tuple[str, ...]
    readings: tuple[str, ...]
    services: tuple[str, ...]
    observations: tuple[ReadingObservation, ...] = ()


@dataclass(frozen=True)
class InvoiceDecision:
    source_sha256: str
    provider: str
    period: str
    amount: str
    concept: str


@dataclass(frozen=True)
class OnboardingConfiguration:
    """Única decisión normalizada que comparten resumen, perfil y publicación."""

    service_decision: str
    active_modules: tuple[str, ...]
    reading_column: str
    reading_bindings: tuple[ReadingBinding, ...]
    source_traces: tuple[tuple[str, str, str], ...]
    invoice_decisions: tuple[InvoiceDecision, ...] = ()
    not_applicable_modules: tuple[str, ...] = ()


@dataclass(frozen=True)
class OnboardingDraft:
    community_code: str
    community_name: str
    sources: tuple[SourceCandidate, ...]
    detected_modules: tuple[str, ...]
    questions: tuple[OnboardingQuestion, ...]
    invoice_evidence: tuple[InvoiceEvidence, ...] = ()
    reading_evidence: tuple[ReadingEvidence, ...] = ()


@dataclass(frozen=True)
class OnboardingResult:
    community_id: int
    period_id: int | None
    case_id: int | None
    profile_path: Path
    template_path: Path
    source_document_ids: tuple[int, ...]


def build_profile_payload(
    draft: OnboardingDraft,
    answers: Mapping[str, str | bool] | OnboardingConfiguration,
) -> dict[str, object]:
    """Build an in-memory, validated-shape profile from explicit answers."""
    configuration = (
        answers
        if isinstance(answers, OnboardingConfiguration)
        else resolve_onboarding_configuration(draft, answers)
    )
    active_modules = list(configuration.active_modules)
    _validate_community_code(draft.community_code)
    key = f"{draft.community_code}_v1"

    required_sheets = ["DATOS"] + [
        sheet
        for module in active_modules
        for sheet in MODULE_REQUIRED_SHEETS[module]
    ]
    return {
        "key": key,
        "version": "1",
        "community_code": draft.community_code,
        "template_relative_path": (
            f"plantillas/comunidades/{draft.community_code}/{key}.xlsx"
        ),
        "active_modules": active_modules,
        "required_sheets": required_sheets,
        "required_formula_cells": [],
        "concepts": [
            concept
            for module in active_modules
            for concept in _concepts_for(module)
        ],
        "onboarding_configuration": _configuration_payload(configuration),
    }


def resolve_onboarding_configuration(
    draft: OnboardingDraft,
    answers: Mapping[str, str | bool],
) -> OnboardingConfiguration:
    """Normaliza respuestas humanas en un contrato único y trazable."""
    service_answer = answers.get("service")
    if isinstance(service_answer, str) and service_answer in _SERVICE_DECISIONS:
        service_decision = service_answer
        active_modules = _SERVICE_DECISIONS[service_answer]
    elif len(draft.detected_modules) == 1:
        # Compatibilidad con borradores construidos antes de que el análisis
        # incorporase siempre la decisión explícita de servicio.
        service_decision = draft.detected_modules[0]
        active_modules = (draft.detected_modules[0],)
    else:
        raise ValueError("Debe confirmarse el servicio o indicar que no aplica")

    unsupported = sorted(set(active_modules).difference(_METER_MODULES))
    if unsupported:
        raise ValueError("Servicio de contador no soportado: " + ", ".join(unsupported))

    not_applicable_modules = _not_applicable_modules(active_modules, answers)
    _require_answers(draft.questions, answers, active_modules=active_modules)
    reading_sources = tuple(
        source for source in draft.sources
        if source.kind.startswith("meter_reading_")
    )
    reading_columns = _confirmed_reading_columns(
        reading_sources, draft.reading_evidence
    )
    if active_modules and not reading_sources:
        raise ValueError("Cada servicio activo necesita una fuente de lecturas")

    bindings = tuple(
        _binding_for(
            module,
            active_modules=active_modules,
            reading_sources=reading_sources,
            reading_evidence=draft.reading_evidence,
            reading_columns=reading_columns,
            questions=draft.questions,
            answers=answers,
        )
        for module in active_modules
    )
    distinct_columns = _unique_labels(binding.column for binding in bindings)
    reading_column = distinct_columns[0] if len(distinct_columns) == 1 else ""
    invoice_decisions = _invoice_decisions(draft.invoice_evidence, answers)
    traces = tuple(
        (source.kind, source.path.name, source.sha256) for source in draft.sources
    )
    return OnboardingConfiguration(
        service_decision=service_decision,
        active_modules=active_modules,
        reading_column=reading_column,
        reading_bindings=bindings,
        source_traces=traces,
        invoice_decisions=invoice_decisions,
        not_applicable_modules=not_applicable_modules,
    )


def _configuration_payload(configuration: OnboardingConfiguration) -> dict[str, object]:
    return {
        "schema_version": 2,
        "service_decision": configuration.service_decision,
        "active_modules": list(configuration.active_modules),
        "reading_column": configuration.reading_column,
        "reading_bindings": [
            {
                "module": binding.module,
                "column": binding.column,
                "meter": binding.meter,
                "date": binding.date,
                "value": binding.value,
                "source_sha256s": list(binding.source_sha256s),
            }
            for binding in configuration.reading_bindings
        ],
        "invoice_decisions": [
            {
                "source_sha256": decision.source_sha256,
                "provider": decision.provider,
                "period": decision.period,
                "amount": decision.amount,
                "concept": decision.concept,
            }
            for decision in configuration.invoice_decisions
        ],
        "not_applicable_modules": list(configuration.not_applicable_modules),
        "sources": [
            {"kind": kind, "name": name, "sha256": sha256}
            for kind, name, sha256 in configuration.source_traces
        ],
    }


def _not_applicable_modules(
    active_modules: tuple[str, ...],
    answers: Mapping[str, str | bool],
) -> tuple[str, ...]:
    traced: list[str] = []
    for key, answer in answers.items():
        if not key.startswith("reading_") or ":" not in key:
            continue
        module = key.rsplit(":", 1)[1]
        if module not in _METER_MODULES:
            continue
        if answer == "NO_APLICA":
            if module in active_modules:
                raise ValueError(f"NO_APLICA no es válido para el módulo activo {module}")
            traced.append(module)
        elif module not in active_modules:
            raise ValueError(f"Sólo el módulo inactivo {module} admite NO_APLICA")
    return _unique_labels(traced)


def _binding_for(
    module: str,
    *,
    active_modules: tuple[str, ...],
    reading_sources: tuple[SourceCandidate, ...],
    reading_evidence: tuple[ReadingEvidence, ...],
    reading_columns: tuple[str, ...],
    questions: tuple[OnboardingQuestion, ...],
    answers: Mapping[str, str | bool],
) -> ReadingBinding:
    key = f"reading_column:{module}"
    selected = answers.get(key)
    if selected is None and len(active_modules) == 1:
        # Compatibility is limited to drafts created before per-module questions.
        selected = answers.get("reading_column")
    requires_explicit = (
        len(active_modules) > 1
        or any(question.key in {key, "reading_column"} for question in questions)
    )
    module_columns = _reading_candidates_for_module(
        reading_evidence, module, "columns"
    )
    admissible_columns = module_columns or reading_columns
    if isinstance(selected, str) and selected.strip():
        column = selected.strip()
    elif not requires_explicit and len(admissible_columns) == 1:
        column = admissible_columns[0]
    else:
        raise ValueError(f"Debe responderse la columna de lectura obligatoria de {module}")
    if column == "NO_APLICA":
        raise ValueError(f"NO_APLICA no es válido para el módulo activo {module}")
    if admissible_columns and column not in admissible_columns:
        raise ValueError(f"La columna de lectura de {module} no consta en la evidencia")

    scoped_evidence = tuple(
        evidence for evidence in reading_evidence
        if (not evidence.services or module in evidence.services)
        and column in evidence.columns
    )
    evidence_sha256s = tuple(
        evidence.source_sha256
        for evidence in scoped_evidence
    )
    source_sha256s = evidence_sha256s or tuple(
        source.sha256 for source in reading_sources
    )
    binding = ReadingBinding(
        module,
        column,
        _unique_labels(source_sha256s),
        meter=_confirmed_reading_detail(
            "reading_meter", module, scoped_evidence, questions, answers
        ),
        date=_confirmed_reading_detail(
            "reading_date", module, scoped_evidence, questions, answers
        ),
        value=_confirmed_reading_detail(
            "reading_value", module, scoped_evidence, questions, answers
        ),
    )
    observations = tuple(item for evidence in scoped_evidence
                         for item in evidence.observations if item.column == column)
    if observations and not any(
        all(not actual or actual == confirmed for actual, confirmed in (
            (item.meter, binding.meter), (item.date, binding.date),
            (item.value, binding.value),
        )) for item in observations
    ):
        raise ValueError(f"La combinación contador, fecha y lectura de {module} no consta en la evidencia")
    return binding


def _confirmed_reading_detail(
    base_key: str,
    module: str,
    evidence_items: tuple[ReadingEvidence, ...],
    questions: tuple[OnboardingQuestion, ...],
    answers: Mapping[str, str | bool],
) -> str:
    attribute = {
        "reading_meter": "meters",
        "reading_date": "dates",
        "reading_value": "readings",
    }[base_key]
    candidates = _reading_candidates_for_module(
        evidence_items, module, attribute
    )
    key = f"{base_key}:{module}"
    answer = answers.get(key)
    if isinstance(answer, str) and answer.strip():
        value = answer.strip()
        if value == "NO_APLICA":
            raise ValueError(f"NO_APLICA no es válido para el módulo activo {module}")
        if candidates and value not in candidates:
            raise ValueError(f"La respuesta {key} no consta en la evidencia")
        return value
    if len(candidates) == 1:
        return candidates[0]
    raise ValueError(f"Debe responderse la decisión obligatoria {key}")


def _invoice_decisions(
    evidence_items: tuple[InvoiceEvidence, ...],
    answers: Mapping[str, str | bool],
) -> tuple[InvoiceDecision, ...]:
    decisions: list[InvoiceDecision] = []
    total = len(evidence_items)
    for index, evidence in enumerate(evidence_items):
        decisions.append(InvoiceDecision(
            source_sha256=evidence.source_sha256,
            provider=_confirmed_evidence_value(
                "invoice_provider", evidence.providers, index, total, answers
            ),
            period=_confirmed_evidence_value(
                "invoice_period", evidence.periods, index, total, answers
            ),
            amount=_confirmed_evidence_value(
                "invoice_amount", evidence.amounts, index, total, answers
            ),
            concept=_confirmed_evidence_value(
                "invoice_concept", evidence.concepts, index, total, answers
            ),
        ))
    return tuple(decisions)


def _confirmed_evidence_value(
    base_key: str,
    candidates: tuple[str, ...],
    index: int,
    total: int,
    answers: Mapping[str, str | bool],
) -> str:
    key = _indexed_question_key(base_key, index, total)
    answer = answers.get(key)
    if isinstance(answer, str) and answer.strip():
        value = answer.strip()
        if candidates and value not in candidates:
            raise ValueError(f"La respuesta {key} no consta en la evidencia")
        return value
    if len(candidates) == 1:
        return candidates[0]
    raise ValueError(f"Debe responderse la decisión obligatoria {key}")


def confirm_onboarding(
    connection: sqlite3.Connection,
    *,
    draft: OnboardingDraft,
    answers: Mapping[str, str | bool],
    project_root: Path,
    archive_root: Path,
    period_name: str | None,
    start_date: date | None,
    end_date: date | None,
    actor: str,
) -> OnboardingResult:
    """Publish one onboarding or compensate every artifact created by it."""
    _validate_period(period_name, start_date, end_date)
    if connection.in_transaction:
        raise RuntimeError("No confirme un alta dentro de una transacción externa")
    if not isinstance(actor, str) or not actor.strip():
        raise ValueError("Debe indicarse quién confirma el alta")

    root = Path(project_root).resolve()
    archives = Path(archive_root).resolve()
    payload = build_profile_payload(draft, answers)
    _verify_source_fingerprints(draft.sources)
    key = str(payload["key"])
    profile_path = runtime_profile_path(key, root)
    template_path = root / str(payload["template_relative_path"])
    if profile_path.exists():
        raise FileExistsError(f"Ya existe el perfil {profile_path}")
    if template_path.exists():
        raise FileExistsError(f"Ya existe la plantilla {template_path}")

    profile_temporary: Path | None = None
    template_temporary: Path | None = None
    profile_created = False
    template_created = False
    community_created = False
    case_id: int | None = None
    period_id: int | None = None
    period_created = False
    document_ids: list[int] = []
    archived_paths: list[Path] = []
    created_directories: list[Path] = []
    existing_community = connection.execute(
        "SELECT id_comunidad FROM comunidades WHERE codigo=?",
        (draft.community_code,),
    ).fetchone()
    community_id: int | None = (
        int(existing_community[0]) if existing_community is not None else None
    )
    period_ids_before: set[int] = set()
    if community_id is not None:
        period_ids_before = {
            int(row[0]) for row in connection.execute(
                "SELECT id_periodo FROM periodos WHERE id_comunidad=?", (community_id,)
            )
        }

    try:
        created_directories.extend(_ensure_directory(profile_path.parent, root))
        created_directories.extend(_ensure_directory(template_path.parent, root))
        with tempfile.NamedTemporaryFile(
            dir=template_path.parent, prefix=f".{template_path.name}.",
            suffix=".tmp.xlsx", delete=False,
        ) as temporary:
            template_temporary = Path(temporary.name)
        layout = excel_generator.create_canonical_community_template(
            template_temporary,
            community_code=draft.community_code,
            community_name=draft.community_name,
            active_modules=tuple(payload["active_modules"]),
        )
        payload["workbook_layout"] = layout
        profile = validate_profile_payload(payload, root)
        _validate_generated_template(template_temporary, profile)

        with tempfile.NamedTemporaryFile(
            mode="w", encoding="utf-8", newline="\n",
            dir=profile_path.parent, prefix=f".{profile_path.name}.",
            suffix=".tmp", delete=False,
        ) as temporary:
            json.dump(payload, temporary, ensure_ascii=False, indent=2)
            temporary.write("\n")
            profile_temporary = Path(temporary.name)
        reloaded_payload = json.loads(profile_temporary.read_text(encoding="utf-8"))
        validate_profile_payload(reloaded_payload, root)

        _publish_exclusive(template_temporary, template_path)
        template_created = True
        _unlink_created_file(template_temporary)
        template_temporary = None
        _publish_exclusive(profile_temporary, profile_path)
        profile_created = True
        _unlink_created_file(profile_temporary)
        profile_temporary = None

        community_id = gestor_bd.obtener_o_crear_comunidad(
            connection, draft.community_code, draft.community_name
        )
        community_created = existing_community is None
        if start_date is not None and end_date is not None:
            case = expedient_service.create_case(
                connection,
                community_id,
                name=str(period_name),
                start_date=start_date,
                end_date=end_date,
            )
            case_id = case.id_case
            period_id = expedient_service.link_case_to_period(connection, case_id)
            period_created = period_id not in period_ids_before
            for source in draft.sources:
                document_ids_before = _source_document_ids(connection, case_id)
                try:
                    document, created = expedient_service.register_source_document(
                        connection,
                        case_id,
                        source_path=source.path,
                        archive_root=archives,
                        document_kind=source.kind,
                    )
                except Exception:
                    for recovered in _new_source_document_rows(
                        connection, case_id, document_ids_before
                    ):
                        document_ids.append(int(recovered[0]))
                        archived_paths.append(Path(recovered[1]))
                    raise
                if created:
                    document_ids.append(document.id_document)
                    archived_paths.append(document.archived_path)

        return OnboardingResult(
            community_id=community_id,
            period_id=period_id,
            case_id=case_id,
            profile_path=profile_path,
            template_path=template_path,
            source_document_ids=tuple(document_ids),
        )
    except Exception:
        _compensate_onboarding(
            connection,
            community_id=community_id if community_created else None,
            period_id=period_id if period_created else None,
            case_id=case_id,
            document_ids=tuple(document_ids),
        )
        for archived_path in reversed(archived_paths):
            _unlink_created_file(archived_path)
            _remove_empty_parents(archived_path.parent, archives)
        if profile_created:
            _unlink_created_file(profile_path)
        if template_created:
            _unlink_created_file(template_path)
        _remove_created_directories(created_directories)
        raise
    finally:
        if profile_temporary is not None:
            _unlink_created_file(profile_temporary)
        if template_temporary is not None:
            _unlink_created_file(template_temporary)


def _validate_period(
    period_name: str | None, start_date: date | None, end_date: date | None
) -> None:
    supplied_dates = (start_date is not None, end_date is not None)
    if any(supplied_dates) and not all(supplied_dates):
        raise ValueError("Las fechas del período deben estar completas o no indicarse")
    if all(supplied_dates) and (not isinstance(period_name, str) or not period_name.strip()):
        raise ValueError("Debe indicarse el nombre del período cuando hay fechas")


def _verify_source_fingerprints(sources: tuple[SourceCandidate, ...]) -> None:
    for source in sources:
        if _sha256(source.path) != source.sha256:
            raise ValueError(f"La fuente cambió desde el análisis: {source.path.name}")


def _validate_generated_template(path: Path, profile) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if workbook.sheetnames != list(profile.required_sheets):
            raise ValueError("La plantilla canónica no coincide con las hojas del perfil")
        for module, table in profile.workbook_layout["tables"].items():
            if module not in profile.active_modules or table["sheet"] not in workbook.sheetnames:
                raise ValueError("El layout de la plantilla no coincide con sus módulos")
    finally:
        workbook.close()


def _compensate_onboarding(
    connection: sqlite3.Connection,
    *,
    community_id: int | None,
    period_id: int | None,
    case_id: int | None,
    document_ids: tuple[int, ...],
) -> None:
    connection.rollback()
    connection.execute("BEGIN IMMEDIATE")
    try:
        if document_ids:
            placeholders = ",".join("?" for _ in document_ids)
            connection.execute(
                f"DELETE FROM source_documents WHERE id_document IN ({placeholders})",
                document_ids,
            )
        if case_id is not None:
            connection.execute("DELETE FROM regularization_cases WHERE id_case=?", (case_id,))
        if period_id is not None:
            connection.execute("DELETE FROM periodos WHERE id_periodo=?", (period_id,))
        if community_id is not None:
            connection.execute("DELETE FROM comunidades WHERE id_comunidad=?", (community_id,))
        connection.commit()
    except Exception:
        connection.rollback()
        raise


def _unlink_created_file(path: Path) -> None:
    try:
        path.unlink()
    except FileNotFoundError:
        pass


def _publish_exclusive(source: Path, destination: Path) -> None:
    """Publish a same-filesystem temporary without ever replacing a destination."""
    os.link(source, destination)


def _source_document_ids(
    connection: sqlite3.Connection, case_id: int
) -> frozenset[int]:
    return frozenset(
        int(row[0]) for row in connection.execute(
            "SELECT id_document FROM source_documents WHERE id_case=?", (case_id,)
        )
    )


def _new_source_document_rows(
    connection: sqlite3.Connection, case_id: int, ids_before: frozenset[int]
):
    return tuple(
        row for row in connection.execute(
            """SELECT id_document, archived_path FROM source_documents
               WHERE id_case=?""",
            (case_id,),
        )
        if int(row[0]) not in ids_before
    )


def _remove_empty_parents(directory: Path, boundary: Path) -> None:
    boundary = boundary.resolve()
    current = directory.resolve()
    while current == boundary or boundary in current.parents:
        try:
            current.rmdir()
        except OSError:
            break
        if current == boundary:
            break
        current = current.parent


def _ensure_directory(directory: Path, boundary: Path) -> tuple[Path, ...]:
    missing: list[Path] = []
    current = directory
    boundary = boundary.resolve()
    while not current.exists() and current != boundary:
        missing.append(current)
        current = current.parent
    directory.mkdir(parents=True, exist_ok=True)
    return tuple(missing)


def _remove_created_directories(directories: list[Path]) -> None:
    for directory in directories:
        try:
            directory.rmdir()
        except FileNotFoundError:
            continue
        except OSError:
            continue


def _require_answers(
    questions: tuple[OnboardingQuestion, ...],
    answers: Mapping[str, str | bool],
    *,
    active_modules: tuple[str, ...],
) -> None:
    for question in questions:
        if not question.required:
            continue
        if ":" in question.key:
            module = question.key.rsplit(":", 1)[1]
            if module in _METER_MODULES and module not in active_modules:
                if answers.get(question.key) == "NO_APLICA":
                    continue
                if question.key in answers:
                    raise ValueError(
                        f"Sólo el módulo inactivo {module} admite NO_APLICA"
                    )
                continue
        answer = answers.get(question.key)
        if (
            not isinstance(answer, str)
            or not answer.strip()
            or (question.candidates and answer not in question.candidates)
        ):
            if question.key == "reading_column" or question.key.startswith("reading_column:"):
                raise ValueError("Debe responderse la columna de lectura obligatoria")
            if question.key == "service":
                raise ValueError("Debe responderse el servicio obligatorio")
            raise ValueError(f"Debe responderse la pregunta obligatoria: {question.prompt}")


def _validate_community_code(community_code: str) -> None:
    if (
        not isinstance(community_code, str)
        or not community_code
        or community_code[-1] in {".", " "}
        or any(character in _WINDOWS_RESERVED_CHARACTERS for character in community_code)
        or any(ord(character) <= 0x1F for character in community_code)
        or community_code.split(".", 1)[0].upper() in _WINDOWS_RESERVED_NAMES
    ):
        raise ValueError("El código de comunidad no permite crear una clave segura")


def _concepts_for(module: str) -> tuple[dict[str, object], ...]:
    return canonical_concepts_for_module(module)


def analyse_sources(
    *,
    community_code: str,
    community_name: str,
    owner_list_path: Path,
    reading_paths: tuple[Path, ...],
    invoice_paths: tuple[Path, ...],
    project_root: Path,
) -> OnboardingDraft:
    """Return a draft from explicit source files without persisting any data."""
    del project_root  # Kept in the public API for later confirmation steps.
    if not reading_paths:
        raise ValueError("Se requiere al menos una fuente de lecturas")

    owner_path = _validated_path(owner_list_path, {".csv"}, "listado de propietarios")
    sources = [_source(owner_path, "owner_list")]
    for reading_path in reading_paths:
        path = _validated_path(reading_path, _EXCEL_SUFFIXES | {".pdf"}, "lectura")
        sources.append(_source(path, "meter_reading_pdf" if path.suffix.lower() == ".pdf" else "meter_reading_excel"))
    for invoice_path in invoice_paths:
        path = _validated_path(invoice_path, {".pdf"}, "factura")
        sources.append(_source(path, "invoice_pdf"))

    readings = tuple(source for source in sources if source.kind.startswith("meter_reading_"))
    invoices = tuple(source for source in sources if source.kind == "invoice_pdf")
    invoice_evidence = _invoice_evidence(invoices)
    reading_evidence = _reading_evidence(readings)
    detected_modules = _detected_modules(invoice_evidence, reading_evidence)
    return OnboardingDraft(
        community_code=community_code,
        community_name=community_name,
        sources=tuple(sources),
        detected_modules=detected_modules,
        questions=_questions(
            readings, detected_modules, invoice_evidence, reading_evidence
        ),
        invoice_evidence=invoice_evidence,
        reading_evidence=reading_evidence,
    )


def _validated_path(path: Path, allowed_suffixes: frozenset[str] | set[str], label: str) -> Path:
    resolved = Path(path)
    if not resolved.is_file():
        raise ValueError(f"No existe el archivo de {label}: {resolved}")
    if resolved.suffix.lower() not in allowed_suffixes:
        allowed = ", ".join(sorted(allowed_suffixes))
        raise ValueError(f"Extensión no permitida para {label}: {resolved.suffix} (se permite {allowed})")
    return resolved


def _source(path: Path, kind: str) -> SourceCandidate:
    headers = _excel_headers(path) if kind == "meter_reading_excel" else ()
    text = extraer_texto(str(path)) if path.suffix.lower() == ".pdf" else ""
    return SourceCandidate(path=path, kind=kind, sha256=_sha256(path), headers=headers, text=text)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(65536), b""):
            digest.update(block)
    return digest.hexdigest()


def _excel_headers(path: Path) -> tuple[str, ...]:
    if path.suffix.lower() == ".xls":
        import xlrd

        sheet = xlrd.open_workbook(path).sheet_by_index(0)
        return tuple(str(sheet.cell_value(0, column)).strip() for column in range(sheet.ncols))
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        return tuple(str(value or "").strip() for value in first_row)
    finally:
        workbook.close()


def _module_candidates(material: str) -> tuple[str, ...]:
    material = material.lower()
    modules = []
    acs_labels = _SERVICE_LABELS["ACS"]
    if any(label in material for label in acs_labels):
        modules.append("ACS")
        for label in acs_labels:
            material = material.replace(label, " ")
    if any(label in material for label in _SERVICE_LABELS["CALEFACCION"]):
        modules.append("CALEFACCION")
    return tuple(modules)


def _detected_modules(
    invoices: tuple[InvoiceEvidence, ...],
    readings: tuple[ReadingEvidence, ...],
) -> tuple[str, ...]:
    return _unique_labels(
        candidate
        for evidence in (*invoices, *readings)
        for candidate in (
            evidence.concepts
            if isinstance(evidence, InvoiceEvidence)
            else evidence.services
        )
        if candidate in _METER_MODULES
    )


def _reading_columns(readings: tuple[SourceCandidate, ...]) -> tuple[str, ...]:
    return _unique_labels(
        header
        for source in readings
        for header in source.headers
        if "lectura" in header.lower()
        and "fecha" not in header.lower()
        and "contador" not in header.lower()
    )


def _confirmed_reading_columns(
    readings: tuple[SourceCandidate, ...],
    evidence_items: tuple[ReadingEvidence, ...],
) -> tuple[str, ...]:
    return _unique_labels((
        *_reading_columns(readings),
        *(column for evidence in evidence_items for column in evidence.columns),
    ))


def _reading_candidates_for_module(
    evidence_items: tuple[ReadingEvidence, ...],
    module: str,
    attribute: str,
) -> tuple[str, ...]:
    return _unique_labels(
        value
        for evidence in evidence_items
        if not evidence.services or module in evidence.services
        for value in getattr(evidence, attribute)
    )


def _invoice_evidence(
    invoices: tuple[SourceCandidate, ...],
) -> tuple[InvoiceEvidence, ...]:
    evidence: list[InvoiceEvidence] = []
    for source in invoices:
        explicit_concepts = _keyed_values(source.text, ("CONCEPTO", "SERVICIO"))
        concepts = _unique_labels(
            canonical
            for value in explicit_concepts
            for canonical in (_canonical_concept(value),)
            if canonical
        )
        if not concepts:
            concepts = _module_candidates(source.text)
        evidence.append(InvoiceEvidence(
            source_sha256=source.sha256,
            providers=_keyed_values(source.text, ("PROVEEDOR", "EMISOR")),
            periods=_keyed_values(source.text, ("PERIODO", "PERÍODO")),
            amounts=_keyed_values(source.text, ("IMPORTE", "TOTAL")),
            concepts=concepts,
        ))
    return tuple(evidence)


def _reading_evidence(
    readings: tuple[SourceCandidate, ...],
) -> tuple[ReadingEvidence, ...]:
    evidence: list[ReadingEvidence] = []
    for source in readings:
        if source.kind == "meter_reading_excel":
            columns = _reading_columns((source,))
            observations = excel_observations(source.path)
            for column in columns:
                selected = tuple(item for item in observations if item.column == column)
                evidence.append(ReadingEvidence(
                    source_sha256=source.sha256,
                    meters=_unique_labels(item.meter for item in selected if item.meter),
                    columns=(column,),
                    dates=_unique_labels(item.date for item in selected if item.date),
                    readings=_unique_labels(item.value for item in selected if item.value),
                    services=_module_candidates(column),
                    observations=selected,
                ))
            continue
        else:
            columns = _keyed_values(source.text, ("COLUMNA",))
            meters = _keyed_values(source.text, ("CONTADOR",))
            dates = _keyed_values(source.text, ("FECHA",))
            values = _keyed_values(source.text, ("LECTURA",))
            material = source.text
        evidence.append(ReadingEvidence(
            source_sha256=source.sha256,
            meters=meters,
            columns=columns,
            dates=dates,
            readings=values,
            services=_module_candidates(material),
        ))
    return tuple(evidence)


def _keyed_values(text: str, labels: tuple[str, ...]) -> tuple[str, ...]:
    alternatives = "|".join(re.escape(label) for label in labels)
    pattern = re.compile(
        rf"(?:^|[;\n])\s*(?:{alternatives})\s*:\s*([^;\n]+)",
        re.IGNORECASE,
    )
    return _unique_labels(match.group(1).strip() for match in pattern.finditer(text))


def _canonical_concept(value: str) -> str:
    modules = _module_candidates(value)
    if len(modules) == 1:
        return modules[0]
    return " ".join(value.split())


def _service_candidates(detected_modules: tuple[str, ...]) -> tuple[str, ...]:
    candidates = list(detected_modules)
    candidates.extend(
        module for module in _METER_MODULES if module not in candidates
    )
    candidates.extend(("ACS+CALEFACCION", "NO_APLICA"))
    return tuple(candidates)


def _questions(
    readings: tuple[SourceCandidate, ...],
    detected_modules: tuple[str, ...],
    invoices: tuple[InvoiceEvidence, ...],
    reading_evidence: tuple[ReadingEvidence, ...],
) -> tuple[OnboardingQuestion, ...]:
    questions: list[OnboardingQuestion] = []
    invoice_fields = (
        ("invoice_provider", "proveedor", "providers"),
        ("invoice_period", "período", "periods"),
        ("invoice_amount", "importe", "amounts"),
        ("invoice_concept", "concepto o servicio", "concepts"),
    )
    for index, evidence in enumerate(invoices):
        for base_key, label, attribute in invoice_fields:
            candidates = getattr(evidence, attribute)
            if len(candidates) != 1:
                questions.append(OnboardingQuestion(
                    key=_indexed_question_key(base_key, index, len(invoices)),
                    prompt=f"Confirma el {label} de la factura seleccionada.",
                    candidates=candidates,
                    required=True,
                ))
    reading_columns = _confirmed_reading_columns(readings, reading_evidence)
    candidate_modules = _METER_MODULES
    for module in candidate_modules:
        module_columns = _reading_candidates_for_module(
            reading_evidence, module, "columns"
        )
        module_values = _reading_candidates_for_module(
            reading_evidence, module, "readings"
        )
        if (
            len(module_columns) != 1
            or len(detected_modules) > 1 and module in detected_modules
            or len(module_values) != 1
        ):
            questions.append(OnboardingQuestion(
                key=f"reading_column:{module}",
                prompt=(
                    f"Selecciona la columna de lectura de {module}."
                    if module_columns or reading_columns
                    else f"Indica la columna o campo de lectura de {module}."
                ),
                candidates=module_columns or reading_columns,
                required=True,
            ))
    questions.extend(_reading_detail_questions(reading_evidence, candidate_modules))
    questions.append(OnboardingQuestion(
        key="service",
        prompt="Confirma el servicio asociado a las lecturas o indica que no aplica.",
        candidates=_service_candidates(detected_modules),
        required=True,
    ))
    return tuple(questions)


def _reading_detail_questions(
    evidence_items: tuple[ReadingEvidence, ...],
    modules: tuple[str, ...],
) -> tuple[OnboardingQuestion, ...]:
    attributes_by_field = {
        "reading_meter": "meters",
        "reading_date": "dates",
        "reading_value": "readings",
    }
    labels = {
        "reading_meter": "contador",
        "reading_date": "fecha de lectura",
        "reading_value": "valor de lectura",
    }
    questions: list[OnboardingQuestion] = []
    for module in modules:
        for base_key, attribute in attributes_by_field.items():
            candidates = _reading_candidates_for_module(
                evidence_items, module, attribute
            )
            if len(candidates) != 1:
                questions.append(OnboardingQuestion(
                    key=f"{base_key}:{module}",
                    prompt=f"Confirma el {labels[base_key]} de {module}.",
                    candidates=candidates,
                    required=True,
                ))
    return tuple(questions)


def _indexed_question_key(base_key: str, index: int, total: int) -> str:
    return base_key if total == 1 else f"{base_key}:{index + 1}"


def _unique_labels(labels) -> tuple[str, ...]:
    unique: list[str] = []
    seen: set[str] = set()
    for label in labels:
        normalized = " ".join(label.split()).casefold()
        if normalized not in seen:
            seen.add(normalized)
            unique.append(label)
    return tuple(unique)
