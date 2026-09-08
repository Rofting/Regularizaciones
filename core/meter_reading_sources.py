"""Read actual meter observations without losing their row/column relationship."""

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
import unicodedata

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def label(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return format(value, ".12g")
    return str(value).strip()


def normalized(value) -> str:
    return "".join(character for character in unicodedata.normalize("NFKD", label(value))
                   if not unicodedata.combining(character)).casefold()


def services(value) -> tuple[str, ...]:
    text = normalized(value)
    return tuple(module for module, names in (
        ("ACS", ("acs", "agua caliente")),
        ("CALEFACCION", ("calefaccion",)),
    ) if any(name in text for name in names))


@dataclass(frozen=True)
class ReadingObservation:
    column: str
    meter: str
    date: str
    value: str
    property_code: str = ""
    sheet: str = ""
    value_cell: str = ""
    meter_cell: str = ""
    date_cell: str = ""


def excel_grid(path: Path):
    if path.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(path)
        sheet = book.sheet_by_index(0)
        rows = []
        for index in range(sheet.nrows):
            rows.append([
                xlrd.xldate_as_datetime(cell.value, book.datemode)
                if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                for cell in sheet.row(index)
            ])
        return sheet.name, rows
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        return sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def excel_observations(path: Path) -> tuple[ReadingObservation, ...]:
    sheet, rows = excel_grid(path)
    if not rows:
        return ()
    headers = [label(value) for value in rows[0]]
    names = [normalized(value) for value in headers]
    property_columns = [index for index, name in enumerate(names)
                        if name in {"vivienda", "propiedad", "fdenominacion", "codigo vivienda"}]

    def related_column(index, field):
        candidates = [position for position, name in enumerate(names) if field in name]
        module = services(headers[index])
        scoped = [position for position in candidates if services(headers[position]) == module]
        candidates = scoped or [position for position in candidates if not services(headers[position])]
        boundary = next((part for part in ("inicial", "final") if part in names[index]), None)
        if boundary:
            scoped = [position for position in candidates if boundary in names[position]]
            candidates = scoped or [position for position in candidates
                                    if not any(part in names[position] for part in ("inicial", "final"))]
        return candidates[0] if len(candidates) == 1 else None

    observations = []
    for index, name in enumerate(names):
        if "lectura" not in name or "fecha" in name or "contador" in name:
            continue
        meter_column = related_column(index, "contador")
        date_column = related_column(index, "fecha")
        for number, row in enumerate(rows[1:], start=2):
            def cell(position):
                return label(row[position]) if position is not None and position < len(row) else ""

            if not any(cell(position) for position in (index, meter_column, date_column)):
                continue
            observations.append(ReadingObservation(
                column=headers[index], meter=cell(meter_column), date=cell(date_column),
                value=cell(index),
                property_code=cell(property_columns[0]) if len(property_columns) == 1 else "",
                sheet=sheet, value_cell=f"{get_column_letter(index + 1)}{number}",
                meter_cell=f"{get_column_letter(meter_column + 1)}{number}" if meter_column is not None else "",
                date_cell=f"{get_column_letter(date_column + 1)}{number}" if date_column is not None else "",
            ))
    return tuple(observations)
