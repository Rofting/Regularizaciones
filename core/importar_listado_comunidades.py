"""
importar_listado_comunidades.py
================================
Lee LISTADO_COMUNIDADES.xlsx e inserta en la BD todas las comunidades
que aún no existan. Las que ya están registradas se conservan sin cambios
(solo se actualiza el CIF si faltaba).

LÓGICA:
  - NO borra comunidades existentes
  - NO toca propietarios ni lecturas
  - SÍ inserta comunidades nuevas (código + nombre + CIF)
  - SÍ actualiza CIF vacíos de comunidades ya registradas

USO:
    python importar_listado_comunidades.py
    python importar_listado_comunidades.py --listado ../LISTADO_COMUNIDADES.xlsx
    python importar_listado_comunidades.py --bd ../data/gestion.db --listado otro.xlsx

COLUMNAS ESPERADAS EN EL EXCEL (en cualquier orden):
    num_empresa / codigo   → código numérico de la comunidad
    Razón social / nombre  → nombre completo
    NIF / CIF              → CIF de la comunidad
"""

import sys
import os
import argparse
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from gestor_bd import conectar, crear_bd, obtener_o_crear_comunidad

BASE_DIR       = Path(__file__).parent.parent
RUTA_BD        = BASE_DIR / "data" / "gestion.db"
RUTA_LISTADO   = BASE_DIR / "LISTADO_COMUNIDADES.xlsx"


# ---------------------------------------------------------------------------
# FUNCIÓN PRINCIPAL
# ---------------------------------------------------------------------------

def importar_listado(ruta_listado: str, ruta_bd: str,
                     solo_listar: bool = False,
                     verbose: bool = True) -> dict:
    """
    Importa comunidades desde el LISTADO_COMUNIDADES.xlsx a la BD.

    Args:
        ruta_listado:  Ruta al Excel con el listado maestro
        ruta_bd:       Ruta a gestion.db
        solo_listar:   Si True, imprime qué haría pero no modifica la BD
        verbose:       Mostrar progreso por pantalla

    Returns:
        {ok, nuevas, ya_existentes, errores, lista_nuevas}
    """
    if not os.path.exists(ruta_listado):
        return {"ok": False, "error": f"Archivo no encontrado: {ruta_listado}"}

    # Crear BD si no existe
    if not os.path.exists(ruta_bd):
        if verbose:
            print(f"  BD no encontrada, creando en: {ruta_bd}")
        crear_bd(ruta_bd)

    con = conectar(ruta_bd)

    # ── LEER EXCEL ──────────────────────────────────────────────────────────
    wb = openpyxl.load_workbook(ruta_listado, read_only=True, data_only=True)
    ws = wb.active

    # Detectar cabecera (fila 1)
    primera_fila = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    cabecera = [str(v).strip().lower() if v is not None else "" for v in primera_fila]

    col_codigo = next(
        (i for i, h in enumerate(cabecera) if any(k in h for k in ("num", "codigo", "empresa"))),
        0
    )
    col_nombre = next(
        (i for i, h in enumerate(cabecera) if any(k in h for k in ("raz", "nombre", "social"))),
        1
    )
    col_cif = next(
        (i for i, h in enumerate(cabecera) if any(k in h for k in ("nif", "cif"))),
        2
    )

    if verbose:
        print(f"\n  Columnas detectadas: código[{col_codigo}], nombre[{col_nombre}], CIF[{col_cif}]")

    # ── PROCESAR FILAS ──────────────────────────────────────────────────────
    nuevas       = 0
    ya_existentes = 0
    errores      = 0
    lista_nuevas = []

    for nfila, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not fila or fila[col_codigo] is None:
            continue

        codigo = str(fila[col_codigo]).strip()
        nombre = str(fila[col_nombre]).strip() if fila[col_nombre] else f"COMUNIDAD {codigo}"
        cif    = str(fila[col_cif]).strip()    if len(fila) > col_cif and fila[col_cif] else None

        if not codigo or codigo.lower() in ("none", "nan", ""):
            continue

        # ¿Ya existe?
        existe = con.execute(
            "SELECT id_comunidad, cif FROM comunidades WHERE codigo=?", (codigo,)
        ).fetchone()

        if existe:
            ya_existentes += 1
            # Actualizar CIF si faltaba
            if cif and (not existe["cif"] or existe["cif"].strip() == ""):
                if not solo_listar:
                    con.execute(
                        "UPDATE comunidades SET cif=? WHERE codigo=?",
                        (cif, codigo)
                    )
                if verbose:
                    print(f"  🔄 {codigo} — CIF actualizado: {cif}")
            continue

        # ── NUEVA COMUNIDAD ─────────────────────────────────────────────────
        if solo_listar:
            print(f"  [NUEVA] {codigo:>6} — {nombre[:55]:<55}  CIF: {cif or 'sin CIF'}")
            nuevas += 1
            lista_nuevas.append(codigo)
            continue

        try:
            id_com = obtener_o_crear_comunidad(con, codigo, nombre, cif=cif)
            if verbose:
                print(f"  ➕ {codigo:>6} — {nombre[:55]:<55}  (id={id_com})")
            nuevas += 1
            lista_nuevas.append(codigo)
        except Exception as e:
            if verbose:
                print(f"  ❌ {codigo}: {e}")
            errores += 1

    if not solo_listar:
        con.commit()

    con.close()
    wb.close()

    return {
        "ok":            errores == 0,
        "nuevas":        nuevas,
        "ya_existentes": ya_existentes,
        "errores":       errores,
        "lista_nuevas":  lista_nuevas,
    }


# ---------------------------------------------------------------------------
# EJECUCIÓN DIRECTA
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Importa todas las comunidades del LISTADO a la BD"
    )
    parser.add_argument("--listado", default=str(RUTA_LISTADO),
                        help="Ruta al LISTADO_COMUNIDADES.xlsx")
    parser.add_argument("--bd",      default=str(RUTA_BD),
                        help="Ruta a gestion.db")
    parser.add_argument("--solo-listar", action="store_true",
                        help="Solo mostrar qué se haría, sin modificar la BD")
    args = parser.parse_args()

    print(f"\n{'='*60}")
    print(f"  IMPORTAR LISTADO DE COMUNIDADES")
    print(f"{'='*60}")
    print(f"  Listado: {args.listado}")
    print(f"  BD:      {args.bd}")
    if args.solo_listar:
        print(f"  Modo:    SOLO LISTAR (sin cambios en BD)")
    print()

    resultado = importar_listado(
        ruta_listado=args.listado,
        ruta_bd=args.bd,
        solo_listar=args.solo_listar,
    )

    print(f"\n{'─'*60}")
    if resultado.get("ok"):
        print(f"  ✅ Completado")
        print(f"     Nuevas comunidades:  {resultado['nuevas']}")
        print(f"     Ya registradas:      {resultado['ya_existentes']}")
    else:
        print(f"  ❌ Error: {resultado.get('error', '')}")
        print(f"     Nuevas:    {resultado.get('nuevas', 0)}")
        print(f"     Errores:   {resultado.get('errores', 0)}")
