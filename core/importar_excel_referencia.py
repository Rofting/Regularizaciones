"""Importa libros económicos validados conservando cada celda de origen."""

from __future__ import annotations

import argparse
import calendar
import hashlib
import re
import sqlite3
import unicodedata
from collections.abc import Callable
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

try:
    from .reference_models import ConceptReference, ExerciseReference, SourceValue
except ImportError:  # ejecución directa desde core/ en la aplicación existente
    from reference_models import ConceptReference, ExerciseReference, SourceValue


ProgressCallback = Callable[[str, dict], None]


class ReferenceValidationError(ValueError):
    """El libro no contiene un dato inequívoco o sus totales no son coherentes."""


def _emit(progress: ProgressCallback | None, stage: str, **details) -> None:
    if progress:
        progress(stage, details)


def _normalize_label(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(character for character in text if not unicodedata.combining(character))
    text = text.upper().replace("�", " ")
    return " ".join(re.sub(r"[^A-Z0-9]+", " ", text).split())


def _unique_cell(
    sheet,
    needle: str,
    *,
    min_row: int = 1,
    max_row: int | None = None,
    exact: bool = False,
):
    normalized_needle = _normalize_label(needle)
    matches = []
    last_row = min(max_row or sheet.max_row, sheet.max_row)
    for row in sheet.iter_rows(min_row=min_row, max_row=last_row):
        for cell in row:
            normalized_value = _normalize_label(cell.value)
            if (
                normalized_value == normalized_needle
                if exact
                else normalized_needle in normalized_value
            ):
                matches.append(cell)
    if not matches:
        raise ReferenceValidationError(
            f"No se encuentra la etiqueta '{needle}' en {sheet.title}"
        )
    if len(matches) > 1:
        addresses = ", ".join(cell.coordinate for cell in matches)
        raise ReferenceValidationError(
            f"Etiqueta duplicada '{needle}' en {sheet.title}: {addresses}"
        )
    return matches[0]


def _paired_cell(formula_sheet, value_sheet, address: str):
    return formula_sheet[address], value_sheet[address]


def _to_decimal(value, *, field: str) -> Decimal:
    if isinstance(value, str) and value.startswith("#"):
        raise ReferenceValidationError(f"{field} contiene el error de Excel {value}")
    try:
        return Decimal(str(value)).quantize(Decimal("0.000001"), ROUND_HALF_UP)
    except (InvalidOperation, TypeError, ValueError):
        raise ReferenceValidationError(f"{field} no es numérico: {value!r}") from None


def _to_cents(value, *, field: str) -> int:
    amount = _to_decimal(value, field=field).quantize(Decimal("0.01"), ROUND_HALF_UP)
    return int(amount * 100)


MONTHS = {
    "ENERO": 1,
    "FEBRERO": 2,
    "MARZO": 3,
    "ABRIL": 4,
    "MAYO": 5,
    "JUNIO": 6,
    "JULIO": 7,
    "AGOSTO": 8,
    "SEPTIEMBRE": 9,
    "OCTUBRE": 10,
    "NOVIEMBRE": 11,
    "DICIEMBRE": 12,
}


def _parse_period(value) -> tuple[str, date, date]:
    text = _normalize_label(value)
    explicit = re.search(
        r"(\d{2}) (\d{2}) (\d{4}).*?(\d{2}) (\d{2}) (\d{4})", text
    )
    if explicit:
        start = date(int(explicit[3]), int(explicit[2]), int(explicit[1]))
        end = date(int(explicit[6]), int(explicit[5]), int(explicit[4]))
        return f"{start.year}-{end.year}", start, end

    named = re.search(
        r"(" + "|".join(MONTHS) + r") (\d{4}).*?("
        + "|".join(MONTHS) + r") (\d{4})",
        text,
    )
    if named:
        start_month, start_year = MONTHS[named[1]], int(named[2])
        end_month, end_year = MONTHS[named[3]], int(named[4])
        start = date(start_year, start_month, 1)
        end = date(end_year, end_month, calendar.monthrange(end_year, end_month)[1])
        return f"{start.year}-{end.year}", start, end

    raise ReferenceValidationError(f"Periodo no reconocido: {value!r}")


def _source(
    formula_sheet,
    value_sheet,
    address: str,
    *,
    entity_type: str,
    entity_key: str,
    field_name: str,
) -> SourceValue:
    formula_cell, value_cell = _paired_cell(formula_sheet, value_sheet, address)
    return SourceValue(
        entity_type=entity_type,
        entity_key=entity_key,
        field_name=field_name,
        sheet_name=formula_sheet.title,
        cell_address=address,
        raw_value=str(formula_cell.value),
        normalized_value=str(value_cell.value),
    )


def _concept_from_column(
    formula_sheet,
    value_sheet,
    *,
    concept_key: str,
    column: int,
    actual_row: int,
    billed_row: int,
    difference_row: int,
    unit_row: int,
) -> ConceptReference:
    addresses = {
        "actual_cents": value_sheet.cell(actual_row, column).coordinate,
        "billed_cents": value_sheet.cell(billed_row, column).coordinate,
        "difference_cents": value_sheet.cell(difference_row, column).coordinate,
        "billed_unit_price": value_sheet.cell(unit_row, 6).coordinate,
        "actual_unit_price": value_sheet.cell(unit_row, 7).coordinate,
    }
    values = {
        name: value_sheet[address].value for name, address in addresses.items()
    }
    actual_cents = _to_cents(values["actual_cents"], field=f"{concept_key} coste real")
    billed_cents = _to_cents(values["billed_cents"], field=f"{concept_key} cobrado")
    difference_cents = _to_cents(
        values["difference_cents"], field=f"{concept_key} diferencia"
    )
    if abs((billed_cents - actual_cents) - difference_cents) > 1:
        raise ReferenceValidationError(
            f"La diferencia de {concept_key} no coincide con cobrado - coste real"
        )

    sources = tuple(
        _source(
            formula_sheet,
            value_sheet,
            address,
            entity_type="concept",
            entity_key=concept_key,
            field_name=field_name,
        )
        for field_name, address in addresses.items()
    )
    return ConceptReference(
        concept_key=concept_key,
        billed_cents=billed_cents,
        actual_cents=actual_cents,
        billed_unit_price=_to_decimal(
            values["billed_unit_price"], field=f"{concept_key} precio cobrado"
        ),
        actual_unit_price=_to_decimal(
            values["actual_unit_price"], field=f"{concept_key} precio real"
        ),
        source_values=sources,
    )


def parse_reference_workbook(path: str | Path) -> ExerciseReference:
    path = Path(path)
    if not path.is_file():
        raise ReferenceValidationError(f"No existe el Excel: {path}")

    formula_book = load_workbook(path, data_only=False, read_only=True)
    value_book = load_workbook(path, data_only=True, read_only=True)
    required = ("DATOS", "LECTURAS ACS M3", "ANALISIS")
    missing = [name for name in required if name not in formula_book.sheetnames]
    if missing:
        raise ReferenceValidationError(f"Falta la hoja {', '.join(missing)}")

    formula_data, value_data = formula_book["DATOS"], value_book["DATOS"]
    formula_analysis = formula_book["ANALISIS"]
    value_analysis = value_book["ANALISIS"]

    period_label = _unique_cell(formula_data, "PERIODO")
    period_name, start_date, end_date = _parse_period(value_data[period_label.coordinate].value)
    homes_label = _unique_cell(formula_data, "VIVIENDAS")
    property_value_cell = None
    for column in range(homes_label.column + 1, formula_data.max_column + 1):
        candidate = value_data.cell(homes_label.row, column)
        if isinstance(candidate.value, (int, float)):
            property_value_cell = candidate
            break
    if property_value_cell is None:
        raise ReferenceValidationError("No se encuentra el número de viviendas")
    property_count = int(property_value_cell.value)

    result_anchor = _unique_cell(formula_analysis, "RESULTADO DEL ANALISIS")
    result_min = result_anchor.row
    result_max = min(result_anchor.row + 20, formula_analysis.max_row)
    total_label = _unique_cell(
        formula_analysis, "TOTAL", min_row=result_min, max_row=result_max, exact=True
    )
    billed_label = _unique_cell(
        formula_analysis,
        "IMPORTE COBRADO",
        min_row=result_min,
        max_row=result_max,
        exact=True,
    )
    difference_label = _unique_cell(
        formula_analysis,
        "DIFERENCIA",
        min_row=result_min,
        max_row=result_max,
        exact=True,
    )
    fixed_header = _unique_cell(
        formula_analysis, "CUOTA FIJA", min_row=result_min, max_row=total_label.row
    )
    variable_header = _unique_cell(
        formula_analysis, "CUOTA VARIABLE", min_row=result_min, max_row=total_label.row
    )
    fixed_unit_label = _unique_cell(
        formula_analysis,
        "CUOTA FIJA VIVIENDA MES",
        min_row=result_min,
        max_row=result_max,
    )
    variable_unit_label = _unique_cell(
        formula_analysis,
        "CUOTA VARIABLE M3 CONSUMIDO",
        min_row=result_min,
        max_row=result_max,
    )

    fixed = _concept_from_column(
        formula_analysis,
        value_analysis,
        concept_key="acs_fixed",
        column=fixed_header.column,
        actual_row=total_label.row,
        billed_row=billed_label.row,
        difference_row=difference_label.row,
        unit_row=fixed_unit_label.row,
    )
    variable = _concept_from_column(
        formula_analysis,
        value_analysis,
        concept_key="acs_variable",
        column=variable_header.column,
        actual_row=total_label.row,
        billed_row=billed_label.row,
        difference_row=difference_label.row,
        unit_row=variable_unit_label.row,
    )

    total_actual_address = value_analysis.cell(total_label.row, 7).coordinate
    total_billed_address = value_analysis.cell(billed_label.row, 7).coordinate
    exercise_sources = (
        _source(
            formula_data,
            value_data,
            "A1",
            entity_type="exercise",
            entity_key=period_name,
            field_name="community_name",
        ),
        _source(
            formula_data,
            value_data,
            period_label.coordinate,
            entity_type="exercise",
            entity_key=period_name,
            field_name="period",
        ),
        _source(
            formula_data,
            value_data,
            property_value_cell.coordinate,
            entity_type="exercise",
            entity_key=period_name,
            field_name="property_count",
        ),
        _source(
            formula_analysis,
            value_analysis,
            total_actual_address,
            entity_type="exercise",
            entity_key=period_name,
            field_name="total_actual_cents",
        ),
        _source(
            formula_analysis,
            value_analysis,
            total_billed_address,
            entity_type="exercise",
            entity_key=period_name,
            field_name="total_billed_cents",
        ),
    )
    return ExerciseReference(
        community_name=str(value_data["A1"].value).strip(),
        period_name=period_name,
        property_count=property_count,
        start_date=start_date,
        end_date=end_date,
        total_billed_cents=_to_cents(
            value_analysis[total_billed_address].value, field="total cobrado"
        ),
        total_actual_cents=_to_cents(
            value_analysis[total_actual_address].value, field="total coste real"
        ),
        concepts=(fixed, variable),
        source_values=exercise_sources,
    )


def import_reference_workbook(
    connection: sqlite3.Connection,
    id_comunidad: int,
    path: str | Path,
    *,
    progress: ProgressCallback | None = None,
) -> tuple[int, int]:
    path = Path(path).resolve()
    _emit(progress, "reading_reference", file=path.name)
    reference = parse_reference_workbook(path)
    _emit(progress, "validating_reference", period=reference.period_name)
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()

    existing = connection.execute(
        """
        SELECT id_batch, id_periodo FROM import_batches
        WHERE id_comunidad=? AND source_sha256=?
        """,
        (id_comunidad, source_hash),
    ).fetchone()
    if existing:
        _emit(progress, "reference_already_imported", period=reference.period_name)
        return int(existing[0]), int(existing[1])

    if connection.in_transaction:
        connection.commit()
    connection.execute("BEGIN IMMEDIATE")
    try:
        period = connection.execute(
            "SELECT id_periodo FROM periodos WHERE id_comunidad=? AND nombre=?",
            (id_comunidad, reference.period_name),
        ).fetchone()
        if period:
            id_periodo = int(period[0])
            connection.execute(
                """
                UPDATE periodos SET fecha_inicio=?, fecha_fin=?, estado='cerrado'
                WHERE id_periodo=?
                """,
                (reference.start_date.isoformat(), reference.end_date.isoformat(), id_periodo),
            )
        else:
            cursor = connection.execute(
                """
                INSERT INTO periodos
                    (id_comunidad, nombre, fecha_inicio, fecha_fin, estado)
                VALUES (?, ?, ?, ?, 'cerrado')
                """,
                (
                    id_comunidad,
                    reference.period_name,
                    reference.start_date.isoformat(),
                    reference.end_date.isoformat(),
                ),
            )
            id_periodo = int(cursor.lastrowid)

        cursor = connection.execute(
            """
            INSERT INTO import_batches
                (id_comunidad, id_periodo, source_kind, source_path,
                 source_sha256, status)
            VALUES (?, ?, 'excel_reference', ?, ?, 'processing')
            """,
            (id_comunidad, id_periodo, str(path), source_hash),
        )
        id_batch = int(cursor.lastrowid)

        all_sources = list(reference.source_values)
        for concept in reference.concepts:
            all_sources.extend(concept.source_values)
        connection.executemany(
            """
            INSERT INTO source_values
                (id_batch, entity_type, entity_key, field_name, sheet_name,
                 cell_address, raw_value, normalized_value)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    id_batch,
                    source.entity_type,
                    source.entity_key,
                    source.field_name,
                    source.sheet_name,
                    source.cell_address,
                    source.raw_value,
                    source.normalized_value,
                )
                for source in all_sources
            ],
        )

        fixed = next(item for item in reference.concepts if item.concept_key == "acs_fixed")
        variable = next(
            item for item in reference.concepts if item.concept_key == "acs_variable"
        )
        connection.execute(
            """
            INSERT INTO config_suministro
                (id_comunidad, id_periodo, tipo_suministro, metodo_reparto,
                 precio_variable_facturado, precio_fijo_facturado,
                 precio_variable_real, precio_fijo_real)
            VALUES (?, ?, 'ACS', 'contador', ?, ?, ?, ?)
            ON CONFLICT(id_comunidad, id_periodo, tipo_suministro) DO UPDATE SET
                metodo_reparto=excluded.metodo_reparto,
                precio_variable_facturado=excluded.precio_variable_facturado,
                precio_fijo_facturado=excluded.precio_fijo_facturado,
                precio_variable_real=excluded.precio_variable_real,
                precio_fijo_real=excluded.precio_fijo_real
            """,
            (
                id_comunidad,
                id_periodo,
                float(variable.billed_unit_price),
                float(fixed.billed_unit_price),
                float(variable.actual_unit_price),
                float(fixed.actual_unit_price),
            ),
        )
        connection.execute(
            """
            UPDATE import_batches
            SET status='validated', validated_at=datetime('now')
            WHERE id_batch=?
            """,
            (id_batch,),
        )
        connection.commit()
    except Exception:
        connection.rollback()
        raise

    _emit(progress, "reference_imported", period=reference.period_name)
    return id_batch, id_periodo


def _main() -> int:
    parser = argparse.ArgumentParser(description="Inspecciona Excel de referencia")
    parser.add_argument("--inspect", nargs="+", type=Path, required=True)
    args = parser.parse_args()
    for path in args.inspect:
        reference = parse_reference_workbook(path)
        print(
            f"{path.name}: periodo={reference.period_name}; "
            f"propiedades={reference.property_count}; conceptos={len(reference.concepts)}"
        )
        for concept in reference.concepts:
            cells = ",".join(source.cell_address for source in concept.source_values)
            print(
                f"  {concept.concept_key}: cobrado={concept.billed_cents}; "
                f"real={concept.actual_cents}; celdas={cells}"
            )
    return 0


if __name__ == "__main__":
    raise SystemExit(_main())
