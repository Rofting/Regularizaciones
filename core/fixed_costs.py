"""Gastos fijos mensuales del estudio (hoja OTROS GASTOS del modelo común).

El modelo reparte cuatro gastos periódicos que no siempre llegan como factura
del expediente: el servicio de lecturas de contadores de ACS y de calefacción
y el mantenimiento de placas solares y de la sala de calderas. Son de cada
comunidad, así que el modelo común los trae vacíos y aquí se guardan por
período para escribirlos al generar el Excel.
"""

from __future__ import annotations

import sqlite3
import unicodedata
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from cuotas_servicio import _importe


SHEET = "OTROS GASTOS"


@dataclass(frozen=True)
class FixedCost:
    key: str
    label: str
    cell: str
    sheet_label: str  # texto de la columna A que confirma que la fila es la esperada


FIXED_COSTS = (
    FixedCost("fixed_cost_meter_reading_acs", "Lecturas de contadores de ACS", "B6",
              "LECTURAS CONTADORES ACS"),
    FixedCost("fixed_cost_meter_reading_heating", "Lecturas de contadores de calefacción", "B7",
              "LECTURAS CONTADORES CALEF"),
    FixedCost("fixed_cost_solar_maintenance", "Mantenimiento de placas solares", "B8",
              "MANTENIMIENTO PLACAS SOLARES"),
    FixedCost("fixed_cost_boiler_maintenance", "Mantenimiento de la sala de calderas", "B9",
              "MANTENIMIENTO SALA CALDERAS"),
)
FIXED_COST_CELLS = {SHEET: tuple(cell for item in FIXED_COSTS for cell in (item.cell, "C" + item.cell[1:]))}


@dataclass(frozen=True)
class UnconfirmedCost:
    cost: FixedCost
    value: object
    cell: str


class TemplateCostReviewError(ValueError):
    """No se pudo comprobar la plantilla; no se puede dar por revisada."""


def unconfirmed_template_costs(template: Path, values: dict[str, Decimal]) -> tuple[UnconfirmedCost, ...]:
    """Detecta importes heredados en las filas conocidas, sin alterar el libro."""
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(template, read_only=True, data_only=False)
    except Exception as error:
        raise TemplateCostReviewError(f"No se pueden revisar los gastos de la plantilla {template.name}: {error}") from error
    try:
        if SHEET not in workbook.sheetnames:
            return ()
        sheet = workbook[SHEET]
        findings = []
        for item in FIXED_COSTS:
            if item.key in values:
                continue
            row = int(item.cell[1:])
            if not _normalise(sheet[f"A{row}"].value).startswith(item.sheet_label):
                continue
            for address in (item.cell, f"C{row}"):
                raw = sheet[address].value
                if raw is None or str(raw).strip() == "":
                    continue
                if address.startswith("C") and str(raw).replace(" ", "").replace("$", "").upper() == f"={item.cell}*12":
                    continue
                try:
                    if _importe(str(raw), "El importe heredado") == 0:
                        continue
                except ValueError:
                    pass  # Una fórmula o un texto sin interpretar también requiere revisión.
                findings.append(UnconfirmedCost(item, raw, address))
        return tuple(findings)
    finally:
        workbook.close()


def inherited_cost_message(findings: tuple[UnconfirmedCost, ...]) -> str:
    details = []
    for item in findings:
        try:
            unit = "€/año" if item.cell.startswith("C") else "€/mes"
            value = f"{_importe(str(item.value), 'El importe heredado'):.2f}".replace(".", ",") + " " + unit
        except ValueError:
            value = str(item.value)
        details.append(f"{SHEET}!{item.cell}: {item.cost.label} ({value})")
    return (
        "La plantilla contiene gastos sin confirmar para esta comunidad y período: "
        + "; ".join(details)
        + ". Abre «Gastos fijos» y guarda los importes correctos; un campo vacío confirma que no hay gasto."
    )


def _normalise(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return " ".join(text.upper().split())


def matching_fixed_cost_cells(workbook) -> dict[str, tuple[str, ...]]:
    """Sólo estas filas del modelo pueden cambiar al confirmar gastos."""
    if SHEET not in workbook.sheetnames:
        return {}
    sheet = workbook[SHEET]
    cells = []
    for item in FIXED_COSTS:
        row = item.cell[1:]
        if _normalise(sheet[f"A{row}"].value).startswith(item.sheet_label):
            cells.extend((item.cell, f"C{row}"))
    return {SHEET: tuple(cells)}


def load_fixed_costs(
    connection: sqlite3.Connection, community_id: int, period_id: int,
) -> dict[str, Decimal]:
    keys = tuple(item.key for item in FIXED_COSTS)
    rows = connection.execute(
        f"""SELECT parameter_key, numeric_value FROM period_parameters
            WHERE id_comunidad=? AND id_periodo=? AND parameter_key IN ({",".join("?" * len(keys))})""",
        (community_id, period_id, *keys),
    ).fetchall()
    return {
        str(row[0]): Decimal(str(row[1])) for row in rows if row[1] is not None
    }


def save_fixed_costs(
    connection: sqlite3.Connection,
    community_id: int,
    period_id: int,
    values: dict[str, object],
) -> dict[str, Decimal]:
    """Guarda importes mensuales; vacío confirma cero y sustituye valores heredados."""
    period = connection.execute(
        "SELECT id_comunidad FROM periodos WHERE id_periodo=?", (period_id,),
    ).fetchone()
    if period is None or int(period[0]) != int(community_id):
        raise ValueError("El período no corresponde a esta comunidad")
    known = {item.key for item in FIXED_COSTS}
    saved: dict[str, Decimal] = {}
    for key, raw in values.items():
        if key not in known:
            raise ValueError(f"Gasto fijo desconocido: {key}")
        text = str(raw if raw is not None else "").strip()
        amount = _importe(text, "El importe mensual") if text else Decimal(0)
        if amount < 0:
            raise ValueError("Los gastos fijos no pueden ser negativos")
        saved[key] = amount
    for key, amount in saved.items():
        connection.execute(
            """INSERT INTO period_parameters
                   (id_comunidad,id_periodo,parameter_key,numeric_value,unit)
               VALUES (?,?,?,?,'EUR/mes')
               ON CONFLICT(id_comunidad,id_periodo,parameter_key)
               DO UPDATE SET numeric_value=excluded.numeric_value, unit=excluded.unit""",
            (community_id, period_id, key, float(amount)),
        )
    connection.commit()
    return saved


def write_fixed_costs(workbook, values: dict[str, Decimal], set_cell) -> list[str]:
    """Escribe los importes en su celda si la fila del libro es la esperada.

    Devuelve las celdas escritas. Una plantilla propia con otra disposición no
    se toca: sólo se escribe donde la etiqueta de la columna A coincide.
    """
    if SHEET not in workbook.sheetnames or not values:
        return []
    sheet = workbook[SHEET]
    written = []
    for item in FIXED_COSTS:
        if item.key not in values:
            continue
        row = sheet[item.cell].row
        if not _normalise(sheet[f"A{row}"].value).startswith(item.sheet_label):
            continue
        set_cell(sheet[item.cell], float(values[item.key]))
        # Algunos maestros viejos tenían el anual escrito a mano en C. La
        # confirmación mensual también sustituye ese total por su fórmula.
        annual = f"C{row}"
        set_cell(sheet[annual], f"={item.cell}*12")
        written.extend((item.cell, annual))
    return written


__all__ = [
    "FIXED_COSTS",
    "FIXED_COST_CELLS",
    "load_fixed_costs",
    "save_fixed_costs",
    "write_fixed_costs",
    "unconfirmed_template_costs",
    "inherited_cost_message",
    "matching_fixed_cost_cells",
]
