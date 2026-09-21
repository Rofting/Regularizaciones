"""Generación reproducible y atómica del Excel oficial de un expediente."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import sqlite3
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Callable, Mapping
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook

import document_review
import cuotas_servicio
from excel_profiles import (
    ExcelProfile,
    calculate_profile_sha256,
    configured_profile_paths,
    load_profile,
)
from excel_validation import validate_workbook, workbook_fingerprint
from office_recalculation import LibreOfficeRecalculator, WorkbookRecalculator


class ExportBlockedError(ValueError):
    """El expediente no reúne las condiciones para una salida oficial."""


@dataclass(frozen=True)
class ExportResult:
    id_export_run: int
    output_path: Path
    backup_path: Path | None
    stage: str
    input_hash: str


ProgressCallback = Callable[[str, dict[str, Any]], None]


def _emit(progress: ProgressCallback | None, stage: str, **payload: Any) -> None:
    if progress is not None:
        progress(stage, payload)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        while chunk := handle.read(64 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _cents(value: Any) -> int:
    amount = Decimal(str(value or 0)).quantize(Decimal("0.01"), ROUND_HALF_UP)
    return int(amount * 100)


def _safe_date(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value[:10])
    except ValueError as error:
        raise ExportBlockedError(f"Fecha normalizada no válida: {value!r}") from error


def _profile_for_community(
    connection: sqlite3.Connection,
    project_root: Path,
    community_code: str,
    community_id: int,
) -> ExcelProfile:
    """Resuelve primero el perfil que ya está activo en la base.

    Un JSON nuevo para la misma comunidad no puede cambiar silenciosamente el
    formato de un expediente histórico. Sólo el primer uso, sin registro,
    puede escoger el único perfil compatible disponible.
    """
    active = connection.execute(
        """SELECT profile_key,profile_version,profile_sha256 FROM excel_template_profiles
           WHERE id_comunidad=? AND status='active'
           ORDER BY id_template_profile""",
        (community_id,),
    ).fetchall()
    if len(active) > 1:
        raise ExportBlockedError(
            f"Hay varios perfiles activos para la comunidad {community_code}"
        )
    if active:
        row = active[0]
        try:
            profile = load_profile(row["profile_key"], project_root)
        except (LookupError, ValueError) as error:
            raise ExportBlockedError(
                f"No se puede cargar el perfil activo {row['profile_key']}: {error}"
            ) from error
        if profile.version != row["profile_version"]:
            raise ExportBlockedError("La versión del perfil activo no coincide con su registro")
        if calculate_profile_sha256(profile, project_root) != row["profile_sha256"]:
            raise ExportBlockedError(
                "La huella del perfil activo no coincide con su registro; reimporte y revalide"
            )
        if profile.community_code != str(community_code):
            raise ExportBlockedError("El perfil activo no corresponde a la comunidad")
        return profile

    matches: list[ExcelProfile] = []
    for path in configured_profile_paths(project_root):
        try:
            profile = load_profile(path.stem, project_root)
        except (LookupError, ValueError):
            continue
        if profile.community_code == str(community_code):
            matches.append(profile)
    if not matches:
        raise ExportBlockedError(
            f"No hay un perfil Excel configurado para la comunidad {community_code}"
        )
    if len(matches) != 1:
        keys = ", ".join(profile.key for profile in matches)
        raise ExportBlockedError(
            f"Hay varios perfiles activos posibles para la comunidad {community_code}: {keys}"
        )
    profile = matches[0]
    if not profile.workbook_layout:
        raise ExportBlockedError(f"El perfil {profile.key} no declara rangos de escritura")
    return profile


def _input_case_context(connection: sqlite3.Connection, id_case: int) -> sqlite3.Row:
    """Contexto canónico de un expediente para calcular su huella de entrada."""
    row = connection.execute(
        """SELECT c.id_case,c.id_comunidad,c.nombre,c.fecha_inicio,c.fecha_fin,
                  c.estado,c.id_periodo,co.codigo,co.nombre AS comunidad_nombre
           FROM regularization_cases c
           JOIN comunidades co ON co.id_comunidad=c.id_comunidad
           WHERE c.id_case=?""",
        (id_case,),
    ).fetchone()
    if row is None:
        raise ExportBlockedError("El expediente no existe")
    return row


def _case_context(connection: sqlite3.Connection, id_case: int) -> sqlite3.Row:
    row = _input_case_context(connection, id_case)
    if row["estado"] != "ready_for_calculation":
        raise ExportBlockedError(
            "El expediente debe estar listo para cálculo antes de generar el Excel"
        )
    if row["id_periodo"] is None:
        raise ExportBlockedError("El expediente no tiene un período ligado")
    period = connection.execute(
        """SELECT id_comunidad,fecha_inicio,fecha_fin FROM periodos WHERE id_periodo=?""",
        (row["id_periodo"],),
    ).fetchone()
    if period is None or int(period["id_comunidad"]) != int(row["id_comunidad"]):
        raise ExportBlockedError("El período ligado no pertenece a la comunidad")
    if period["fecha_inicio"] != row["fecha_inicio"] or period["fecha_fin"] != row["fecha_fin"]:
        raise ExportBlockedError("Las fechas del período no coinciden con el expediente")
    issue_count = connection.execute(
        "SELECT COUNT(*) FROM review_issues WHERE id_case=? AND status='open'",
        (id_case,),
    ).fetchone()[0]
    if issue_count:
        raise ExportBlockedError(
            f"Hay {issue_count} incidencia(s) abierta(s); revísalas antes de exportar"
        )
    invalid_documents = connection.execute(
        """SELECT COUNT(*) FROM source_documents
           WHERE id_case=? AND status NOT IN ('validated','not_applicable')""",
        (id_case,),
    ).fetchone()[0]
    if invalid_documents:
        raise ExportBlockedError("Hay fuentes del expediente pendientes de validación")
    if document_review.case_has_unapplied_sources(connection, id_case):
        raise ExportBlockedError("Hay fuentes pendientes de confirmar y aplicar antes de exportar")
    return row


def calculate_case_input_hash(
    connection: sqlite3.Connection,
    *,
    id_case: int,
    project_root: Path,
    profile: ExcelProfile | None = None,
) -> str:
    """Devuelve la huella canónica de las entradas actuales de un expediente.

    Esta es la única puerta pública para comparar las entradas usadas por el
    Excel oficial con las que pretende usar un reparto posterior. Si el
    llamante conoce el perfil que empleó una salida histórica debe pasarlo para
    impedir que un perfil activo más nuevo cambie la huella silenciosamente.
    """
    case = _input_case_context(connection, id_case)
    root = Path(project_root).resolve()
    selected_profile = profile or _profile_for_community(
        connection, root, case["codigo"], int(case["id_comunidad"])
    )
    if selected_profile.community_code != str(case["codigo"]):
        raise ExportBlockedError("El perfil no corresponde a la comunidad del expediente")
    profile_hash = calculate_profile_sha256(selected_profile, root)
    return _input_hash(connection, case, selected_profile, profile_hash)


def _required_parameter_keys(profile: ExcelProfile) -> set[str]:
    result: set[str] = set()
    for concept in profile.concepts:
        if not concept.required:
            continue
        for source in (concept.actual_source, concept.billed_source):
            if source and source.startswith("period_parameters."):
                result.add(source.split(".", 1)[1])
    return result


def _validate_normalized_inputs(
    connection: sqlite3.Connection,
    case: sqlite3.Row,
    profile: ExcelProfile,
) -> None:
    period_id = int(case["id_periodo"])
    for module in ("GAS", "ELECTRICIDAD", "AGUA"):
        if module not in profile.active_modules:
            continue
        count = connection.execute(
            """SELECT COUNT(*) FROM facturas
               WHERE id_comunidad=? AND id_periodo=? AND tipo_suministro=?""",
            (case["id_comunidad"], period_id, module),
        ).fetchone()[0]
        if not count:
            raise ExportBlockedError(f"Faltan facturas normalizadas de {module}")
        incomplete = connection.execute(
            """SELECT COUNT(*) FROM facturas f
               WHERE f.id_comunidad=? AND f.id_periodo=? AND f.tipo_suministro=?
                 AND (f.fecha_factura IS NULL OR f.fecha_inicio IS NULL OR f.fecha_fin IS NULL
                      OR f.importe_total IS NULL
                      OR NOT EXISTS (
                          SELECT 1 FROM invoice_components c
                          WHERE c.id_factura=f.id_factura AND c.component_key='total'
                      ))""",
            (case["id_comunidad"], period_id, module),
        ).fetchone()[0]
        if incomplete:
            raise ExportBlockedError(f"Hay facturas de {module} con datos obligatorios incompletos")
        invoices = connection.execute(
            """SELECT f.id_factura,f.termino_fijo,f.termino_variable,f.importe_total,
                      MAX(CASE WHEN c.component_key='fixed' THEN c.amount END) AS fixed_component,
                      MAX(CASE WHEN c.component_key='variable' THEN c.amount END) AS variable_component,
                      MAX(CASE WHEN c.component_key='total' THEN c.amount END) AS total_component
               FROM facturas f JOIN invoice_components c ON c.id_factura=f.id_factura
               WHERE f.id_comunidad=? AND f.id_periodo=? AND f.tipo_suministro=?
               GROUP BY f.id_factura""",
            (case["id_comunidad"], period_id, module),
        ).fetchall()
        for invoice in invoices:
            fixed = _cents(invoice["termino_fijo"])
            variable = _cents(invoice["termino_variable"])
            total = _cents(invoice["importe_total"])
            if (
                fixed != _cents(invoice["fixed_component"])
                or variable != _cents(invoice["variable_component"])
                or total != _cents(invoice["total_component"])
            ):
                raise ExportBlockedError(
                    f"Los componentes de la factura {invoice['id_factura']} de {module} "
                    "no coinciden con los guardados en la factura"
                )
            # El desglose de término fijo y variable se extrae del PDF y no es
            # fiable: unas facturas no lo traen, en otras deja fuera impuestos
            # o alquiler de equipos, y en las de varias páginas mezcla cifras.
            # El importe total sí es de fiar y es el que usa el reparto, así
            # que un desglose que no cuadre no detiene el expediente: se
            # escribe sólo el total (ver _write_derived_formula).

    parameters = {
        row[0]
        for row in connection.execute(
            """SELECT parameter_key FROM period_parameters
               WHERE id_comunidad=? AND id_periodo=? AND numeric_value IS NOT NULL""",
            (case["id_comunidad"], period_id),
        )
    }
    # Los importes de ACS y calefacción no son entradas: el libro los calcula
    # en su hoja de análisis a partir de las facturas (coste real) y de la hoja
    # de cobros (importe cobrado). Exigirlos aquí pedía a mano justo lo que el
    # modelo deduce, y escribirlos habría machacado sus fórmulas. Sólo se
    # reclaman cuando la comunidad no tiene cuotas anotadas de las que partir.
    calculados = {
        clave
        for servicio, prefijo in (("ACS", "acs"), ("CALEFACCION", "heating"))
        if servicio in profile.active_modules
        and cuotas_servicio.resumen(
            connection, community_id=case["id_comunidad"],
            period_id=period_id, servicio=servicio,
        ).apuntes
        for clave in (
            f"{prefijo}_fixed_actual", f"{prefijo}_variable_actual",
            f"{prefijo}_fixed_billed", f"{prefijo}_variable_billed",
        )
    }
    missing_parameters = sorted(
        _required_parameter_keys(profile).difference(parameters).difference(calculados)
    )
    if missing_parameters:
        servicios = sorted({
            "ACS" if clave.startswith("acs_") else "calefacción"
            for clave in missing_parameters
            if clave.startswith(("acs_", "heating_"))
        })
        if servicios:
            raise ExportBlockedError(
                "Faltan las cuotas cobradas de " + " y ".join(servicios)
                + ". Anótalas en «Cuotas cobradas» y vuelve a generar: de ellas "
                "sale el importe cobrado del análisis."
            )
        raise ExportBlockedError(
            "Faltan parámetros normalizados: " + ", ".join(missing_parameters)
        )
    for reading_type in _profile_reading_types(profile):
        # Garajes y locales no tienen contador de ACS ni de calefacción: pedirles
        # lecturas de apertura y cierre hacía fallar a toda la comunidad.
        owners = connection.execute(
            """SELECT id_propietario FROM propietarios
               WHERE id_comunidad=? AND activo=1 AND tipo_unidad='vivienda'""",
            (case["id_comunidad"],),
        ).fetchall()
        if not owners:
            raise ExportBlockedError(
                f"No hay viviendas activas para las lecturas {reading_type}"
            )
        for owner in owners:
            readings = connection.execute(
                """SELECT fecha_lectura,valor_acumulado,estado FROM period_readings
                   WHERE id_propietario=? AND id_periodo=? AND tipo=?
                   ORDER BY fecha_lectura""",
                (owner[0], period_id, reading_type),
            ).fetchall()
            dates = {reading["fecha_lectura"] for reading in readings}
            # Mismas fechas que usará la hoja: la última lectura antes de que
            # arranque el período y la última que cabe dentro, no el día exacto
            # de inicio y fin del ejercicio.
            fecha_inicial, fecha_final = _boundary_reading_dates(
                connection, case, reading_type,
            )
            if fecha_inicial not in dates or fecha_final not in dates:
                raise ExportBlockedError(
                    f"Faltan lecturas {reading_type} de {fecha_inicial} o {fecha_final}"
                )
            if any(reading["estado"] in ("sin_lectura", "contador_averiado") for reading in readings):
                raise ExportBlockedError(f"Hay una lectura {reading_type} sin validar")


def _template_registration(
    connection: sqlite3.Connection,
    case: sqlite3.Row,
    profile: ExcelProfile,
    project_root: Path,
) -> tuple[int, Path, str]:
    root = project_root.resolve()
    template = (root / profile.template_relative_path).resolve()
    try:
        template.relative_to(root)
    except ValueError:
        raise ExportBlockedError("La plantilla queda fuera del proyecto") from None
    if not template.is_file():
        raise ExportBlockedError(f"No existe la plantilla registrada: {template}")
    template_hash = _sha256(template)
    active = connection.execute(
        """SELECT id_template_profile,profile_key,profile_version,
                  template_relative_path,template_sha256,profile_sha256
           FROM excel_template_profiles
           WHERE id_comunidad=? AND status='active'
           ORDER BY id_template_profile DESC""",
        (case["id_comunidad"],),
    ).fetchall()
    if active:
        matching = [
            row for row in active
            if row["profile_key"] == profile.key and row["profile_version"] == profile.version
        ]
        if len(active) != 1 or not matching:
            raise ExportBlockedError("El perfil activo registrado no coincide con el expediente")
        row = matching[0]
        if (
            row["template_relative_path"] != profile.template_relative_path
            or row["template_sha256"] != template_hash
        ):
            raise ExportBlockedError(
                "La plantilla cambió respecto del hash registrado; registra una nueva versión"
            )
        if row["profile_sha256"] != calculate_profile_sha256(profile, project_root):
            raise ExportBlockedError(
                "La huella del perfil cambió respecto del registro; reimporte y revalide"
            )
        return int(row["id_template_profile"]), template, template_hash
    profile_hash = calculate_profile_sha256(profile, project_root)
    cursor = connection.execute(
        """INSERT INTO excel_template_profiles
           (id_comunidad,profile_key,profile_version,template_relative_path,
            template_sha256,profile_sha256,status)
           VALUES (?,?,?,?,?,?,'active')""",
        (
            case["id_comunidad"], profile.key, profile.version,
            profile.template_relative_path, template_hash, profile_hash,
        ),
    )
    connection.commit()
    return int(cursor.lastrowid), template, template_hash


def _query_dicts(connection: sqlite3.Connection, query: str, parameters: tuple) -> list[dict]:
    cursor = connection.execute(query, parameters)
    names = [description[0] for description in cursor.description]
    return [dict(zip(names, row)) for row in cursor.fetchall()]


def _profile_reading_types(profile: ExcelProfile) -> tuple[str, ...]:
    """Tipos de contador que forman parte de las entradas del perfil.

    Los libros de ACS no deben invalidarse por una lectura de calefacción que
    no utilizan, mientras que un perfil de calefacción sí debe fijar esa
    lectura en su huella de entrada. Se reconocen tanto módulos declarados
    como conceptos que usan la convención de nombre de cada servicio.
    """
    if profile.onboarding_configuration:
        return tuple(
            binding["module"]
            for binding in profile.onboarding_configuration["reading_bindings"]
        )
    types: set[str] = set()
    concept_keys = {concept.key for concept in profile.concepts}
    if "ACS" in profile.active_modules or any(key.startswith("acs_") for key in concept_keys):
        types.add("ACS")
    if "CALEFACCION" in profile.active_modules or any(
        key.startswith("heating_") for key in concept_keys
    ):
        types.add("CALEFACCION")
    return tuple(sorted(types))


def _input_hash(
    connection: sqlite3.Connection,
    case: sqlite3.Row,
    profile: ExcelProfile,
    profile_hash: str | None = None,
) -> str:
    profile_hash = profile_hash or profile.source_sha256
    if not profile_hash:
        raise ExportBlockedError("No se puede determinar la huella del perfil")
    period_id = int(case["id_periodo"])
    reading_types = _profile_reading_types(profile)
    reading_rows = []
    if reading_types:
        placeholders = ",".join("?" for _ in reading_types)
        reading_rows = _query_dicts(
            connection,
            f"""SELECT p.codigo_vivienda,l.tipo,l.fecha_lectura,l.valor_acumulado,l.estado,
                       l.approved_by,l.approved_at
                FROM period_readings l JOIN propietarios p ON p.id_propietario=l.id_propietario
                WHERE p.id_comunidad=? AND l.id_periodo=? AND l.tipo IN ({placeholders})
                ORDER BY p.codigo_vivienda,l.tipo,l.fecha_lectura""",
            (case["id_comunidad"], period_id, *reading_types),
        )
    payload = {
        "case": {key: case[key] for key in (
            "id_case", "id_comunidad", "codigo", "comunidad_nombre", "nombre",
            "fecha_inicio", "fecha_fin", "id_periodo"
        )},
        "profile": [profile.key, profile.version, profile_hash],
        "invoices": _query_dicts(
            connection,
            """SELECT id_factura,tipo_suministro,proveedor,num_factura,fecha_factura,
                      fecha_inicio,fecha_fin,dias_facturados,consumo_total,unidad_consumo,
                      termino_fijo,termino_variable,importe_total
               FROM facturas WHERE id_comunidad=? AND id_periodo=? ORDER BY id_factura""",
            (case["id_comunidad"], period_id),
        ),
        "components": _query_dicts(
            connection,
            """SELECT c.id_factura,c.component_key,c.amount,c.unit
               FROM invoice_components c JOIN facturas f ON f.id_factura=c.id_factura
               WHERE f.id_comunidad=? AND f.id_periodo=?
               ORDER BY c.id_factura,c.component_key""",
            (case["id_comunidad"], period_id),
        ),
        "parameters": _query_dicts(
            connection,
            """SELECT parameter_key,numeric_value,text_value,unit
               FROM period_parameters WHERE id_comunidad=? AND id_periodo=?
               ORDER BY parameter_key""",
            (case["id_comunidad"], period_id),
        ),
        "readings": reading_rows,
        "eligible_expenses": _query_dicts(
            connection,
            """SELECT id_gasto,fecha,descripcion,importe_total,activo
               FROM gastos_extra
               WHERE id_comunidad=? AND activo=1 AND fecha BETWEEN ? AND ?
               ORDER BY id_gasto""",
            (case["id_comunidad"], case["fecha_inicio"], case["fecha_fin"]),
        ),
        "owners": _query_dicts(
            connection,
            """SELECT id_propietario,codigo_vivienda,nombre_propietario,coeficiente,activo
               FROM propietarios WHERE id_comunidad=? ORDER BY id_propietario""",
            (case["id_comunidad"],),
        ),
    }
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def _input_columns(table: Mapping[str, Any]) -> Mapping[str, str]:
    """Columnas cuyo valor pertenece a la base de datos.

    Los perfiles nuevos separan explícitamente las columnas derivadas. La
    compatibilidad con el formato inicial se mantiene para perfiles ya
    registrados, pero nunca se toma una celda con fórmula como fila inválida.
    """
    return table.get("input_columns", table.get("columns", {}))


def _derived_columns(table: Mapping[str, Any]) -> Mapping[str, str]:
    return table.get("derived_columns", {})


def _set_cell_value(cell, value: Any) -> None:
    """Cambia contenido sin alterar el estilo efectivo de la plantilla."""
    style = copy(cell._style)
    cell.value = value
    cell._style = style


def _clear_table(sheet, table: Mapping[str, Any]) -> None:
    season = table.get("season_columns") or {}
    seasonal = [
        str(season[key]) for key in ("winter", "summer") if season.get(key)
    ]
    for row in range(int(table["start_row"]), int(table["end_row"]) + 1):
        for column in (*_input_columns(table).values(),
                       *_derived_columns(table).values(), *seasonal):
            cell = sheet[f"{column}{row}"]
            cell.value = None


def _available_table_rows(sheet, table: Mapping[str, Any]) -> list[int]:
    """Todas las filas declaradas son entradas; las fórmulas no las invalidan."""
    return list(range(int(table["start_row"]), int(table["end_row"]) + 1))


def _write_derived_formula(
    sheet, row: int, columns: Mapping[str, str], total_value: float | None = None,
) -> None:
    """Escribe únicamente fórmulas que se derivan de insumos normalizados.

    Con desglose, el total es la suma de sus términos y se escribe como
    fórmula viva. Sin desglose —facturas que sólo traen el importe— se escribe
    el total tal cual: la fórmula daría cero y borraría el importe real.
    """
    total = columns.get("total")
    fixed = columns.get("fixed")
    variable = columns.get("variable")
    if not total:
        return
    celda = sheet[f"{total}{row}"]
    if isinstance(celda.value, str) and celda.value.startswith("="):
        return
    if fixed and variable:
        terminos = [sheet[f"{columna}{row}"].value for columna in (fixed, variable)]
        numericos = [valor for valor in terminos if isinstance(valor, (int, float))]
        suma = sum(numericos)
        if numericos and (
            total_value is None or abs(suma - float(total_value)) <= 0.01
        ):
            _set_cell_value(celda, f"={fixed}{row}+{variable}{row}")
            return
        # Si el desglose no cuadra con el total se conserva —es lo que dice la
        # factura y con ello concilia la base—, pero el total deja de ser su
        # suma y se escribe el importe real cobrado.
    if total_value is not None:
        _set_cell_value(celda, total_value)


_TEMPORADA_POR_DEFECTO = ("12-01", "05-31")


def _season_boundary(value: Any, fallback: str) -> tuple[int, int]:
    """Lee un 'MM-DD' de configuración y lo deja como (mes, día)."""
    texto = str(value or fallback).strip()
    try:
        month, _, day = texto.partition("-")
        boundary = (int(month), int(day))
    except ValueError:
        raise ExportBlockedError(f"Límite de temporada no válido: {texto!r}") from None
    if not (1 <= boundary[0] <= 12 and 1 <= boundary[1] <= 31):
        raise ExportBlockedError(f"Límite de temporada fuera de rango: {texto!r}")
    return boundary


def _season_columns(
    connection: sqlite3.Connection,
    case: sqlite3.Row,
    module: str,
    table: Mapping[str, Any],
) -> tuple[str, str, tuple[int, int], tuple[int, int]] | None:
    """Columnas y fechas de la temporada de calefacción de un suministro.

    Cada suministro tiene la suya: en el modelo del despacho el gas cuenta como
    invierno lo que cierra entre el 1 de diciembre y el 31 de mayo, y la
    electricidad entre el 1 de noviembre y el 30 de abril. El período puede
    redefinirlas sin tocar el perfil.
    """
    # El perfil congela sus mapas (MappingProxyType), así que no basta con
    # comprobar dict: sin esto el reparto verano/invierno se saltaba entero.
    columnas = table.get("season_columns")
    if not isinstance(columnas, Mapping):
        return None
    invierno, verano = columnas.get("winter"), columnas.get("summer")
    if not invierno or not verano:
        return None
    parametros = {
        row["parameter_key"]: row["text_value"]
        for row in connection.execute(
            """SELECT parameter_key,text_value FROM period_parameters
               WHERE id_comunidad=? AND id_periodo=?""",
            (case["id_comunidad"], case["id_periodo"]),
        )
    }
    inicio = parametros.get(f"winter_season_start:{module}") or columnas.get("winter_start")
    fin = parametros.get(f"winter_season_end:{module}") or columnas.get("winter_end")
    return (
        str(invierno), str(verano),
        _season_boundary(inicio, _TEMPORADA_POR_DEFECTO[0]),
        _season_boundary(fin, _TEMPORADA_POR_DEFECTO[1]),
    )


def _is_winter(day: date | None, start: tuple[int, int], end: tuple[int, int]) -> bool:
    """La temporada cruza el cambio de año, así que se comprueba en dos tramos."""
    if day is None:
        return False
    moment = (day.month, day.day)
    if start <= end:
        return start <= moment <= end
    return moment >= start or moment <= end


def _write_invoice_table(
    connection: sqlite3.Connection,
    workbook,
    case: sqlite3.Row,
    module: str,
    table: Mapping[str, Any],
) -> None:
    sheet = workbook[table["sheet"]]
    available_rows = _available_table_rows(sheet, table)
    _clear_table(sheet, table)
    invoices = connection.execute(
        """SELECT fecha_factura,fecha_inicio,fecha_fin,dias_facturados,
                  consumo_total,termino_fijo,termino_variable,importe_total,proveedor
           FROM facturas
           WHERE id_comunidad=? AND id_periodo=? AND tipo_suministro=?
           ORDER BY fecha_factura,id_factura""",
        (case["id_comunidad"], case["id_periodo"], module),
    ).fetchall()
    if len(invoices) > len(available_rows):
        raise ExportBlockedError(f"La plantilla no tiene filas suficientes para {module}")
    columns = _input_columns(table)
    derived = _derived_columns(table)
    season = _season_columns(connection, case, module, table)
    for row_number, invoice in zip(available_rows, invoices):
        values = {
            "invoice_date": _safe_date(invoice["fecha_factura"]),
            "days": int(invoice["dias_facturados"] or 0),
            "start_date": _safe_date(invoice["fecha_inicio"]),
            "end_date": _safe_date(invoice["fecha_fin"]),
            "consumption": float(invoice["consumo_total"] or 0),
            "fixed": float(invoice["termino_fijo"] or 0),
            "variable": float(invoice["termino_variable"] or 0),
            "total": float(invoice["importe_total"]),
            "provider": invoice["proveedor"] or "",
        }
        for key, column in columns.items():
            if key in values:
                _set_cell_value(sheet[f"{column}{row_number}"], values[key])
        _write_derived_formula(
            sheet, row_number, {**columns, **derived}, total_value=values["total"],
        )
        if season is not None and columns.get("consumption"):
            winter_column, summer_column, start, end = season
            # El consumo entero va a una de las dos columnas, nunca a las dos:
            # de ahí salen los % de ACS y calefacción de la fila de sumas.
            destino = winter_column if _is_winter(values["end_date"], start, end) else summer_column
            otra = summer_column if destino == winter_column else winter_column
            _set_cell_value(
                sheet[f"{destino}{row_number}"],
                f"={columns['consumption']}{row_number}",
            )
            _set_cell_value(sheet[f"{otra}{row_number}"], None)


def _write_other_expenses(connection, workbook, case, table) -> None:
    sheet = workbook[table["sheet"]]
    available_rows = _available_table_rows(sheet, table)
    _clear_table(sheet, table)
    expenses = connection.execute(
        """SELECT fecha,descripcion,importe_total FROM gastos_extra
           WHERE id_comunidad=? AND activo=1 AND fecha BETWEEN ? AND ?
           ORDER BY fecha,id_gasto""",
        (case["id_comunidad"], case["fecha_inicio"], case["fecha_fin"]),
    ).fetchall()
    if not expenses:
        parameter = connection.execute(
            """SELECT numeric_value FROM period_parameters
               WHERE id_comunidad=? AND id_periodo=?
                 AND parameter_key='extraordinary_expense_actual'""",
            (case["id_comunidad"], case["id_periodo"]),
        ).fetchone()
        if parameter is not None:
            expenses = [(case["fecha_fin"], "Total de otros gastos", parameter[0])]
    if len(expenses) > len(available_rows):
        raise ExportBlockedError("La plantilla no tiene filas suficientes para otros gastos")
    columns = _input_columns(table)
    for row_number, expense in zip(available_rows, expenses):
        _set_cell_value(sheet[f"{columns['date']}{row_number}"], _safe_date(expense[0]))
        _set_cell_value(sheet[f"{columns['description']}{row_number}"], expense[1])
        _set_cell_value(sheet[f"{columns['amount']}{row_number}"], float(expense[2]))


def _boundary_reading_dates(
    connection: sqlite3.Connection, case: sqlite3.Row, reading_type: str,
) -> tuple[str, str]:
    """Fechas de lectura que abren y cierran el período.

    Los contadores no se leen el día que empieza el ejercicio: la lectura
    inicial es la última tomada antes de que arranque —normalmente el cierre
    del período anterior— y la final, la última que cabe dentro. Exigir que
    coincidieran exactamente con las fechas del expediente hacía fallar a
    cualquier comunidad por un día de diferencia.
    """
    fila = connection.execute(
        """SELECT
             (SELECT MAX(fecha_lectura) FROM period_readings r
               JOIN propietarios p ON p.id_propietario=r.id_propietario
               WHERE p.id_comunidad=? AND p.tipo_unidad='vivienda' AND r.tipo=?
                 AND r.fecha_lectura<=?) AS inicial,
             (SELECT MAX(fecha_lectura) FROM period_readings r
               JOIN propietarios p ON p.id_propietario=r.id_propietario
               WHERE p.id_comunidad=? AND p.tipo_unidad='vivienda' AND r.tipo=?
                 AND r.fecha_lectura<=?) AS final""",
        (
            case["id_comunidad"], reading_type, case["fecha_inicio"],
            case["id_comunidad"], reading_type, case["fecha_fin"],
        ),
    ).fetchone()
    inicial, final = fila["inicial"], fila["final"]
    if inicial is None or final is None or inicial >= final:
        raise ExportBlockedError(
            f"Faltan lecturas {reading_type} que abran y cierren el período"
        )
    return inicial, final



def _write_meter_readings(
    connection,
    workbook,
    case,
    module,
    table,
    *,
    materialize_consumption: bool = False,
) -> None:
    sheet = workbook[table["sheet"]]
    available_rows = _available_table_rows(sheet, table)
    _clear_table(sheet, table)
    readings = connection.execute(
        """SELECT p.codigo_vivienda,l.fecha_lectura,l.valor_acumulado
           FROM period_readings l JOIN propietarios p ON p.id_propietario=l.id_propietario
           WHERE p.id_comunidad=? AND p.activo=1 AND p.tipo_unidad='vivienda'
             AND l.id_periodo=? AND l.tipo=?
           ORDER BY p.codigo_vivienda,l.fecha_lectura""",
        (case["id_comunidad"], case["id_periodo"], module),
    ).fetchall()
    by_owner: dict[str, dict[str, float]] = {}
    for reading in readings:
        by_owner.setdefault(reading["codigo_vivienda"], {})[reading["fecha_lectura"]] = float(
            reading["valor_acumulado"]
        )
    fecha_inicial, fecha_final = _boundary_reading_dates(connection, case, module)
    sin_lectura = sorted(
        vivienda for vivienda, values in by_owner.items()
        if fecha_inicial not in values or fecha_final not in values
    )
    if sin_lectura:
        muestra = ", ".join(sin_lectura[:5])
        resto = f" y {len(sin_lectura) - 5} más" if len(sin_lectura) > 5 else ""
        raise ExportBlockedError(
            f"Faltan lecturas {module} de {fecha_inicial} o {fecha_final} en: {muestra}{resto}"
        )
    if len(available_rows) < 2:
        raise ExportBlockedError(f"La plantilla no tiene filas suficientes para {module}")
    columns = _input_columns(table)
    derived = _derived_columns(table)
    parameters = {
        row["parameter_key"]: row["numeric_value"]
        for row in connection.execute(
            """SELECT parameter_key,numeric_value FROM period_parameters
               WHERE id_comunidad=? AND id_periodo=?""",
            (case["id_comunidad"], case["id_periodo"]),
        )
    }
    prefix = "acs" if module == "ACS" else "heating"
    tiene_cuotas = cuotas_servicio.resumen(
        connection, community_id=case["id_comunidad"],
        period_id=case["id_periodo"], servicio=module,
    ).apuntes > 0
    required = (f"{prefix}_variable_actual", f"{prefix}_fixed_actual")
    missing = [key for key in required if parameters.get(key) is None]
    if missing and not tiene_cuotas:
        raise ExportBlockedError(
            f"Faltan las cuotas cobradas de {module}, o los parámetros "
            + ", ".join(missing)
        )

    initial = sum(values[fecha_inicial] for values in by_owner.values())
    final = sum(values[fecha_final] for values in by_owner.values())

    # La hoja de cobros lleva dos clases de apunte, como en el libro del
    # despacho: una fila por cada cuota fija mensual y una fila por cada
    # liquidación de consumo. De su suma sale el 'importe cobrado' del
    # análisis. Si la comunidad aún no tiene cuotas anotadas se conserva el
    # resumen de dos filas de siempre, para no dejar la hoja vacía.
    apuntes = cuotas_servicio.listar(
        connection, community_id=case["id_comunidad"],
        period_id=case["id_periodo"], servicio=module,
    )
    filas: list[dict] = []
    if apuntes:
        for apunte in apuntes:
            if apunte["concepto"] == "fija":
                filas.append({
                    "charge_date": _safe_date(apunte["fecha"]),
                    "fixed_fee": float(apunte["importe"]),
                })
            else:
                fila = {
                    "charge_date": _safe_date(apunte["fecha"]),
                    "variable_fee": float(apunte["importe"]),
                }
                if apunte["consumo"] is not None:
                    fila["consumption"] = float(apunte["consumo"])
                if apunte["fecha_inicio"]:
                    fila["initial_date"] = _safe_date(apunte["fecha_inicio"])
                if apunte["fecha_fin"]:
                    fila["final_date"] = _safe_date(apunte["fecha_fin"])
                filas.append(fila)
        # El consumo leído del período se apunta en la última liquidación, que
        # es la que cierra las lecturas de los contadores. Si la comunidad sólo
        # gira cuota fija y no tiene liquidaciones, se añade igualmente la fila
        # de cierre: sin ella el consumo del período no llegaría a la hoja.
        cierre = next((fila for fila in reversed(filas) if "variable_fee" in fila), None)
        if cierre is None:
            cierre = {
                "charge_date": _safe_date(fecha_final),
                "variable_fee": float(parameters.get(f"{prefix}_variable_actual") or 0),
            }
            filas.append(cierre)
        cierre.setdefault("final_date", _safe_date(fecha_final))
        cierre.setdefault("initial_date", _safe_date(fecha_inicial))
        cierre["final"], cierre["initial"] = final, initial
    else:
        filas = [
            {
                "charge_date": _safe_date(fecha_final),
                "final_date": _safe_date(fecha_final), "final": final,
                "initial_date": _safe_date(fecha_inicial), "initial": initial,
                "variable_fee": float(parameters[f"{prefix}_variable_actual"]),
            },
            {
                "charge_date": _safe_date(case["fecha_fin"]),
                "fixed_fee": float(parameters[f"{prefix}_fixed_actual"]),
            },
        ]
    if len(filas) > len(available_rows):
        raise ExportBlockedError(
            f"La plantilla no tiene filas suficientes para las cuotas de {module}"
        )
    escritas = list(zip(available_rows, filas))
    for row, values in escritas:
        for key, column in columns.items():
            if key in values:
                _set_cell_value(sheet[f"{column}{row}"], values[key])

    consumption = derived.get("consumption")
    total = derived.get("total")
    variable_unit = derived.get("variable_unit")
    fixed_unit = derived.get("fixed_unit")
    # Cada apunte lleva sus propias fórmulas, según sea cuota fija mensual o
    # liquidación por consumo, igual que en el libro del despacho.
    for row, values in escritas:
        if total and columns.get("variable_fee") and columns.get("fixed_fee"):
            _set_cell_value(
                sheet[f"{total}{row}"],
                f"=SUM({columns['variable_fee']}{row}:{columns['fixed_fee']}{row})",
            )
        es_variable = "variable_fee" in values
        if es_variable and consumption:
            if "final" in values and "initial" in values:
                _set_cell_value(
                    sheet[f"{consumption}{row}"],
                    final - initial
                    if materialize_consumption
                    else f"={columns['final']}{row}-{columns['initial']}{row}",
                )
            elif values.get("consumption") is not None:
                # El consumo liquidado es una columna derivada de la hoja, no
                # una entrada, así que no lo escribe el bucle de columnas.
                _set_cell_value(sheet[f"{consumption}{row}"], values["consumption"])
        if es_variable and variable_unit and consumption and sheet[f"{consumption}{row}"].value:
            _set_cell_value(
                sheet[f"{variable_unit}{row}"],
                f"={columns['variable_fee']}{row}/{consumption}{row}",
            )
        if "fixed_fee" in values and fixed_unit:
            _set_cell_value(
                sheet[f"{fixed_unit}{row}"],
                f"={columns['fixed_fee']}{row}/'DATOS'!$D$4",
            )


def _write_workbook(
    connection: sqlite3.Connection,
    path: Path,
    case: sqlite3.Row,
    profile: ExcelProfile,
    progress: ProgressCallback | None,
) -> None:
    workbook = load_workbook(path)
    try:
        metadata = profile.workbook_layout["metadata_cells"]
        metadata_values = {
            "community_name": case["comunidad_nombre"],
            "period_label": (
                f"PERIODO: EJERCICIO "
                f"{_safe_date(case['fecha_inicio']):%d/%m/%Y} - "
                f"{_safe_date(case['fecha_fin']):%d/%m/%Y}"
            ),
            # Garajes y locales no entran en la regularización: el modelo
            # divide cuotas fijas entre viviendas, así que contarlos desviaría
            # todos los €/vivienda de la hoja de análisis.
            "owner_count": connection.execute(
                """SELECT COUNT(*) FROM propietarios
                   WHERE id_comunidad=? AND activo=1 AND tipo_unidad='vivienda'""",
                (case["id_comunidad"],),
            ).fetchone()[0],
        }
        for key, value in metadata_values.items():
            if key in metadata:
                sheet_name, address = metadata[key]
                _set_cell_value(workbook[sheet_name][address], value)

        tables = profile.workbook_layout["tables"]
        for module in profile.active_modules:
            _emit(progress, f"write_{module.lower()}", module=module)
            table = tables.get(module)
            if table is None:
                raise ExportBlockedError(f"El perfil no declara la tabla de {module}")
            if module in ("GAS", "ELECTRICIDAD", "AGUA"):
                _write_invoice_table(connection, workbook, case, module, table)
            elif module == "OTROS_GASTOS":
                _write_other_expenses(connection, workbook, case, table)
            elif module in ("ACS", "CALEFACCION"):
                _write_meter_readings(
                    connection,
                    workbook,
                    case,
                    module,
                    table,
                    materialize_consumption=bool(profile.onboarding_configuration),
                )

        parameters = {
            row["parameter_key"]: row["numeric_value"]
            for row in connection.execute(
                """SELECT parameter_key,numeric_value FROM period_parameters
                   WHERE id_comunidad=? AND id_periodo=?""",
                (case["id_comunidad"], case["id_periodo"]),
            )
        }
        for key, binding in profile.workbook_layout["parameter_cells"].items():
            if key in parameters and parameters[key] is not None:
                sheet_name, address = binding
                celda = workbook[sheet_name][address]
                # El modelo calcula estos importes con sus propias fórmulas.
                # Escribir encima las borraría y congelaría el análisis.
                if isinstance(celda.value, str) and celda.value.startswith("="):
                    continue
                _set_cell_value(celda, float(parameters[key]))
        for check, binding in profile.workbook_layout["total_checks"].items():
            if not check.startswith("parameter:"):
                continue
            key = check.split(":", 1)[1]
            if key in profile.workbook_layout["parameter_cells"]:
                continue
            if key in parameters and parameters[key] is not None:
                sheet_name, address = binding
                if ":" not in address:
                    _set_cell_value(workbook[sheet_name][address], float(parameters[key]))
        workbook.save(path)
    finally:
        workbook.close()


def _restore_missing_ooxml_parts(template: Path, generated: Path) -> None:
    """Restaura el diseño OOXML ajeno a las celdas mutables.

    Esta exportación nunca añade hojas, gráficos ni relaciones: sólo escribe
    entradas ya declaradas. Por ello es seguro restaurar desde la plantilla
    las partes visuales y sus relaciones, aunque openpyxl las haya recreado.
    """
    with ZipFile(template, "r") as source:
        source_parts = {name: source.read(name) for name in source.namelist()}
    with ZipFile(generated, "r") as current:
        generated_parts = {name: current.read(name) for name in current.namelist()}
    design_prefixes = (
        "xl/charts/", "xl/drawings/", "xl/media/", "customXml/",
        "xl/theme/", "xl/printerSettings/",
    )
    # styles.xml tampoco: las celdas apuntan a sus formatos por índice, así
    # que traer la tabla de la plantilla desplazaba todos los formatos y los
    # importes aparecían como fechas. LibreOffice ya conserva el formato.
    design_exact = {"_rels/.rels"}
    # Nunca se restauran las partes que llevan el contenido: las hojas, el
    # libro y, sobre todo, la tabla de cadenas compartidas. Las celdas apuntan
    # a ella por índice, así que traer la de la plantilla sobre un libro ya
    # reescrito descoloca todos los textos y el fichero deja de poder leerse
    # ("shared_strings: list index out of range").
    content_prefixes = ("xl/worksheets/",)
    content_exact = {
        "xl/sharedStrings.xml", "xl/workbook.xml", "xl/calcChain.xml",
        # El mapa de relaciones del libro empareja cada hoja con su XML. La
        # plantilla numera rId1..rId7 y el libro recalculado rId3..rId9, así
        # que traer el de la plantilla descolocaba todas las hojas: 'ANALISIS'
        # apuntaba a una que no existía y aparecía vacía, sin sus fórmulas.
        "xl/_rels/workbook.xml.rels",
    }
    restore = {
        name: data
        for name, data in source_parts.items()
        if not name.startswith(content_prefixes)
        and name not in content_exact
        and (
            name not in generated_parts
            or name.startswith(design_prefixes)
            or name in design_exact
            or name.endswith(".rels")
        )
    }
    if not restore:
        return
    rebuilt = generated.with_name(generated.stem + ".ooxml" + generated.suffix)
    tipos = _merged_content_types(
        source_parts.get("[Content_Types].xml"),
        generated_parts.get("[Content_Types].xml"),
    )
    with ZipFile(rebuilt, "w", ZIP_DEFLATED) as target:
        for name, data in generated_parts.items():
            if name == "[Content_Types].xml" and tipos is not None:
                target.writestr(name, tipos)
            elif name not in restore:
                target.writestr(name, data)
        for name, data in restore.items():
            target.writestr(name, data)
    os.replace(rebuilt, generated)


def _merged_content_types(plantilla: bytes | None, generado: bytes | None) -> bytes | None:
    """Declaraciones de la plantilla sin perder las del libro generado.

    La plantilla guarda sus textos dentro de cada celda y no declara
    'xl/sharedStrings.xml'; el libro que sale del recálculo sí lo usa. Copiar
    tal cual el catálogo de la plantilla dejaba esa tabla sin declarar, con lo
    que las celdas de texto apuntaban a una lista vacía y el fichero ya no se
    podía abrir. Se parte del catálogo del generado y se le añade lo que sólo
    traiga la plantilla.
    """
    if plantilla is None or generado is None:
        return generado or plantilla
    import xml.etree.ElementTree as ET

    espacio = "http://schemas.openxmlformats.org/package/2006/content-types"
    ET.register_namespace("", espacio)
    raiz = ET.fromstring(generado)
    presentes = {
        (hijo.tag, hijo.get("PartName") or hijo.get("Extension"))
        for hijo in raiz
    }
    for hijo in ET.fromstring(plantilla):
        clave = (hijo.tag, hijo.get("PartName") or hijo.get("Extension"))
        if clave not in presentes:
            raiz.append(hijo)
    return ET.tostring(raiz, encoding="UTF-8", xml_declaration=True)


def _expected_totals(connection, case, profile) -> dict[str, int]:
    calculados_por_el_libro = {
        clave
        for servicio, prefijo in (("ACS", "acs"), ("CALEFACCION", "heating"))
        if servicio in profile.active_modules
        and cuotas_servicio.resumen(
            connection, community_id=case["id_comunidad"],
            period_id=case["id_periodo"], servicio=servicio,
        ).apuntes
        for clave in (
            f"{prefijo}_fixed_actual", f"{prefijo}_variable_actual",
            f"{prefijo}_fixed_billed", f"{prefijo}_variable_billed",
        )
    }
    result: dict[str, int] = {}
    required_parameters = _required_parameter_keys(profile)
    for key in profile.workbook_layout["total_checks"]:
        if key.startswith("invoice_total:"):
            module = key.split(":", 1)[1]
            value = connection.execute(
                """SELECT COALESCE(SUM(importe_total),0) FROM facturas
                   WHERE id_comunidad=? AND id_periodo=? AND tipo_suministro=?""",
                (case["id_comunidad"], case["id_periodo"], module),
            ).fetchone()[0]
        elif key.startswith("invoice_component:"):
            _, module, component = key.split(":", 2)
            value = connection.execute(
                """SELECT COALESCE(SUM(c.amount),0)
                   FROM invoice_components c
                   JOIN facturas f ON f.id_factura=c.id_factura
                   WHERE f.id_comunidad=? AND f.id_periodo=?
                     AND f.tipo_suministro=? AND c.component_key=?""",
                (case["id_comunidad"], case["id_periodo"], module, component),
            ).fetchone()[0]
        elif key.startswith("parameter:"):
            parameter_key = key.split(":", 1)[1]
            row = connection.execute(
                """SELECT numeric_value FROM period_parameters
                   WHERE id_comunidad=? AND id_periodo=? AND parameter_key=?""",
                (case["id_comunidad"], case["id_periodo"], parameter_key),
            ).fetchone()
            if row is None or row[0] is None:
                # Los importes que el libro calcula en su análisis no se
                # concilian contra la base: no hay nada con lo que compararlos,
                # son su resultado. Se comprueban las entradas que los
                # alimentan (facturas, lecturas y cuotas cobradas).
                if parameter_key in calculados_por_el_libro:
                    continue
                if parameter_key in required_parameters:
                    raise ExportBlockedError(
                        f"Falta el parámetro de conciliación {parameter_key}"
                    )
                # Un concepto opcional inactivo deja intacta su sección de la
                # plantilla y concilia contra cero, no contra un dato inventado.
                value = 0
            else:
                value = row[0]
        elif key.startswith("reading_total:"):
            reading_type = key.split(":", 1)[1]
            if reading_type not in {"ACS", "CALEFACCION"}:
                raise ExportBlockedError(
                    f"Tipo de lectura no soportado en la conciliación: {reading_type}"
                )
            # Las mismas fechas de lectura que usa la hoja, no las del
            # ejercicio: si no, el total esperado no cuadraría con lo escrito.
            fecha_inicial, fecha_final = _boundary_reading_dates(
                connection, case, reading_type,
            )
            rows = connection.execute(
                """SELECT p.id_propietario,l.fecha_lectura,l.valor_acumulado
                   FROM propietarios p JOIN period_readings l
                     ON l.id_propietario=p.id_propietario
                   WHERE p.id_comunidad=? AND p.activo=1 AND p.tipo_unidad='vivienda'
                     AND l.id_periodo=? AND l.tipo=?
                     AND l.fecha_lectura IN (?,?)
                   ORDER BY p.id_propietario,l.fecha_lectura""",
                (
                    case["id_comunidad"], case["id_periodo"], reading_type,
                    fecha_inicial, fecha_final,
                ),
            ).fetchall()
            values_by_owner: dict[int, dict[str, float]] = {}
            for row in rows:
                values_by_owner.setdefault(int(row["id_propietario"]), {})[
                    row["fecha_lectura"]
                ] = float(row["valor_acumulado"])
            value = sum(
                readings[fecha_final] - readings[fecha_inicial]
                for readings in values_by_owner.values()
                if fecha_inicial in readings and fecha_final in readings
            )
        else:
            raise ExportBlockedError(f"Comprobación de total no soportada: {key}")
        result[key] = _cents(value)
    return result


def _set_run(
    connection: sqlite3.Connection,
    run_id: int,
    status: str,
    *,
    output_path: Path | None = None,
    diagnostic_path: Path | None = None,
    error_message: str | None = None,
    published: bool = False,
) -> None:
    connection.execute(
        """UPDATE excel_export_runs
           SET status=?,output_path=COALESCE(?,output_path),
               diagnostic_path=COALESCE(?,diagnostic_path),error_message=?,
               updated_at=datetime('now'),
               published_at=CASE WHEN ? THEN datetime('now') ELSE published_at END
           WHERE id_export_run=?""",
        (
            status,
            str(output_path) if output_path else None,
            str(diagnostic_path) if diagnostic_path else None,
            error_message,
            1 if published else 0,
            run_id,
        ),
    )
    connection.commit()


def _store_computed_parameters(
    connection: sqlite3.Connection,
    case: sqlite3.Row,
    profile: ExcelProfile,
    workbook_path: Path,
) -> dict[str, float]:
    """Devuelve a la base los importes que el libro acaba de calcular.

    El análisis del Excel es quien reparte el coste entre ACS y calefacción y
    quien suma lo cobrado, siguiendo las fórmulas del despacho. El reparto por
    propietario y las cartas necesitan esas mismas cifras, así que se leen del
    libro ya recalculado y se guardan como parámetros del período. Así el Excel
    sigue siendo el motor de cálculo y la base, su registro.
    """
    workbook = load_workbook(workbook_path, data_only=True)
    try:
        guardados: dict[str, float] = {}
        for clave, (hoja, direccion) in profile.workbook_layout["parameter_cells"].items():
            if hoja not in workbook.sheetnames:
                continue
            valor = workbook[hoja][direccion].value
            if not isinstance(valor, (int, float)):
                continue
            connection.execute(
                """INSERT INTO period_parameters
                   (id_comunidad,id_periodo,parameter_key,numeric_value,unit,source_sheet,source_cell)
                   VALUES (?,?,?,?,'EUR',?,?)
                   ON CONFLICT(id_comunidad,id_periodo,parameter_key) DO UPDATE SET
                       numeric_value=excluded.numeric_value,
                       source_sheet=excluded.source_sheet,
                       source_cell=excluded.source_cell""",
                (
                    case["id_comunidad"], case["id_periodo"], clave,
                    float(valor), hoja, direccion,
                ),
            )
            guardados[clave] = float(valor)
        connection.commit()
        return guardados
    finally:
        workbook.close()



def generate_official_excel(
    connection: sqlite3.Connection,
    *,
    id_case: int,
    project_root: Path,
    output_root: Path,
    progress: ProgressCallback | None = None,
    recalculator: WorkbookRecalculator | None = None,
) -> ExportResult:
    """Genera, recalcula, valida y publica el libro de un expediente listo."""
    _emit(progress, "validate_case", id_case=id_case)
    case = _case_context(connection, id_case)
    project_root = Path(project_root).resolve()
    output_root = Path(output_root).resolve()
    profile = _profile_for_community(
        connection, project_root, case["codigo"], int(case["id_comunidad"])
    )
    if profile.community_code != str(case["codigo"]):
        raise ExportBlockedError("El perfil no corresponde a la comunidad del expediente")
    _validate_normalized_inputs(connection, case, profile)
    profile_id, template, template_hash = _template_registration(
        connection, case, profile, project_root
    )
    input_hash = calculate_case_input_hash(
        connection, id_case=id_case, project_root=project_root, profile=profile,
    )
    cursor = connection.execute(
        """INSERT INTO excel_export_runs
           (id_case,id_periodo,id_template_profile,input_sha256,template_sha256,status)
           VALUES (?,?,?,?,?,'validating')""",
        (id_case, case["id_periodo"], profile_id, input_hash, template_hash),
    )
    run_id = int(cursor.lastrowid)
    connection.commit()

    official_directory = output_root / "Excels_Maestros"
    output_path = official_directory / f"Comunidad_{case['codigo']}.xlsx"
    work_directory = output_root / ".excel_exports" / f"run_{run_id}"
    temporary_path = work_directory / f"Comunidad_{case['codigo']}.xlsx"
    backup_path: Path | None = None
    predecessor_existed = output_path.is_file()
    published = False
    try:
        _emit(progress, "prepare_template", template=str(template))
        work_directory.mkdir(parents=True, exist_ok=False)
        shutil.copy2(template, temporary_path)
        template_fingerprint = workbook_fingerprint(template, profile)
        _set_run(connection, run_id, "generating", diagnostic_path=temporary_path)
        _write_workbook(connection, temporary_path, case, profile, progress)
        _restore_missing_ooxml_parts(template, temporary_path)

        _emit(progress, "recalculate", id_export_run=run_id)
        _set_run(connection, run_id, "recalculating")
        engine = recalculator or LibreOfficeRecalculator()
        engine.recalculate(temporary_path, work_directory)
        _restore_missing_ooxml_parts(template, temporary_path)

        _emit(progress, "reconcile", id_export_run=run_id)
        validate_workbook(
            temporary_path,
            profile,
            _expected_totals(connection, case, profile),
            expected_fingerprint=template_fingerprint,
        )

        calculados = _store_computed_parameters(
            connection, case, profile, temporary_path,
        )
        if calculados:
            _emit(progress, "parameters", stored=len(calculados))
            # Guardar lo que el propio libro acaba de calcular cambia la huella
            # de entradas del expediente. Sin actualizarla, el reparto creería
            # que los datos se tocaron después de validar el Excel y exigiría
            # regenerarlo una y otra vez.
            input_hash = calculate_case_input_hash(
                connection, id_case=id_case, project_root=project_root, profile=profile,
            )
            connection.execute(
                "UPDATE excel_export_runs SET input_sha256=? WHERE id_export_run=?",
                (input_hash, run_id),
            )
            connection.commit()

        _emit(progress, "publish", output=str(output_path))
        official_directory.mkdir(parents=True, exist_ok=True)
        if predecessor_existed:
            backup_directory = official_directory / "backups"
            backup_directory.mkdir(parents=True, exist_ok=True)
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
            backup_path = backup_directory / f"Comunidad_{case['codigo']}_{timestamp}.xlsx"
            shutil.copy2(output_path, backup_path)
        os.replace(temporary_path, output_path)
        published = True
        _set_run(
            connection, run_id, "validated", output_path=output_path,
            diagnostic_path=output_path, published=True,
        )
        try:
            work_directory.rmdir()
        except OSError:
            pass
        return ExportResult(run_id, output_path, backup_path, "validated", input_hash)
    except Exception as error:
        diagnostic = temporary_path if temporary_path.is_file() else None
        if published:
            if diagnostic is None and output_path.is_file():
                diagnostic = work_directory / "failed_after_publish.xlsx"
                diagnostic.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(output_path, diagnostic)
            if backup_path is not None and backup_path.is_file():
                shutil.copy2(backup_path, output_path)
            elif not predecessor_existed and output_path.is_file():
                output_path.unlink()
        _set_run(
            connection, run_id, "failed", diagnostic_path=diagnostic,
            error_message=f"{type(error).__name__}: {error}",
        )
        raise
