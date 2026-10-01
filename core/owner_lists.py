"""Listados de propietarios en cualquier formato habitual.

Los importadores trabajan con un CSV canónico (``Fdenominacion;Nombre;
Coeficiente;Email``), que es la exportación de un programa concreto. Los
despachos tienen sus listados en Excel o en CSV con otras columnas y
separadores; aquí se reconocen por el significado de sus cabeceras y se
convierten a ese CSV canónico, sin tocar el archivo original.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

import keywords


CANONICAL_HEADERS = ("Fdenominacion", "Nombre", "Coeficiente", "Email")
_EMAIL = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.IGNORECASE)


@dataclass(frozen=True)
class OwnerRow:
    vivienda: str
    nombre: str
    coeficiente: str
    email: str


def _read_rows(path: Path) -> list[list[object]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t|")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";" if text.count(";") > text.count(",") else ","
        return [row for row in csv.reader(io.StringIO(text), dialect)]
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            best: list[list[object]] = []
            for sheet in book.worksheets:
                rows = [list(row) for row in sheet.iter_rows(values_only=True)]
                if _find_header(rows) is not None:
                    return rows
                best = best or rows
            return best
        finally:
            book.close()
    if suffix == ".xls":
        import xlrd
        book = xlrd.open_workbook(path)
        for index in range(book.nsheets):
            sheet = book.sheet_by_index(index)
            rows = [sheet.row_values(row) for row in range(sheet.nrows)]
            if _find_header(rows) is not None:
                return rows
        return []
    raise ValueError(f"Formato de listado no compatible: {suffix}")


def _find_header(rows: Sequence[Sequence[object]]):
    roles = keywords.section("cabeceras_propietarios")
    for index, row in enumerate(rows[:30]):
        columns: dict[str, int] = {}
        for column, cell in enumerate(row):
            match = keywords.best_role(cell, roles, threshold=90)
            if match and match[0] not in columns:
                columns[match[0]] = column
        if "vivienda" in columns and "nombre" in columns:
            return index, columns
    return None


def _clean(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value or "").split())


def _coefficient(cells: list[object], column: int | None) -> str:
    """Coeficiente con coma decimal, que es lo que espera el importador.

    Un número de Excel (2.5) escrito tal cual se leería como 25.
    """
    if column is None or column >= len(cells) or cells[column] in (None, ""):
        return ""
    value = cells[column]
    if isinstance(value, (int, float)):
        text = f"{float(value):.6f}".rstrip("0").rstrip(".")
        return text.replace(".", ",")
    text = _clean(value).replace("%", "").strip()
    if "." in text and "," not in text and re.fullmatch(r"\d+\.\d+", text):
        return text.replace(".", ",")
    return text


def read_owner_rows(path: str | Path) -> list[OwnerRow]:
    """Filas de propietarios del listado; ``ValueError`` si no se reconoce."""
    rows = _read_rows(Path(path))
    found = _find_header(rows)
    if found is None:
        raise ValueError(
            "No se reconocen las columnas del listado: se necesita al menos vivienda y nombre."
        )
    header_index, columns = found
    result: list[OwnerRow] = []
    for raw in rows[header_index + 1:]:
        cells = list(raw)

        def cell(role: str) -> str:
            column = columns.get(role)
            return _clean(cells[column]) if column is not None and column < len(cells) else ""

        vivienda, nombre = cell("vivienda"), cell("nombre")
        email_match = _EMAIL.search(cell("email") or " ".join(_clean(value) for value in cells))
        email = email_match.group(0).lower() if email_match else ""
        if not vivienda and (not nombre or _EMAIL.fullmatch(nombre)):
            # Algunos programas imprimen el correo en la línea siguiente.
            if email and result and not result[-1].email:
                last = result[-1]
                result[-1] = OwnerRow(last.vivienda, last.nombre, last.coeficiente, email)
            continue
        if not vivienda or not nombre:
            continue
        result.append(OwnerRow(vivienda, nombre, _coefficient(cells, columns.get("coeficiente")), email))
    if not result:
        raise ValueError("El listado no contiene propietarios con vivienda y nombre.")
    return result


def is_canonical(path: str | Path) -> bool:
    path = Path(path)
    if path.suffix.lower() != ".csv":
        return False
    with path.open("r", encoding="utf-8-sig", errors="replace") as handle:
        first = handle.readline()
    return [item.strip() for item in first.strip().split(";")][:2] == list(CANONICAL_HEADERS[:2])


def normalise_owner_list(path: str | Path, destination_dir: str | Path) -> Path:
    """Devuelve un CSV canónico: el original si ya lo es, o una conversión."""
    path = Path(path)
    if not path.is_file() or is_canonical(path):
        # Un archivo inexistente lo rechaza después el análisis con su mensaje.
        return path
    rows = read_owner_rows(path)
    destination = Path(destination_dir) / f"{path.stem}_propietarios.csv"
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(CANONICAL_HEADERS)
        for row in rows:
            writer.writerow((row.vivienda, row.nombre, row.coeficiente, row.email))
    return destination


__all__ = ["CANONICAL_HEADERS", "OwnerRow", "is_canonical", "normalise_owner_list", "read_owner_rows"]
