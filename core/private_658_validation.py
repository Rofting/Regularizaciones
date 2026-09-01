"""Runner explícito para validar fuentes privadas fuera del proyecto.

No forma parte del flujo ordinario ni se ejecuta desde la interfaz.  Requiere
las cuatro variables de entorno documentadas y prepara una raíz nueva para que
la comprobación de una comunidad real no pueda alterar ni las fuentes ni el
proyecto de trabajo.
"""

from __future__ import annotations

import argparse
import os
import re
import secrets
import shutil
import sqlite3
import subprocess
import sys
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import load_workbook

import case_workflow_actions
import document_review
import expedient_service
import gestor_bd
from excel_profiles import load_profile


REQUIRED_ENVIRONMENT = (
    "REGULARIZACION_658_MASTER",
    "REGULARIZACION_658_PROPIETARIOS",
    "REGULARIZACION_658_LECTURAS",
    "REGULARIZACION_658_VALIDATION_ROOT",
)
PUBLIC_PROFILE_PATH = Path("config") / "excel_profiles" / "658_acs_v1.json"


class PrivateValidationBlockedError(ValueError):
    """No se han cumplido las guardas de la validación privada."""


@dataclass(frozen=True)
class ValidationEnvironment:
    master: Path
    owners: Path
    readings: Path
    validation_root: Path


@dataclass(frozen=True)
class PrivateValidationResult:
    exit_code: int
    status: str
    run_root: Path
    issue_counts: Mapping[str, int]
    generated_letters: int = 0


def _is_relative_to(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False


def validation_environment(environ: Mapping[str, str] | None = None) -> ValidationEnvironment:
    """Lee únicamente la configuración explícita y valida que son archivos."""
    values = os.environ if environ is None else environ
    missing = [name for name in REQUIRED_ENVIRONMENT if not str(values.get(name, "")).strip()]
    if missing:
        raise PrivateValidationBlockedError(
            "Faltan variables obligatorias: " + ", ".join(missing)
        )
    master = Path(str(values["REGULARIZACION_658_MASTER"])).expanduser().resolve()
    owners = Path(str(values["REGULARIZACION_658_PROPIETARIOS"])).expanduser().resolve()
    readings = Path(str(values["REGULARIZACION_658_LECTURAS"])).expanduser().resolve()
    for source in (master, owners, readings):
        if not source.is_file():
            raise PrivateValidationBlockedError("Cada fuente privada debe ser un archivo existente")
    root = Path(str(values["REGULARIZACION_658_VALIDATION_ROOT"])).expanduser().resolve()
    for source in (master, owners, readings):
        if _is_relative_to(root, source.parent) or _is_relative_to(source, root):
            raise PrivateValidationBlockedError("La raíz de validación no es segura")
    return ValidationEnvironment(master, owners, readings, root)


def _require_safe_validation_root(validation_root: Path, project_root: Path) -> None:
    root = validation_root.resolve()
    project = project_root.resolve()
    home = Path.home().resolve()
    volume_root = Path(root.anchor).resolve()
    if root in {volume_root, home, home.parent.resolve()}:
        raise PrivateValidationBlockedError("La raíz de validación no es segura")
    if root == project or _is_relative_to(project, root) or _is_relative_to(root, project):
        raise PrivateValidationBlockedError("La raíz de validación no es segura")
    if root.exists() and not root.is_dir():
        raise PrivateValidationBlockedError("La raíz de validación no es segura")
    if root.exists() and any(root.iterdir()):
        raise PrivateValidationBlockedError("La raíz de validación debe ser nueva o estar vacía")


def prepare_validation_run(validation_root: Path, *, project_root: Path) -> Path:
    """Crea una ejecución aislada copiando solamente configuración y Word público."""
    destination = Path(validation_root).resolve()
    source_project = Path(project_root).resolve()
    _require_safe_validation_root(destination, source_project)
    profile_source = source_project / PUBLIC_PROFILE_PATH
    if not profile_source.is_file():
        raise PrivateValidationBlockedError("No existe el perfil público requerido")
    public_word = source_project / "plantillas" / "Plantilla_Cartas.docx"
    if not public_word.is_file():
        raise PrivateValidationBlockedError("No existe la plantilla Word pública requerida")
    destination.mkdir(parents=True, exist_ok=True)
    run = destination / ("validacion_" + datetime.now().strftime("%Y%m%d_%H%M%S") + "_" + secrets.token_hex(4))
    run.mkdir()
    profile_destination = run / PUBLIC_PROFILE_PATH
    profile_destination.parent.mkdir(parents=True)
    shutil.copy2(profile_source, profile_destination)
    (run / "plantillas").mkdir()
    shutil.copy2(public_word, run / "plantillas" / public_word.name)
    return run


def _extract_master_period(master: Path) -> tuple[date, date]:
    """Obtiene las fechas estructurales del encabezado sin registrar su texto."""
    workbook = load_workbook(master, read_only=True, data_only=True)
    try:
        pattern = re.compile(r"(\d{1,2}[/-]\d{1,2}[/-]\d{4})")
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if not isinstance(value, str) or "PERIODO" not in value.upper():
                        continue
                    found = pattern.findall(value)
                    if len(found) >= 2:
                        dates = [datetime.strptime(item.replace("-", "/"), "%d/%m/%Y").date() for item in found[:2]]
                        return min(dates), max(dates)
    finally:
        workbook.close()
    raise PrivateValidationBlockedError("No se pudo determinar el período estructural del maestro")


def _issue_counts(connection: sqlite3.Connection, id_case: int) -> dict[str, int]:
    rows = connection.execute(
        """SELECT code,COUNT(*) AS total FROM review_issues
           WHERE id_case=? AND status='open' GROUP BY code ORDER BY code""",
        (id_case,),
    ).fetchall()
    return {str(row["code"]): int(row["total"]) for row in rows}


def render_sanitized_report(
    *,
    status: str,
    run_root: Path,
    issue_counts: Mapping[str, int],
    generated_letters: int,
    source_paths: Iterable[Path] = (),
    failures: Iterable[str] = (),
) -> str:
    """Produce evidencia operativa sin identificar archivos ni propietarios."""
    del source_paths, failures
    lines = [
        "VALIDACION_PRIVADA_658",
        f"estado: {status}",
        f"salida: {Path(run_root).resolve()}",
        f"cartas_generadas: {int(generated_letters)}",
        "incidencias_abiertas:",
    ]
    if issue_counts:
        lines.extend(f"- {code}: {int(total)}" for code, total in sorted(issue_counts.items()))
    else:
        lines.append("- ninguna")
    return "\n".join(lines) + "\n"


def _write_report(result: PrivateValidationResult) -> Path:
    report = result.run_root / "informe_validacion.txt"
    report.write_text(
        render_sanitized_report(
            status=result.status,
            run_root=result.run_root,
            issue_counts=result.issue_counts,
            generated_letters=result.generated_letters,
        ),
        encoding="utf-8",
    )
    return report


def locate_libreoffice() -> Path:
    """Encuentra una instalación local compatible, sin descargar ni instalar nada."""
    configured = os.environ.get("LIBREOFFICE_PATH", "").strip()
    candidates = [Path(configured)] if configured else []
    executable = shutil.which("soffice")
    if executable:
        candidates.append(Path(executable))
    candidates.extend((
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
    ))
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    raise PrivateValidationBlockedError("LibreOffice no está disponible para validar el documento")


def _render_one_page(document: Path, output_directory: Path) -> None:
    office = locate_libreoffice()
    output_directory.mkdir(parents=True, exist_ok=True)
    completed = subprocess.run(
        [str(office), "--headless", "--convert-to", "pdf", "--outdir", str(output_directory), str(document)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=90,
        check=False,
    )
    pdf = output_directory / f"{document.stem}.pdf"
    if completed.returncode != 0 or not pdf.is_file():
        raise PrivateValidationBlockedError("No se pudo renderizar una carta con LibreOffice")
    import pypdfium2

    document_pdf = pypdfium2.PdfDocument(str(pdf))
    if len(document_pdf) != 1:
        raise PrivateValidationBlockedError("La carta renderizada no cabe en una página")


def _require_safe_existing_run(run_root: Path, project_root: Path) -> Path:
    """Valida una ejecución aislada existente sin crear ni borrar contenido."""
    run = Path(run_root).resolve()
    project = Path(project_root).resolve()
    home = Path.home().resolve()
    volume_root = Path(run.anchor).resolve()
    if (
        run in {volume_root, home, home.parent.resolve(), project}
        or _is_relative_to(run, project)
        or _is_relative_to(project, run)
    ):
        raise PrivateValidationBlockedError("La ruta de reanudación no es segura")
    if not run.is_dir() or not run.name.startswith("validacion_"):
        raise PrivateValidationBlockedError("La ruta no identifica una ejecución de validación")
    required = (
        run / PUBLIC_PROFILE_PATH,
        run / "plantillas" / "Plantilla_Cartas.docx",
        run / "data" / "gestion.db",
    )
    if not all(path.is_file() and _is_relative_to(path.resolve(), run) for path in required):
        raise PrivateValidationBlockedError("La ejecución existente está incompleta")
    return run


def _single_validation_case(
    connection: sqlite3.Connection,
) -> sqlite3.Row:
    rows = connection.execute(
        """SELECT c.id_case,c.id_comunidad,c.estado,co.codigo
           FROM regularization_cases c
           JOIN comunidades co ON co.id_comunidad=c.id_comunidad
           ORDER BY c.id_case"""
    ).fetchall()
    if len(rows) != 1 or str(rows[0]["codigo"]) != "658":
        raise PrivateValidationBlockedError(
            "La base de validación no contiene un único expediente compatible"
        )
    return rows[0]


def _validated_export_exists(
    connection: sqlite3.Connection,
    id_case: int,
    run_root: Path,
) -> bool:
    row = connection.execute(
        """SELECT output_path FROM excel_export_runs
           WHERE id_case=? AND status='validated'
           ORDER BY created_at DESC,id_export_run DESC LIMIT 1""",
        (id_case,),
    ).fetchone()
    if row is None or not row["output_path"]:
        return False
    output = Path(str(row["output_path"])).resolve()
    return _is_relative_to(output, run_root) and output.is_file()


def _completed_letter_output(
    connection: sqlite3.Connection,
    id_case: int,
    run_root: Path,
) -> tuple[int, Path] | None:
    run = connection.execute(
        """SELECT id_letter_run,output_path FROM letter_generation_runs
           WHERE id_case=? AND status='completed'
           ORDER BY completed_at DESC,id_letter_run DESC LIMIT 1""",
        (id_case,),
    ).fetchone()
    if run is None or not run["output_path"]:
        return None
    output = Path(str(run["output_path"])).resolve()
    if not _is_relative_to(output, run_root) or not output.is_dir():
        return None
    letters = connection.execute(
        """SELECT output_path,status FROM generated_letters
           WHERE id_letter_run=? ORDER BY id_generated_letter""",
        (run["id_letter_run"],),
    ).fetchall()
    verified = [
        Path(str(row["output_path"])).resolve()
        for row in letters
        if row["status"] == "generated" and row["output_path"]
    ]
    if not verified or len(verified) != len(letters):
        return None
    if not all(_is_relative_to(path, output) and path.is_file() for path in verified):
        return None
    return len(verified), verified[0]


def resume_private_658_validation(
    run_root: Path,
    *,
    project_root: Path | None = None,
) -> PrivateValidationResult:
    """Continúa una ejecución aislada usando su misma base y sus mismas salidas.

    No vuelve a importar fuentes. Si quedan incidencias abiertas, sólo actualiza
    el informe saneado. Cada etapa posterior se ejecuta únicamente cuando falta
    y las salidas completadas se aceptan sólo si conservan su auditoría y están
    dentro de la carpeta de ejecución.
    """
    source_project = Path(project_root or Path(__file__).resolve().parents[1]).resolve()
    run = _require_safe_existing_run(Path(run_root), source_project)
    database_path = run / "data" / "gestion.db"
    connection = gestor_bd.conectar(str(database_path))
    try:
        case = _single_validation_case(connection)
        id_case = int(case["id_case"])
        community_id = int(case["id_comunidad"])
        issues = _issue_counts(connection, id_case)
        if issues:
            result = PrivateValidationResult(2, "blocked_open_issues", run, issues)
            _write_report(result)
            return result
        status = str(case["estado"])
        if status in {"draft", "gathering_sources", "under_review"}:
            status = document_review.validate_case_ready(connection, id_case).status
        has_export = _validated_export_exists(connection, id_case, run)
    finally:
        connection.close()

    if status == "ready_for_calculation":
        if not has_export:
            export = case_workflow_actions.run_generate_excel(
                database_path,
                id_case=id_case,
                active_community_id=community_id,
                project_root=run,
                output_root=run / "salidas",
            )
            if not export.output_path.is_file():
                raise PrivateValidationBlockedError("No se publicó el Excel oficial")
        case_workflow_actions.run_calculate_distribution(
            database_path,
            id_case=id_case,
            active_community_id=community_id,
            project_root=run,
        )
        status = "reconciled"
    elif status == "calculated":
        case_workflow_actions.run_calculate_distribution(
            database_path,
            id_case=id_case,
            active_community_id=community_id,
            project_root=run,
        )
        status = "reconciled"

    if status == "reconciled":
        concepts = case_workflow_actions.available_case_letter_concepts(
            database_path,
            id_case=id_case,
            active_community_id=community_id,
            project_root=run,
        )
        letters = case_workflow_actions.run_generate_letters(
            database_path,
            id_case=id_case,
            active_community_id=community_id,
            project_root=run,
            selected_concepts=tuple(key for key, _label in concepts),
        )
        if letters.failures:
            result = PrivateValidationResult(
                1, "incomplete_letters", run, {}, letters.generated_count
            )
            _write_report(result)
            return result
        generated_count = letters.generated_count
        document = next(iter(sorted(letters.output_path.glob("*.docx"))), None)
    elif status in {"deliveries_generated", "closed"}:
        connection = gestor_bd.conectar(str(database_path))
        try:
            completed = _completed_letter_output(connection, id_case, run)
        finally:
            connection.close()
        if completed is None:
            raise PrivateValidationBlockedError(
                "La salida de cartas completada no conserva una auditoría válida"
            )
        generated_count, document = completed
    else:
        raise PrivateValidationBlockedError(
            f"El expediente no puede reanudarse desde el estado {status}"
        )

    if document is None or not document.is_file():
        raise PrivateValidationBlockedError("No se generó ninguna carta para validar")
    rendered = run / "render" / f"{document.stem}.pdf"
    if not rendered.is_file():
        _render_one_page(document, run / "render")
    result = PrivateValidationResult(0, "completed", run, {}, generated_count)
    _write_report(result)
    return result


def run_private_658_validation(
    *,
    environ: Mapping[str, str] | None = None,
    project_root: Path | None = None,
) -> PrivateValidationResult:
    """Ejecuta la validación aislada; nunca se invoca desde la suite ordinaria."""
    environment = validation_environment(environ)
    original_root = Path(project_root or Path(__file__).resolve().parents[1]).resolve()
    run_root = prepare_validation_run(environment.validation_root, project_root=original_root)
    database_path = run_root / "data" / "gestion.db"
    profile = load_profile("658_acs_v1", run_root)
    start_date, end_date = _extract_master_period(environment.master)
    gestor_bd.crear_bd(str(database_path))
    connection = gestor_bd.conectar(str(database_path))
    try:
        community_id = gestor_bd.obtener_o_crear_comunidad(
            connection, profile.community_code, f"Comunidad {profile.community_code}"
        )
        case = expedient_service.create_case(
            connection, community_id,
            name=f"Regularización {start_date.isoformat()}_{end_date.isoformat()}",
            start_date=start_date, end_date=end_date,
        )
    finally:
        connection.close()

    case_workflow_actions.run_bootstrap_import(
        database_path, id_case=case.id_case, active_community_id=community_id,
        project_root=run_root, master_path=environment.master,
        owner_list_path=environment.owners, readings_path=environment.readings,
        actor="validacion_privada",
    )
    connection = gestor_bd.conectar(str(database_path))
    try:
        issues = _issue_counts(connection, case.id_case)
    finally:
        connection.close()
    if issues:
        result = PrivateValidationResult(2, "blocked_open_issues", run_root, issues)
        _write_report(result)
        return result

    export = case_workflow_actions.run_generate_excel(
        database_path, id_case=case.id_case, active_community_id=community_id,
        project_root=run_root, output_root=run_root / "salidas",
    )
    # El exportador ya valida estructura, fórmulas y publicación antes de devolver.
    if not export.output_path.is_file():
        raise PrivateValidationBlockedError("No se publicó el Excel oficial")
    case_workflow_actions.run_calculate_distribution(
        database_path, id_case=case.id_case, active_community_id=community_id,
        project_root=run_root,
    )
    concepts = case_workflow_actions.available_case_letter_concepts(
        database_path, id_case=case.id_case, active_community_id=community_id,
        project_root=run_root,
    )
    letters = case_workflow_actions.run_generate_letters(
        database_path, id_case=case.id_case, active_community_id=community_id,
        project_root=run_root, selected_concepts=tuple(key for key, _label in concepts),
    )
    if letters.failures:
        result = PrivateValidationResult(1, "incomplete_letters", run_root, {}, letters.generated_count)
        _write_report(result)
        return result
    document = next(iter(sorted(letters.output_path.glob("*.docx"))), None)
    if document is None:
        raise PrivateValidationBlockedError("No se generó ninguna carta para validar")
    _render_one_page(document, run_root / "render")
    result = PrivateValidationResult(0, "completed", run_root, {}, letters.generated_count)
    _write_report(result)
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Validación privada aislada de la comunidad 658")
    parser.add_argument(
        "--resume",
        type=Path,
        metavar="CARPETA_EJECUCION",
        help="continúa una carpeta validacion_* existente sin reimportar fuentes",
    )
    args = parser.parse_args(argv)
    try:
        result = (
            resume_private_658_validation(args.resume)
            if args.resume is not None
            else run_private_658_validation()
        )
    except PrivateValidationBlockedError as error:
        print(f"VALIDACION_PRIVADA_658: bloqueada ({error})")
        return 2
    except Exception as error:
        print(f"VALIDACION_PRIVADA_658: error técnico ({type(error).__name__})")
        return 1
    print(render_sanitized_report(
        status=result.status, run_root=result.run_root,
        issue_counts=result.issue_counts, generated_letters=result.generated_letters,
    ), end="")
    return result.exit_code


if __name__ == "__main__":
    raise SystemExit(main())
