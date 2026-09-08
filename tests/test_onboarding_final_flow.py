"""Acceptance checks use source files and the public onboarding/import path."""

import sqlite3
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from datetime import date
from io import StringIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "core"))

import case_workflow_actions
import community_onboarding
import excel_generator
import gestor_bd
from case_distribution import calculate_case_distribution
from excel_export_service import generate_official_excel
from office_recalculation import DeterministicRecalculator


class OnboardingSourceFlowTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.database = self.root / "gestion.db"
        with redirect_stdout(StringIO()):
            gestor_bd.crear_bd(str(self.database))
        self.connection = sqlite3.connect(self.database)
        self.connection.row_factory = sqlite3.Row
        self.addCleanup(self.connection.close)
        self.owners = self.root / "owners.csv"
        self.owners.write_text(
            "Fdenominacion;Nombre;Coeficiente\nA-01;Vecino sintetico;100\n",
            encoding="utf-8",
        )
        self.readings = self.root / "readings.xlsx"
        self.headers = (
            "Vivienda", "Contador ACS", "Fecha", "Lectura ACS",
            "Contador Calefacción", "Lectura Calefacción",
        )
        self.write_readings((
            ("A-01", "C-ACS", "2026-01-01", 100, "C-CAL", 200),
            ("A-01", "C-ACS", "2026-12-31", 130, "C-CAL", 500),
        ))

    def write_readings(self, rows):
        workbook = Workbook()
        workbook.active.append(self.headers)
        for row in rows:
            workbook.active.append(row)
        workbook.save(self.readings)
        workbook.close()

    def analyse(self, code="FINAL"):
        return community_onboarding.analyse_sources(
            community_code=code, community_name="Comunidad sintetica",
            owner_list_path=self.owners, reading_paths=(self.readings,),
            invoice_paths=(), project_root=self.root,
        )

    def answers(self, service):
        answers = {"service": service}
        for module in service.split("+"):
            heating = module == "CALEFACCION"
            answers.update({
                f"reading_column:{module}": "Lectura Calefacción" if heating else "Lectura ACS",
                f"reading_meter:{module}": "C-CAL" if heating else "C-ACS",
                f"reading_date:{module}": "2026-12-31",
                f"reading_value:{module}": "500" if heating else "130",
            })
        return answers

    def publish(self, service):
        return community_onboarding.confirm_onboarding(
            self.connection, draft=self.analyse(), answers=self.answers(service),
            project_root=self.root, archive_root=self.root / "expedientes",
            period_name="2026", start_date=date(2026, 1, 1),
            end_date=date(2026, 12, 31), actor="Prueba",
        )

    def test_excel_evidence_relates_real_values_to_each_module(self):
        configuration = community_onboarding.resolve_onboarding_configuration(
            self.analyse(), self.answers("ACS+CALEFACCION")
        )
        self.assertEqual(
            [("ACS", "C-ACS", "2026-12-31", "130"),
             ("CALEFACCION", "C-CAL", "2026-12-31", "500")],
            [(item.module, item.meter, item.date, item.value)
             for item in configuration.reading_bindings],
        )

    def test_headers_without_data_require_meter_date_and_value_answers(self):
        self.write_readings(())
        draft = self.analyse()
        required = {question.key for question in draft.questions if question.required}
        for module in ("ACS", "CALEFACCION"):
            self.assertTrue({f"reading_{field}:{module}" for field in
                             ("meter", "date", "value")}.issubset(required))

    def test_cannot_confirm_a_value_from_the_other_module_or_date(self):
        for field, value in (("reading_value:ACS", "500"),
                             ("reading_date:ACS", "2026-01-01")):
            with self.subTest(field=field):
                answers = self.answers("ACS")
                answers[field] = value
                with self.assertRaises(ValueError):
                    community_onboarding.resolve_onboarding_configuration(
                        self.analyse(), answers
                    )

    def exercise_flow(self, service):
        result = self.publish(service)
        workbook = load_workbook(result.template_path)
        values = {
            "ACS": ("LECTURAS ACS M3", 80, 20, 40, 10),
            "CALEFACCION": ("LECTURAS CALEF KWH", 120, 50, 60, 25),
        }
        for module in service.split("+"):
            name, variable, fixed, billed_variable, billed_fixed = values[module]
            sheet = workbook[name]
            sheet["H8"], sheet["I9"] = variable, fixed
            sheet["H4"], sheet["I4"] = billed_variable, billed_fixed
        workbook.save(result.template_path)
        workbook.close()
        bootstrap, companions = case_workflow_actions.run_bootstrap_import(
            self.database, id_case=result.case_id,
            active_community_id=result.community_id, project_root=self.root,
            master_path=result.template_path, owner_list_path=self.owners,
            readings_path=self.readings, actor="Prueba",
        )
        self.assertEqual(0, bootstrap.open_issue_count)
        self.assertEqual(0, companions.open_issue_count)
        self.assertEqual(2 * len(service.split("+")), companions.imported_reading_count)
        readings = [tuple(row) for row in self.connection.execute(
            "SELECT tipo,fecha_lectura,valor_acumulado FROM lecturas_vecino "
            "ORDER BY tipo,fecha_lectura"
        )]
        expected = []
        if "ACS" in service:
            expected += [("ACS", "2026-01-01", 100), ("ACS", "2026-12-31", 130)]
        if "CALEFACCION" in service:
            expected += [("CALEFACCION", "2026-01-01", 200),
                         ("CALEFACCION", "2026-12-31", 500)]
        self.assertEqual(expected, readings)
        exported = generate_official_excel(
            self.connection, id_case=result.case_id, project_root=self.root,
            output_root=self.root / "salidas", recalculator=DeterministicRecalculator(),
        )
        self.assertTrue(exported.output_path.is_file())
        distribution = calculate_case_distribution(
            self.connection, id_case=result.case_id, project_root=self.root
        )
        expected_totals = {}
        if "ACS" in service:
            expected_totals.update(acs_fixed=2000, acs_variable=8000)
        if "CALEFACCION" in service:
            expected_totals.update(heating_fixed=5000, heating_variable=12000)
        self.assertEqual(expected_totals, dict(distribution.concept_totals_cents))
        self.assertEqual(
            {key: amount // 2 for key, amount in expected_totals.items()},
            dict(self.connection.execute(
                "SELECT concept_key,difference_cents FROM owner_concept_results WHERE id_periodo=?",
                (result.period_id,),
            )),
        )
        book = load_workbook(exported.output_path, data_only=True)
        try:
            for module in service.split("+"):
                sheet = book[values[module][0]]
                self.assertEqual(("Consumo", "Variable", "Fijo"),
                                 (sheet["G7"].value, sheet["H7"].value, sheet["I7"].value))
                self.assertEqual(30 if module == "ACS" else 300, sheet["G8"].value)
        finally:
            book.close()
        _, repeated = case_workflow_actions.run_bootstrap_import(
            self.database, id_case=result.case_id,
            active_community_id=result.community_id, project_root=self.root,
            master_path=result.template_path, owner_list_path=self.owners,
            readings_path=self.readings, actor="Prueba",
        )
        self.assertEqual(0, repeated.imported_reading_count)

    def test_acs_sources_reach_export_and_distribution(self):
        self.exercise_flow("ACS")

    def test_heating_sources_reach_export_and_distribution(self):
        self.exercise_flow("CALEFACCION")

    def test_combined_sources_reach_export_and_distribution(self):
        self.exercise_flow("ACS+CALEFACCION")


class CanonicalTemplateHeadersTest(unittest.TestCase):
    def test_headers_match_the_field_bindings_for_all_modules(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "template.xlsx"
            excel_generator.create_canonical_community_template(
                path, community_code="TEST", community_name="Sintetica",
                active_modules=("ACS", "CALEFACCION", "GAS", "AGUA", "ELECTRICIDAD"),
            )
            workbook = load_workbook(path)
            try:
                for sheet in ("LECTURAS ACS M3", "LECTURAS CALEF KWH"):
                    self.assertEqual(
                        ("Consumo", "Variable", "Fijo", "Total"),
                        tuple(workbook[sheet][f"{column}7"].value for column in "GHIJ"),
                    )
                for sheet in ("GAS", "AGUA", "ELECTRICIDAD"):
                    self.assertEqual(("Total", "Proveedor"),
                                     (workbook[sheet]["I9"].value, workbook[sheet]["J9"].value))
            finally:
                workbook.close()
