import json
import os
import sqlite3
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from datetime import date
from io import StringIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
CORE_DIR = PROJECT_ROOT / "core"
if str(CORE_DIR) not in sys.path:
    sys.path.insert(0, str(CORE_DIR))

import community_onboarding
import case_workflow_actions
import document_review
import expedient_service
import excel_profiles
import gestor_bd
from case_distribution import calculate_case_distribution
from excel_export_service import generate_official_excel
from office_recalculation import DeterministicRecalculator


def make_readings_workbook(path: Path, headers: tuple[str, ...]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    for column, header in enumerate(headers, start=1):
        sheet.cell(1, column, header)
    sheet.append(("A-01", "100", "120"))
    workbook.save(path)
    return path


def make_pdf(path: Path, text: str) -> Path:
    payload = (
        b"%PDF-1.4\n"
        b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
        b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n"
        b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792]"
        b" /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n"
    )
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    payload += b"4 0 obj\n<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
    payload += stream + b"\nendstream\nendobj\n"
    payload += b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n"
    payload += b"xref\n0 6\n0000000000 65535 f \n"
    offsets = []
    offset = len(b"%PDF-1.4\n")
    for object_number in range(1, 6):
        marker = f"{object_number} 0 obj".encode()
        offsets.append(payload.index(marker))
    payload += b"".join(f"{value:010d} 00000 n \n".encode() for value in offsets)
    payload += b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n"
    payload += str(payload.index(b"xref")).encode() + b"\n%%EOF\n"
    path.write_bytes(payload)
    return path


class CommunityOnboardingTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.project_root = Path(self.directory.name) / "project"
        self.project_root.mkdir()
        self.database = self.project_root / "gestion.db"
        with redirect_stdout(StringIO()):
            gestor_bd.crear_bd(str(self.database))
        self.connection = sqlite3.connect(self.database)
        self.connection.row_factory = sqlite3.Row
        self.connection.execute("PRAGMA foreign_keys = ON")
        self.connection.execute("CREATE TABLE audit (id INTEGER PRIMARY KEY)")
        self.connection.commit()
        self.archive_root = self.project_root / "expedientes"
        self.owners_csv = self.project_root / "owners.csv"
        self.owners_csv.write_text(
            "Codigo;Nombre;Fdenominacion\n1;VECINO SINTETICO;A-01\n",
            encoding="utf-8",
        )
        self.readings_xlsx = make_readings_workbook(
            self.project_root / "readings.xlsx", ("Vivienda", "Inicial", "Final")
        )
        self.readings_pdf = make_pdf(
            self.project_root / "readings.pdf", "LECTURAS ACS CONTADOR"
        )
        self.invoice_pdf = make_pdf(
            self.project_root / "invoice.pdf", "FACTURA SINTETICA"
        )
        self.ambiguous_draft = community_onboarding.OnboardingDraft(
            community_code="900",
            community_name="Comunidad prueba",
            sources=(),
            detected_modules=("ACS",),
            questions=(community_onboarding.OnboardingQuestion(
                key="reading_column",
                prompt="Selecciona la columna de lectura.",
                candidates=("Lectura inicial", "Lectura final"),
                required=True,
            ),),
        )
        self.module_draft = community_onboarding.OnboardingDraft(
            community_code="900",
            community_name="Comunidad prueba",
            sources=(community_onboarding.SourceCandidate(
                self.readings_xlsx, "meter_reading_excel",
                community_onboarding._sha256(self.readings_xlsx),
                headers=("Vivienda", "Lectura"),
            ),),
            detected_modules=("ACS", "CALEFACCION"),
            questions=(community_onboarding.OnboardingQuestion(
                key="service",
                prompt="Confirma el servicio de las lecturas.",
                candidates=("ACS", "CALEFACCION", "ACS+CALEFACCION", "NO_APLICA"),
                required=True,
            ),),
        )
        self.valid_draft = community_onboarding.OnboardingDraft(
            community_code="900",
            community_name="Comunidad prueba",
            sources=(
                community_onboarding.SourceCandidate(
                    self.owners_csv, "owner_list",
                    community_onboarding._sha256(self.owners_csv),
                ),
                community_onboarding.SourceCandidate(
                    self.readings_xlsx, "meter_reading_excel",
                    community_onboarding._sha256(self.readings_xlsx),
                    headers=("Vivienda", "Lectura"),
                ),
                community_onboarding.SourceCandidate(
                    self.invoice_pdf, "invoice_pdf",
                    community_onboarding._sha256(self.invoice_pdf),
                ),
            ),
            detected_modules=("ACS",),
            questions=(community_onboarding.OnboardingQuestion(
                key="service",
                prompt="Confirma el servicio de las lecturas.",
                candidates=("ACS", "NO_APLICA"),
                required=True,
            ),),
        )
        self.answers = {"service": "ACS"}

    def tearDown(self):
        self.connection.close()
        self.directory.cleanup()

    def _row_count(self):
        return self.connection.execute("SELECT COUNT(*) FROM audit").fetchone()[0]

    def _community_count(self):
        return self.connection.execute(
            "SELECT COUNT(*) FROM comunidades WHERE codigo='900'"
        ).fetchone()[0]

    def _registered_source_count(self):
        return self.connection.execute(
            "SELECT COUNT(*) FROM source_documents"
        ).fetchone()[0]

    def _confirmation_arguments(self):
        return {
            "draft": self.valid_draft,
            "answers": self.answers,
            "project_root": self.project_root,
            "archive_root": self.archive_root,
            "period_name": "2026",
            "start_date": date(2026, 1, 1),
            "end_date": date(2026, 12, 31),
            "actor": "Prueba",
        }

    def test_analyse_sources_classifies_owner_list_excel_readings_and_pdf_invoice_without_writes(self):
        draft = community_onboarding.analyse_sources(
            community_code="900", community_name="Comunidad prueba",
            owner_list_path=self.owners_csv, reading_paths=(self.readings_xlsx,),
            invoice_paths=(self.invoice_pdf,), project_root=self.project_root,
        )
        self.assertEqual(("owner_list", "meter_reading_excel", "invoice_pdf"),
                         tuple(item.kind for item in draft.sources))
        self.assertEqual(0, self._row_count())
        self.assertFalse((self.project_root / "config/excel_profiles/runtime/900_v1.json").exists())

    def test_analyse_sources_accepts_pdf_readings_as_an_alternative(self):
        draft = community_onboarding.analyse_sources(
            community_code="900", community_name="Comunidad prueba",
            owner_list_path=self.owners_csv, reading_paths=(self.readings_pdf,),
            invoice_paths=(), project_root=self.project_root,
        )
        self.assertEqual("meter_reading_pdf", draft.sources[1].kind)

    def test_analyse_sources_records_sha256_and_requires_confirmation_for_ambiguous_reading_columns(self):
        ambiguous_readings = make_readings_workbook(
            self.project_root / "ambiguous.xlsx", ("Vivienda", "Lectura ACS", "Lectura Agua")
        )
        draft = community_onboarding.analyse_sources(
            community_code="900", community_name="Comunidad prueba",
            owner_list_path=self.owners_csv, reading_paths=(ambiguous_readings,),
            invoice_paths=(), project_root=self.project_root,
        )
        reading = draft.sources[1]
        self.assertEqual(64, len(reading.sha256))
        self.assertTrue(any(question.required for question in draft.questions))

    def test_analyse_sources_does_not_treat_identical_reading_columns_as_ambiguous(self):
        second_readings = make_readings_workbook(
            self.project_root / "same-readings.xlsx", ("Vivienda", "Lectura ACS")
        )
        first_readings = make_readings_workbook(
            self.project_root / "first-readings.xlsx", ("Vivienda", "Lectura ACS")
        )
        draft = community_onboarding.analyse_sources(
            community_code="900", community_name="Comunidad prueba",
            owner_list_path=self.owners_csv,
            reading_paths=(first_readings, second_readings),
            invoice_paths=(), project_root=self.project_root,
        )
        self.assertFalse(any(question.key == "reading_column" for question in draft.questions))

    def test_analyse_sources_prioritizes_acs_over_partial_water_match_but_keeps_real_service_ambiguity(self):
        acs_readings = make_readings_workbook(
            self.project_root / "acs.xlsx", ("Vivienda", "Lectura Agua Caliente ACS")
        )
        acs_draft = community_onboarding.analyse_sources(
            community_code="900", community_name="Comunidad prueba",
            owner_list_path=self.owners_csv, reading_paths=(acs_readings,),
            invoice_paths=(), project_root=self.project_root,
        )
        self.assertEqual(("ACS",), acs_draft.detected_modules)
        self.assertTrue(any(question.key == "service" and question.required
                            for question in acs_draft.questions))

        ambiguous_readings = make_readings_workbook(
            self.project_root / "two-services.xlsx", ("Vivienda", "Lectura ACS", "Lectura Calefacción")
        )
        ambiguous_draft = community_onboarding.analyse_sources(
            community_code="900", community_name="Comunidad prueba",
            owner_list_path=self.owners_csv, reading_paths=(ambiguous_readings,),
            invoice_paths=(), project_root=self.project_root,
        )
        self.assertEqual(("ACS", "CALEFACCION"), ambiguous_draft.detected_modules)
        self.assertTrue(any(question.key == "service" and question.required
                            for question in ambiguous_draft.questions))

    def test_invoice_contributes_service_but_never_supplies_a_reading_column(self):
        readings = make_readings_workbook(
            self.project_root / "neutral-readings.xlsx", ("Vivienda", "Consumo")
        )
        invoice = make_pdf(
            self.project_root / "heating-invoice.pdf",
            "FACTURA CALEFACCION LECTURA CONTADOR",
        )

        draft = community_onboarding.analyse_sources(
            community_code="900", community_name="Comunidad prueba",
            owner_list_path=self.owners_csv, reading_paths=(readings,),
            invoice_paths=(invoice,), project_root=self.project_root,
        )

        self.assertEqual(("CALEFACCION",), draft.detected_modules)
        self.assertEqual("invoice_pdf", draft.sources[-1].kind)
        self.assertTrue(any(
            question.key == "reading_column" and question.required
            for question in draft.questions
        ))

    def test_absent_service_and_reading_column_require_explicit_decisions(self):
        readings = make_readings_workbook(
            self.project_root / "undetected-readings.xlsx", ("Vivienda", "Consumo")
        )

        draft = community_onboarding.analyse_sources(
            community_code="900", community_name="Comunidad prueba",
            owner_list_path=self.owners_csv, reading_paths=(readings,),
            invoice_paths=(), project_root=self.project_root,
        )

        questions = {question.key: question for question in draft.questions}
        self.assertTrue(questions["service"].required)
        self.assertTrue(questions["reading_column"].required)
        with self.assertRaisesRegex(ValueError, "servicio"):
            community_onboarding.build_profile_payload(
                draft, answers={"reading_column": "Consumo"}
            )
        payload = community_onboarding.build_profile_payload(
            draft,
            answers={"service": "NO_APLICA", "reading_column": "Consumo"},
        )
        self.assertEqual([], payload["active_modules"])

    def test_service_and_column_answers_are_the_single_payload_configuration(self):
        readings = make_readings_workbook(
            self.project_root / "ambiguous-services.xlsx",
            ("Vivienda", "Lectura ACS", "Lectura Calefacción"),
        )
        draft = community_onboarding.analyse_sources(
            community_code="900", community_name="Comunidad prueba",
            owner_list_path=self.owners_csv, reading_paths=(readings,),
            invoice_paths=(), project_root=self.project_root,
        )

        acs = community_onboarding.build_profile_payload(
            draft, answers={"service": "ACS", "reading_column": "Lectura ACS"}
        )
        heating = community_onboarding.build_profile_payload(
            draft,
            answers={
                "service": "CALEFACCION",
                "reading_column": "Lectura Calefacción",
            },
        )

        self.assertEqual(["ACS"], acs["active_modules"])
        self.assertEqual(["acs_fixed", "acs_variable"], [
            concept["key"] for concept in acs["concepts"]
        ])
        self.assertEqual("Lectura ACS", acs["onboarding_configuration"]["reading_column"])
        self.assertEqual(["CALEFACCION"], heating["active_modules"])
        self.assertEqual(["heating_fixed", "heating_variable"], [
            concept["key"] for concept in heating["concepts"]
        ])
        self.assertEqual(
            "Lectura Calefacción",
            heating["onboarding_configuration"]["reading_bindings"][0]["column"],
        )

    def test_a_detected_service_can_be_corrected_before_building_the_profile(self):
        readings = make_readings_workbook(
            self.project_root / "detected-acs.xlsx", ("Vivienda", "Lectura ACS")
        )
        draft = community_onboarding.analyse_sources(
            community_code="900", community_name="Comunidad prueba",
            owner_list_path=self.owners_csv, reading_paths=(readings,),
            invoice_paths=(), project_root=self.project_root,
        )

        payload = community_onboarding.build_profile_payload(
            draft, answers={"service": "CALEFACCION"}
        )

        self.assertEqual(["CALEFACCION"], payload["active_modules"])
        self.assertEqual(
            ["heating_fixed", "heating_variable"],
            [concept["key"] for concept in payload["concepts"]],
        )

    def test_profile_payload_requires_an_answer_for_ambiguous_reading_column(self):
        with self.assertRaisesRegex(ValueError, "columna de lectura"):
            community_onboarding.build_profile_payload(self.ambiguous_draft, answers={})

    def test_profile_payload_rejects_empty_false_or_unknown_required_answers(self):
        for answer in ("", False, "Otra columna"):
            with self.subTest(answer=answer):
                with self.assertRaisesRegex(ValueError, "columna de lectura"):
                    community_onboarding.build_profile_payload(
                        self.ambiguous_draft, answers={"reading_column": answer}
                    )

    def test_profile_payload_contains_only_confirmed_modules(self):
        payload = community_onboarding.build_profile_payload(
            self.module_draft,
            answers={"service": "ACS"},
        )
        self.assertIn("ACS", payload["active_modules"])
        self.assertNotIn("CALEFACCION", payload["active_modules"])

    def test_profile_payload_accepts_safe_community_codes(self):
        draft = community_onboarding.OnboardingDraft(
            community_code="644", community_name="Comunidad prueba", sources=(),
            detected_modules=(), questions=(),
        )

        self.assertEqual("644_v1", community_onboarding.build_profile_payload(
            draft, answers={"service": "NO_APLICA"}
        )["key"])

    def test_profile_payload_rejects_unsafe_windows_community_codes(self):
        for community_code in (
            "900/archivo", "900\\archivo", "900:", "900*", "900?", "900<",
            "900>", "900|", '900"', "900.", "900 ", "CON", "PRN", "AUX",
            "NUL", "COM1", "COM9", "LPT1", "LPT9", "CON.txt", "COM1.csv",
            "LPT9.foo", "900\x00", "900\x1f",
        ):
            draft = community_onboarding.OnboardingDraft(
                community_code=community_code, community_name="Comunidad prueba",
                sources=(), detected_modules=(), questions=(),
            )
            with self.subTest(community_code=community_code):
                with self.assertRaisesRegex(ValueError, "código de comunidad"):
                    community_onboarding.build_profile_payload(
                        draft, answers={"service": "NO_APLICA"}
                    )

    def _exercise_meter_service_flow(self, service: str, code: str) -> tuple[str, ...]:
        reading_label = "Lectura ACS" if service == "ACS" else "Lectura Calefacción"
        readings = make_readings_workbook(
            self.project_root / f"{code}-readings.xlsx", ("Vivienda", reading_label)
        )
        draft = community_onboarding.analyse_sources(
            community_code=code, community_name=f"Comunidad {service}",
            owner_list_path=self.owners_csv, reading_paths=(readings,),
            invoice_paths=(), project_root=self.project_root,
        )
        result = community_onboarding.confirm_onboarding(
            self.connection,
            draft=draft,
            answers={"service": service},
            project_root=self.project_root,
            archive_root=self.archive_root,
            period_name="2026",
            start_date=date(2026, 1, 1),
            end_date=date(2026, 12, 31),
            actor="Prueba",
        )
        owner = self.connection.execute(
            """INSERT INTO propietarios
               (id_comunidad,codigo_vivienda,nombre_propietario,coeficiente,activo)
               VALUES (?,?,?,?,1)""",
            (result.community_id, "A-01", "VECINO SINTETICO", 100),
        )
        self.connection.executemany(
            """INSERT INTO lecturas_vecino
               (id_propietario,id_periodo,tipo,fecha_lectura,valor_acumulado,estado)
               VALUES (?,?,?,?,?,'real')""",
            (
                (owner.lastrowid, result.period_id, service, "2026-01-01", 100),
                (owner.lastrowid, result.period_id, service, "2026-12-31", 125),
            ),
        )
        prefix = "acs" if service == "ACS" else "heating"
        persisted = json.loads(result.profile_path.read_text(encoding="utf-8"))
        self.assertEqual(
            {
                f"{prefix}_fixed_actual", f"{prefix}_fixed_billed",
                f"{prefix}_variable_actual", f"{prefix}_variable_billed",
            },
            set(persisted["workbook_layout"]["parameter_cells"]),
        )
        for key, value in (
            (f"{prefix}_fixed_actual", 20),
            (f"{prefix}_fixed_billed", 10),
            (f"{prefix}_variable_actual", 80),
            (f"{prefix}_variable_billed", 40),
        ):
            self.connection.execute(
                """INSERT INTO period_parameters
                   (id_comunidad,id_periodo,parameter_key,numeric_value,unit)
                   VALUES (?,?,?,?, 'EUR')""",
                (result.community_id, result.period_id, key, value),
            )
        self.connection.commit()
        document_review.validate_case_ready(self.connection, result.case_id)

        exported = generate_official_excel(
            self.connection,
            id_case=result.case_id,
            project_root=self.project_root,
            output_root=self.project_root / "salidas",
            recalculator=DeterministicRecalculator(),
        )
        self.assertTrue(exported.output_path.is_file())
        distributed = calculate_case_distribution(
            self.connection, id_case=result.case_id, project_root=self.project_root
        )
        return tuple(distributed.concept_totals_cents)

    def test_acs_onboarding_profile_reaches_canonical_export_and_distribution(self):
        self.assertEqual(
            ("acs_fixed", "acs_variable"),
            self._exercise_meter_service_flow("ACS", "901"),
        )

    def test_heating_onboarding_profile_reaches_canonical_export_and_distribution(self):
        self.assertEqual(
            ("heating_fixed", "heating_variable"),
            self._exercise_meter_service_flow("CALEFACCION", "902"),
        )

    def test_confirm_onboarding_writes_profile_template_archives_sources_and_creates_case(self):
        result = community_onboarding.confirm_onboarding(
            self.connection, **self._confirmation_arguments()
        )

        profile_path = self.project_root / "config/excel_profiles/runtime/900_v1.json"
        template_path = self.project_root / "plantillas/comunidades/900/900_v1.xlsx"
        self.assertIsNotNone(result.case_id)
        self.assertTrue(profile_path.is_file())
        self.assertTrue(template_path.is_file())
        self.assertEqual(3, self._registered_source_count())
        payload = json.loads(profile_path.read_text(encoding="utf-8"))
        workbook = load_workbook(template_path, read_only=True)
        try:
            self.assertEqual(
                ["DATOS", "LECTURAS ACS M3", "ANALISIS"], workbook.sheetnames
            )
        finally:
            workbook.close()
        self.assertEqual("LECTURAS ACS M3", payload["workbook_layout"]["tables"]["ACS"]["sheet"])
        self.assertEqual("900_v1", excel_profiles.load_profile(
            "900_v1", self.project_root
        ).key)
        self.assertEqual(
            "Lectura",
            excel_profiles.load_profile(
                "900_v1", self.project_root
            ).onboarding_configuration["reading_column"],
        )
        self.assertEqual("900_v1", case_workflow_actions.resolve_case_profile(
            self.connection,
            id_case=result.case_id,
            active_community_id=result.community_id,
            project_root=self.project_root,
        ).key)

    def test_onboarded_runtime_profile_installs_bootstrap_template_for_long_code(self):
        draft = community_onboarding.OnboardingDraft(
            community_code="1234",
            community_name="Comunidad prueba larga",
            sources=self.valid_draft.sources,
            detected_modules=self.valid_draft.detected_modules,
            questions=self.valid_draft.questions,
        )
        result = community_onboarding.confirm_onboarding(
            self.connection,
            **{**self._confirmation_arguments(), "draft": draft},
        )
        template = self.project_root / "plantillas/comunidades/1234/1234_v1.xlsx"

        bootstrap, companions = case_workflow_actions.run_bootstrap_import(
            self.database,
            id_case=result.case_id,
            active_community_id=result.community_id,
            project_root=self.project_root,
            master_path=template,
            actor="Prueba",
        )

        self.assertIsNone(companions)
        self.assertEqual(template, bootstrap.installed_template_path)
        self.assertTrue(template.is_file())

    def test_confirm_onboarding_rolls_back_json_database_template_and_archives_after_error(self):
        real_register = expedient_service.register_source_document
        calls = 0
        sentinel_path = None

        def fail_after_two_real_registrations(*args, **kwargs):
            nonlocal calls, sentinel_path
            calls += 1
            if calls == 1:
                case_id = args[1]
                sentinel_path = self.archive_root / str(case_id) / "fuentes" / "ajeno.txt"
                sentinel_path.parent.mkdir(parents=True, exist_ok=True)
                sentinel_path.write_bytes(b"no pertenece al alta")
            result = real_register(*args, **kwargs)
            if calls == 3:
                raise OSError("fallo sintético tardío")
            return result

        with patch(
            "community_onboarding.expedient_service.register_source_document",
            side_effect=fail_after_two_real_registrations,
        ):
            with self.assertRaisesRegex(OSError, "sintético tardío"):
                community_onboarding.confirm_onboarding(
                    self.connection, **self._confirmation_arguments()
                )

        self.assertFalse((self.project_root / "config/excel_profiles/runtime/900_v1.json").exists())
        self.assertFalse((self.project_root / "plantillas/comunidades/900/900_v1.xlsx").exists())
        self.assertFalse((self.project_root / "config/excel_profiles").exists())
        self.assertFalse((self.project_root / "plantillas/comunidades/900").exists())
        self.assertEqual(0, self._community_count())
        self.assertEqual(0, self._registered_source_count())
        self.assertIsNotNone(sentinel_path)
        self.assertEqual(b"no pertenece al alta", sentinel_path.read_bytes())
        self.assertEqual(
            [sentinel_path],
            [path for path in self.archive_root.rglob("*") if path.is_file()],
        )

    def test_confirm_onboarding_never_replaces_a_concurrent_template(self):
        template_path = self.project_root / "plantillas/comunidades/900/900_v1.xlsx"
        sentinel = b"plantilla concurrente"
        real_link = os.link

        def publish_after_concurrent_create(source, destination):
            destination = Path(destination)
            if destination == template_path:
                destination.write_bytes(sentinel)
            return real_link(source, destination)

        with patch("community_onboarding.os.link", side_effect=publish_after_concurrent_create):
            with self.assertRaises(FileExistsError):
                community_onboarding.confirm_onboarding(
                    self.connection, **self._confirmation_arguments()
                )

        self.assertEqual(sentinel, template_path.read_bytes())
        self.assertFalse((self.project_root / "config/excel_profiles/runtime/900_v1.json").exists())
        self.assertEqual(0, self._community_count())

    def test_confirm_onboarding_never_replaces_a_concurrent_profile(self):
        profile_path = self.project_root / "config/excel_profiles/runtime/900_v1.json"
        template_path = self.project_root / "plantillas/comunidades/900/900_v1.xlsx"
        sentinel = b'{"propietario":"ajeno"}'
        real_link = os.link
        calls = 0

        def publish_after_concurrent_create(source, destination):
            nonlocal calls
            calls += 1
            destination = Path(destination)
            if calls == 2:
                self.assertEqual(profile_path, destination)
                destination.write_bytes(sentinel)
            return real_link(source, destination)

        with patch("community_onboarding.os.link", side_effect=publish_after_concurrent_create):
            with self.assertRaises(FileExistsError):
                community_onboarding.confirm_onboarding(
                    self.connection, **self._confirmation_arguments()
                )

        self.assertEqual(sentinel, profile_path.read_bytes())
        self.assertFalse(template_path.exists())
        self.assertEqual(0, self._community_count())

    def test_confirm_onboarding_rejects_partial_period_dates_without_side_effects(self):
        arguments = self._confirmation_arguments()
        arguments["end_date"] = None

        with self.assertRaisesRegex(ValueError, "fechas.*completas"):
            community_onboarding.confirm_onboarding(self.connection, **arguments)

        self.assertEqual(0, self._community_count())
        self.assertFalse((self.project_root / "config/excel_profiles/runtime/900_v1.json").exists())

    def test_confirm_onboarding_cleans_the_actual_source_when_it_changes_during_registration(self):
        real_register = expedient_service.register_source_document
        changed_source = self.owners_csv

        def change_then_fail(*args, **kwargs):
            changed_source.write_bytes(b"fuente cambiada durante el registro")
            result = real_register(*args, **kwargs)
            raise OSError("fallo posterior al registro cambiado")

        with patch(
            "community_onboarding.expedient_service.register_source_document",
            side_effect=change_then_fail,
        ):
            with self.assertRaisesRegex(OSError, "registro cambiado"):
                community_onboarding.confirm_onboarding(
                    self.connection, **self._confirmation_arguments()
                )

        self.assertEqual(0, self._registered_source_count())
        self.assertEqual(
            [], [path for path in self.archive_root.rglob("*") if path.is_file()]
            if self.archive_root.exists() else [],
        )
        self.assertEqual(0, self._community_count())


if __name__ == "__main__":
    unittest.main()
