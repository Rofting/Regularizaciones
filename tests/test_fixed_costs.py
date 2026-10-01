import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from decimal import Decimal
from io import StringIO
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
CORE_DIR = PROJECT_ROOT / "core"
if str(CORE_DIR) not in sys.path:
    sys.path.insert(0, str(CORE_DIR))

from openpyxl import Workbook, load_workbook

import fixed_costs
import gestor_bd


class FixedCostsTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        with redirect_stdout(StringIO()):
            gestor_bd.crear_bd(str(Path(self.directory.name) / "g.db"))
        self.connection = gestor_bd.conectar(str(Path(self.directory.name) / "g.db"))
        self.community = gestor_bd.obtener_o_crear_comunidad(self.connection, "1", "C")
        self.period = self.connection.execute(
            "INSERT INTO periodos(id_comunidad,nombre,fecha_inicio,fecha_fin) VALUES (?,?,?,?)",
            (self.community, "2026", "2026-01-01", "2026-12-31"),
        ).lastrowid

    def tearDown(self):
        self.connection.close()
        self.directory.cleanup()

    def test_save_load_and_clear(self):
        fixed_costs.save_fixed_costs(self.connection, self.community, self.period, {
            "fixed_cost_boiler_maintenance": "1.095,50", "fixed_cost_meter_reading_acs": "82",
        })
        self.assertEqual(
            {"fixed_cost_boiler_maintenance": Decimal("1095.5"), "fixed_cost_meter_reading_acs": Decimal("82")},
            fixed_costs.load_fixed_costs(self.connection, self.community, self.period),
        )
        fixed_costs.save_fixed_costs(self.connection, self.community, self.period,
                                     {"fixed_cost_meter_reading_acs": ""})
        self.assertEqual(Decimal(0), fixed_costs.load_fixed_costs(
            self.connection, self.community, self.period)["fixed_cost_meter_reading_acs"])
        with self.assertRaises(ValueError):
            fixed_costs.save_fixed_costs(self.connection, self.community, self.period,
                                         {"fixed_cost_boiler_maintenance": "-3"})

    def test_values_are_written_only_on_matching_rows(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "OTROS GASTOS"
        sheet["A9"] = "MANTENIMIENTO SALA CALDERAS"
        sheet["A6"] = "Otra cosa distinta"
        written = fixed_costs.write_fixed_costs(
            workbook,
            {"fixed_cost_boiler_maintenance": Decimal("95.5"), "fixed_cost_meter_reading_acs": Decimal("80")},
            lambda cell, value: setattr(cell, "value", value),
        )
        self.assertEqual(["B9", "C9"], written)
        self.assertEqual(95.5, sheet["B9"].value)
        self.assertEqual("=B9*12", sheet["C9"].value)
        self.assertIsNone(sheet["B6"].value)

    def test_common_model_carries_no_community_costs(self):
        sheet = load_workbook(PROJECT_ROOT / "plantillas" / "modelo" / "modelo_acs_v1.xlsx")["OTROS GASTOS"]
        for item in fixed_costs.FIXED_COSTS:
            with self.subTest(cell=item.cell):
                self.assertIn(sheet[item.cell].value, (None, ""))
                self.assertTrue(str(sheet[f"A{sheet[item.cell].row}"].value).upper().startswith(item.sheet_label))
        self.assertEqual("=B6*12", sheet["C6"].value)

    def test_inherited_costs_require_confirmation_only_on_recognized_rows(self):
        path = Path(self.directory.name) / "antigua.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = fixed_costs.SHEET
        sheet["A6"], sheet["B6"] = "LECTURAS CONTADORES ACS", "993,44"
        sheet["A7"], sheet["B7"] = "LECTURAS CONTADORES CALEF", "=80+2"
        sheet["A8"], sheet["B8"] = "MANTENIMIENTO PLACAS SOLARES", 0
        sheet["A9"], sheet["B9"] = "Fila diferente", 150
        workbook.save(path)
        workbook.close()
        before = path.read_bytes()
        findings = fixed_costs.unconfirmed_template_costs(path, {})
        self.assertEqual(["B6", "B7"], [f.cost.cell for f in findings])
        self.assertIn("993,44", fixed_costs.inherited_cost_message(findings))
        self.assertEqual((findings[1],), fixed_costs.unconfirmed_template_costs(
            path, {"fixed_cost_meter_reading_acs": Decimal(0)}))
        self.assertEqual(before, path.read_bytes())

    def test_invalid_form_does_not_save_partial_values(self):
        with self.assertRaises(ValueError):
            fixed_costs.save_fixed_costs(self.connection, self.community, self.period, {
                "fixed_cost_meter_reading_acs": "80", "fixed_cost_boiler_maintenance": "-1",
            })
        self.assertEqual({}, fixed_costs.load_fixed_costs(self.connection, self.community, self.period))

    def test_hardcoded_annual_cost_is_detected_even_if_monthly_cell_is_empty(self):
        path = Path(self.directory.name) / "anual-antiguo.xlsx"
        workbook = Workbook()
        workbook.active.title = fixed_costs.SHEET
        workbook.active["A6"] = "LECTURAS CONTADORES ACS"
        workbook.active["C6"] = 993.44
        workbook.save(path)
        workbook.close()
        findings = fixed_costs.unconfirmed_template_costs(path, {})
        self.assertEqual(["C6"], [item.cell for item in findings])
        self.assertIn("993,44 €/año", fixed_costs.inherited_cost_message(findings))

    def test_fixed_costs_cannot_be_saved_for_another_communitys_period(self):
        other = gestor_bd.obtener_o_crear_comunidad(self.connection, "2", "Otra")
        with self.assertRaisesRegex(ValueError, "no corresponde"):
            fixed_costs.save_fixed_costs(self.connection, other, self.period, {
                "fixed_cost_meter_reading_acs": 0,
            })


if __name__ == "__main__":
    unittest.main()
