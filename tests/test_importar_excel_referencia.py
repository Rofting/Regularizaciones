import sqlite3
import subprocess
import sys
import tempfile
import unittest
from contextlib import closing, redirect_stdout
from io import StringIO
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
CORE_DIR = PROJECT_ROOT / "core"
if str(CORE_DIR) not in sys.path:
    sys.path.insert(0, str(CORE_DIR))

import gestor_bd
from importar_excel_referencia import (
    ReferenceValidationError,
    import_reference_workbook,
    parse_reference_workbook,
)
from tests.helpers import make_reference_workbook


class ReferenceWorkbookParserTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.path = Path(self.directory.name) / "referencia.xlsx"
        make_reference_workbook(
            self.path,
            period="2024-2025",
            actual_cents=100_000,
            billed_cents=90_000,
        )

    def tearDown(self):
        self.directory.cleanup()

    def test_parses_period_concepts_and_source_addresses_by_labels(self):
        reference = parse_reference_workbook(self.path)
        concepts = {concept.concept_key: concept for concept in reference.concepts}

        self.assertEqual("COMUNIDAD DE PRUEBA", reference.community_name)
        self.assertEqual("2024-2025", reference.period_name)
        self.assertEqual("2024-08-01", reference.start_date.isoformat())
        self.assertEqual("2025-07-31", reference.end_date.isoformat())
        self.assertEqual(2, reference.property_count)
        self.assertEqual(30_000, concepts["acs_fixed"].actual_cents)
        self.assertEqual(36_000, concepts["acs_fixed"].billed_cents)
        self.assertEqual(70_000, concepts["acs_variable"].actual_cents)
        self.assertEqual(54_000, concepts["acs_variable"].billed_cents)
        fixed_cells = {source.cell_address for source in concepts["acs_fixed"].source_values}
        self.assertTrue({"H64", "H66", "H67", "F70", "G70"}.issubset(fixed_cells))

    def test_rejects_missing_required_sheet(self):
        workbook = load_workbook(self.path)
        del workbook["ANALISIS"]
        workbook.save(self.path)

        with self.assertRaisesRegex(ReferenceValidationError, "ANALISIS"):
            parse_reference_workbook(self.path)

    def test_rejects_duplicate_anchor(self):
        workbook = load_workbook(self.path)
        workbook["ANALISIS"]["F68"] = "IMPORTE COBRADO"
        workbook.save(self.path)

        with self.assertRaisesRegex(ReferenceValidationError, "duplicada"):
            parse_reference_workbook(self.path)

    def test_rejects_incoherent_difference(self):
        workbook = load_workbook(self.path)
        workbook["ANALISIS"]["H67"] = 99
        workbook.save(self.path)

        with self.assertRaisesRegex(ReferenceValidationError, "diferencia"):
            parse_reference_workbook(self.path)

    def test_import_is_idempotent_by_file_hash(self):
        database_path = Path(self.directory.name) / "gestion.db"
        with redirect_stdout(StringIO()):
            gestor_bd.crear_bd(str(database_path))

        with closing(gestor_bd.conectar(str(database_path))) as connection:
            community_id = gestor_bd.obtener_o_crear_comunidad(
                connection, "TEST", "COMUNIDAD DE PRUEBA"
            )
            first = import_reference_workbook(connection, community_id, self.path)
            second = import_reference_workbook(connection, community_id, self.path)
            batch_count = connection.execute(
                "SELECT COUNT(*) FROM import_batches"
            ).fetchone()[0]
            source_count = connection.execute(
                "SELECT COUNT(*) FROM source_values"
            ).fetchone()[0]
            period_count = connection.execute("SELECT COUNT(*) FROM periodos").fetchone()[0]

        self.assertEqual(first, second)
        self.assertEqual(1, batch_count)
        self.assertGreater(source_count, 8)
        self.assertEqual(1, period_count)

    def test_module_cli_inspects_workbook_without_personal_data(self):
        result = subprocess.run(
            [
                sys.executable,
                "-m",
                "core.importar_excel_referencia",
                "--inspect",
                str(self.path),
            ],
            cwd=PROJECT_ROOT,
            capture_output=True,
            text=True,
            encoding="utf-8",
        )

        self.assertEqual(0, result.returncode, result.stderr)
        self.assertIn("periodo=2024-2025", result.stdout)
        self.assertIn("propiedades=2", result.stdout)
        self.assertNotIn("VECINO", result.stdout)


if __name__ == "__main__":
    unittest.main()
