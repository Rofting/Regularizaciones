"""Regressions for the complete analysed-source producer/consumer workflow."""
import json
import sqlite3
import unittest
from datetime import date
from types import SimpleNamespace
from unittest.mock import patch

from tests import test_expedient_flow
from tests.test_expedient_ui import FakeWidget
import case_ingestion
import case_distribution
import case_letter_service
import database_reset
import db_migrations
import document_review
import excel_export_service
import expedient_service
import expedient_ui
import gestor_bd
import source_analysis
from excel_profiles import ConceptRule
from source_analysis import SourceAnalysis, SourceLocator


class UnifiedIngestionRegressionTest(unittest.TestCase):
    setUp = test_expedient_flow.ExpedientFlowTest.setUp
    tearDown = test_expedient_flow.ExpedientFlowTest.tearDown
    add_confirmed_reading = test_expedient_flow.ExpedientFlowTest.add_confirmed_reading

    def ingest(self, analysis, path=None, case=None):
        self.connection.commit()
        return case_ingestion.add_analysed_document_to_case(
            self.connection, (case or self.case).id_case,
            source_path=path or self.source_path, archive_root=self.archive_root,
            analysis=analysis,
        ).document

    def confirm(self, document, case=None):
        return case_ingestion.confirm_source_candidates(
            self.connection, (case or self.case).id_case, document.id_document,
            confirmed_by="Jose",
        )

    def invoice_analysis(self):
        return source_analysis.analyse_pdf(self.source_path, pdf_processor=lambda *_: {
            "ok": True, "tipo": "FACTURA", "datos": {
                "tipo_suministro": "GAS", "fecha_inicio": "2026-01-01",
                "fecha_fin": "2026-01-31", "importe_total": 128.10,
                "termino_fijo": 28.10, "termino_variable": 100.0,
            },
        })

    def test_clean_analysed_invoice_is_applied_automatically(self):
        self.ingest(self.invoice_analysis())
        self.assertEqual("ready_for_calculation", document_review.validate_case_ready(
            self.connection, self.case.id_case,
        ).status)
        self.assertEqual(1, self.connection.execute("SELECT COUNT(*) FROM facturas").fetchone()[0])

    def test_confirmation_applies_invoice_and_export_input_changes_after_correction(self):
        document = self.ingest(self.invoice_analysis())
        self.assertTrue(callable(getattr(case_ingestion, "confirm_source_candidates", None)),
                        "Production ingestion needs a reachable confirmation operation")
        self.confirm(document)
        profile = SimpleNamespace(key="test", version="1", source_sha256="test",
                                  active_modules=(), concepts=(), onboarding_configuration={})
        context = excel_export_service._input_case_context(self.connection, self.case.id_case)
        before = excel_export_service._input_hash(self.connection, context, profile)
        self.assertEqual(128.1, self.connection.execute("SELECT importe_total FROM facturas").fetchone()[0])
        issue = document_review.create_review_issue(
            self.connection, self.case.id_case, document.id_document,
            code="CHECK", field_name="importe_total", message="Comprobar total",
        )
        document_review.resolve_issue(self.connection, issue.id_issue, value="129.10", reason="Original")
        self.assertEqual(129.1, self.connection.execute("SELECT importe_total FROM facturas").fetchone()[0])
        self.assertNotEqual(before, excel_export_service._input_hash(self.connection, context, profile))
        self.assertEqual("ready_for_calculation", document_review.validate_case_ready(
            self.connection, self.case.id_case).status)

    def test_reanalysis_with_new_candidates_blocks_export_until_confirmation(self):
        document = self.ingest(self.invoice_analysis())
        self.confirm(document)
        document_review.validate_case_ready(self.connection, self.case.id_case)
        candidates = dict(self.invoice_analysis().candidates)
        candidates["proveedor"] = "Proveedor detectado"
        case_ingestion.reanalyze_case_documents(self.connection, self.case.id_case,
            analyser=lambda _: SourceAnalysis.invoice(candidates))
        with self.assertRaisesRegex(excel_export_service.ExportBlockedError, "confirm|valid|revis"):
            excel_export_service._case_context(self.connection, self.case.id_case)

    def test_reanalysis_of_manual_reading_classification_recreates_required_issues(self):
        document = self.ingest(SourceAnalysis.unknown())
        classification_issue = document_review.list_open_issues(
            self.connection, self.case.id_case,
        )[0]
        document_review.resolve_issue(
            self.connection, classification_issue.id_issue,
            value="reading", reason="Clasificación comprobada manualmente",
        )
        document_review.record_candidates(
            self.connection, document.id_document, {"tipo": "ACS"},
            source="manual", validation_status="validated",
        )

        case_ingestion.reanalyze_case_documents(
            self.connection, self.case.id_case,
            analyser=lambda _: SourceAnalysis.unknown(),
        )

        issues = document_review.list_open_issues(self.connection, self.case.id_case)
        self.assertEqual(
            {"fecha_inicio", "fecha_fin", "vecinos"},
            {issue.field_name for issue in issues},
        )
        stored = self.connection.execute(
            """SELECT value, source, validation_status FROM extraction_candidates
               WHERE id_document=? AND field_name='tipo'""",
            (document.id_document,),
        ).fetchone()
        self.assertEqual(("ACS", "manual", "validated"), tuple(stored))

    def test_reanalysis_of_invoice_candidates_misclassified_as_reading_requests_one_reclassification(self):
        document = self.ingest(SourceAnalysis.unknown())
        issue = document_review.list_open_issues(self.connection, self.case.id_case)[0]
        document_review.resolve_issue(
            self.connection, issue.id_issue,
            value="reading", reason="Clasificación inicial",
        )

        case_ingestion.reanalyze_case_documents(
            self.connection, self.case.id_case,
            analyser=lambda _: SourceAnalysis.invoice({
                "tipo_suministro": "AGUA", "fecha_inicio": "2026-01-01",
                "fecha_fin": "2026-01-31", "importe_total": "177.84",
            }, confidence="medium"),
        )

        issues = document_review.list_open_issues(self.connection, self.case.id_case)
        self.assertEqual(1, len(issues))
        self.assertEqual("DOCUMENT_CLASSIFICATION_REQUIRED", issues[0].code)

    def test_resolving_classification_updates_effective_document_kind_immediately(self):
        document = self.ingest(SourceAnalysis.unknown())
        issue = document_review.list_open_issues(self.connection, self.case.id_case)[0]

        document_review.resolve_issue(
            self.connection, issue.id_issue,
            value="invoice", reason="Contenido de factura comprobado",
        )

        stored_kind = self.connection.execute(
            "SELECT document_kind FROM source_documents WHERE id_document=?",
            (document.id_document,),
        ).fetchone()[0]
        self.assertEqual("invoice", stored_kind)

    def test_invoice_candidates_misclassified_as_reading_request_reclassification(self):
        document = case_ingestion.add_document_to_case(
            self.connection, self.case.id_case,
            source_path=self.source_path, archive_root=self.archive_root,
            document_kind="reading",
            candidates={
                "tipo_suministro": "AGUA", "fecha_inicio": "2026-01-01",
                "fecha_fin": "2026-01-31", "importe_total": "177.84",
            },
            required_fields=(),
        ).document
        self.connection.execute(
            """UPDATE source_documents
               SET classification_confidence='low', status='under_review'
               WHERE id_document=?""",
            (document.id_document,),
        )

        case_ingestion.ensure_pending_source_issues(self.connection, self.case.id_case)

        issues = document_review.list_open_issues(self.connection, self.case.id_case)
        self.assertEqual(1, len(issues))
        self.assertEqual("DOCUMENT_CLASSIFICATION_REQUIRED", issues[0].code)
        self.assertEqual("document_kind", issues[0].field_name)
        document_review.resolve_issue(
            self.connection, issues[0].id_issue,
            value="invoice", reason="El contenido corresponde a una factura",
        )
        confirmed = case_ingestion.confirm_source_candidates(
            self.connection, self.case.id_case, document.id_document,
            confirmed_by="Jose",
        )
        self.assertEqual("validated", confirmed.status)
        self.assertEqual(1, self.connection.execute("SELECT COUNT(*) FROM facturas").fetchone()[0])

    def test_pending_reading_with_invalid_type_becomes_actionable(self):
        rows = [{
            "vivienda": "A", "fecha_ant": "2026-01-01", "val_ant": 0,
            "fecha_act": "2026-01-31", "val_act": 0,
        }]
        document = case_ingestion.add_document_to_case(
            self.connection, self.case.id_case,
            source_path=self.reading_file, archive_root=self.archive_root,
            document_kind="reading",
            candidates={"tipo": "12123", "vecinos": json.dumps(rows)},
            required_fields=(),
        ).document
        self.connection.execute(
            """UPDATE source_documents
               SET classification_confidence='low', status='under_review'
               WHERE id_document=?""",
            (document.id_document,),
        )

        ensure = getattr(case_ingestion, "ensure_pending_source_issues", None)
        self.assertTrue(callable(ensure), "Pending sources need an actionable review operation")
        ensure(self.connection, self.case.id_case)

        issues = document_review.list_open_issues(self.connection, self.case.id_case)
        self.assertEqual(["tipo"], [issue.field_name for issue in issues])
        with self.assertRaisesRegex(ValueError, "ACS|CALEFACCION"):
            document_review.resolve_issue(
                self.connection, issues[0].id_issue,
                value="12123", reason="Valor copiado del documento",
            )

    def test_pending_source_with_invalid_structured_data_becomes_actionable(self):
        document = case_ingestion.add_document_to_case(
            self.connection, self.case.id_case,
            source_path=self.reading_file, archive_root=self.archive_root,
            document_kind="reading",
            candidates={"vecinos": "{contenido dañado"},
            required_fields=(),
        ).document
        self.connection.execute(
            """UPDATE source_documents
               SET classification_confidence='medium', status='under_review'
               WHERE id_document=?""",
            (document.id_document,),
        )

        case_ingestion.ensure_pending_source_issues(self.connection, self.case.id_case)

        issues = document_review.list_open_issues(self.connection, self.case.id_case)
        self.assertEqual(["vecinos"], [issue.field_name for issue in issues])
        self.assertIn("formato", issues[0].message.lower())

    def test_pending_invoice_with_invalid_amount_becomes_actionable(self):
        document = case_ingestion.add_document_to_case(
            self.connection, self.case.id_case,
            source_path=self.source_path, archive_root=self.archive_root,
            document_kind="invoice",
            candidates={
                "tipo_suministro": "GAS", "fecha_inicio": "2026-01-01",
                "fecha_fin": "2026-01-31", "importe_total": "no-numérico",
            },
            required_fields=(),
        ).document
        self.connection.execute(
            """UPDATE source_documents
               SET classification_confidence='medium', status='under_review'
               WHERE id_document=?""",
            (document.id_document,),
        )

        case_ingestion.ensure_pending_source_issues(self.connection, self.case.id_case)

        issues = document_review.list_open_issues(self.connection, self.case.id_case)
        self.assertEqual(["importe_total"], [issue.field_name for issue in issues])

    def test_invoice_confirmation_does_not_apply_rejected_required_dates(self):
        document = self.ingest(self.invoice_analysis())
        self.connection.execute("UPDATE extraction_candidates SET validation_status='rejected' WHERE id_document=? AND field_name='fecha_inicio'", (document.id_document,))
        with self.assertRaisesRegex(ValueError, "fecha_inicio"):
            self.confirm(document)
        # La factura limpia se publicó al analizarla; rechazar una fecha después
        # impide confirmarla de nuevo, pero no borra el dato ya aplicado.
        self.assertEqual(1, self.connection.execute("SELECT COUNT(*) FROM facturas").fetchone()[0])

    def test_ui_source_review_reaches_canonical_invoice(self):
        self.ingest(self.invoice_analysis())
        self.assertTrue(callable(getattr(expedient_ui, "open_confirm_sources_dialog", None)),
                        "Complete extracted candidates need a reachable UI confirmation")
        app = SimpleNamespace(ruta_bd_expedientes=self.database_path,
            _refrescar_lista_expedientes=lambda **_: None, _refrescar_expediente=lambda: None,
            log=lambda *_: None, after=lambda *_: None)
        panel = FakeWidget()
        with patch.object(expedient_ui, "_dialog", return_value=panel), \
            patch.multiple(expedient_ui.ctk, CTkFrame=FakeWidget, CTkScrollableFrame=FakeWidget,
                           CTkLabel=FakeWidget, CTkButton=FakeWidget), \
             patch.object(expedient_ui.messagebox, "showwarning"):
            expedient_ui.open_confirm_sources_dialog(app, self.case.id_case)
        self.assertEqual(128.1, self.connection.execute("SELECT importe_total FROM facturas").fetchone()[0])

    def test_confirm_sources_dialog_can_confirm_a_complete_medium_confidence_source(self):
        analysis = SourceAnalysis.invoice({
            "tipo_suministro": "GAS", "fecha_inicio": "2026-01-01",
            "fecha_fin": "2026-01-31", "importe_total": "128.10",
        }, confidence="medium")
        document = self.ingest(analysis)
        self.assertTrue(document_review.case_has_unapplied_sources(
            self.connection, self.case.id_case,
        ))
        app = SimpleNamespace(
            ruta_bd_expedientes=self.database_path,
            _refrescar_lista_expedientes=lambda **_: None,
            _refrescar_expediente=lambda: None,
            log=lambda *_: None,
            after=lambda *_: None,
        )
        dialog = FakeWidget()
        with patch.object(expedient_ui, "_dialog", return_value=dialog), \
            patch.object(expedient_ui.UIM, "fuente", return_value=None), \
            patch.multiple(
                expedient_ui.ctk,
                CTkFrame=FakeWidget, CTkScrollableFrame=FakeWidget,
                CTkLabel=FakeWidget, CTkButton=FakeWidget,
            ), \
            patch.object(expedient_ui.messagebox, "showinfo"), \
            patch.object(expedient_ui.messagebox, "showwarning"):
            expedient_ui.open_confirm_sources_dialog(app, self.case.id_case)
            button = next(
                widget for widget in dialog.descendants()
                if widget.options.get("text") == "Confirmar fuente"
            )
            button.options["command"]()

        self.assertFalse(document_review.case_has_unapplied_sources(
            self.connection, self.case.id_case,
        ))
        status = self.connection.execute(
            "SELECT status FROM source_documents WHERE id_document=?",
            (document.id_document,),
        ).fetchone()[0]
        self.assertEqual("validated", status)

    def test_pdf_structured_readings_survive_analysis_storage_and_application(self):
        self.connection.execute("INSERT INTO propietarios (id_comunidad,codigo_vivienda,nombre_propietario) VALUES (?,'A','Vecino')", (self.community_id,))
        rows = [{"vivienda": "A", "tipo": "ACS", "fecha_ant": "2026-01-01", "val_ant": 100,
                 "fecha_act": "2026-01-31", "val_act": 120}]
        analysis = source_analysis.analyse_pdf(self.source_path, pdf_processor=lambda *_: {
            "ok": True, "tipo": "LECTURA_METRIGEST", "datos": {"vecinos": rows}})
        self.assertEqual(rows, json.loads(analysis.candidates["vecinos"]))
        self.confirm(self.ingest(analysis))
        self.assertEqual([100.0, 120.0], [r[0] for r in self.connection.execute(
            "SELECT valor_acumulado FROM lecturas_vecino ORDER BY fecha_lectura")])

    def test_reading_csv_with_complete_rows_is_applied(self):
        self.connection.execute("INSERT INTO propietarios (id_comunidad,codigo_vivienda,nombre_propietario) VALUES (?,'A','Vecino')", (self.community_id,))
        self.reading_file.write_text("vivienda;tipo;fecha_ant;val_ant;fecha_act;val_act\nA;ACS;2026-01-01;100;2026-01-31;120\n", encoding="utf-8")
        analysis = source_analysis.analyse_tabular(self.reading_file)
        self.assertIn("vecinos", analysis.candidates)
        self.confirm(self.ingest(analysis, self.reading_file))
        self.assertEqual(2, self.connection.execute("SELECT COUNT(*) FROM lecturas_vecino").fetchone()[0])

    def test_incomplete_reading_csv_stays_reviewable(self):
        analysis = source_analysis.analyse_tabular(self.reading_file)
        self.assertTrue(analysis.kind == "unknown" or analysis.required_fields,
                        "Recognized headers without usable rows must require review")
        self.ingest(analysis, self.reading_file)
        self.assertTrue(document_review.list_open_issues(self.connection, self.case.id_case))

    def test_owner_csv_rows_can_be_confirmed_and_applied(self):
        self.reading_file.write_text("Fdenominacion;Nombre;Coeficiente;Email\nA;Vecino A;1,25;a@example.com\n", encoding="utf-8")
        analysis = source_analysis.analyse_tabular(self.reading_file)
        self.assertIn("propietarios", analysis.candidates)
        self.confirm(self.ingest(analysis, self.reading_file))
        owner = self.connection.execute("SELECT codigo_vivienda,coeficiente FROM propietarios").fetchone()
        self.assertEqual(("A", 1.25), tuple(owner))

    def test_legacy_204_generated_issues_are_reconciled_without_losing_manual_outcomes(self):
        for index in range(68):
            path = self.source_path.with_name(f"legacy-{index}.pdf")
            path.write_text(f"Legacy {index}")
            case_ingestion.add_document_to_case(
                self.connection, self.case.id_case, source_path=path, archive_root=self.archive_root,
                document_kind="invoice", candidates={},
                required_fields=("fecha_inicio", "fecha_fin", "importe_total"))
        document_id = self.connection.execute("SELECT id_document FROM source_documents LIMIT 1").fetchone()[0]
        manual = document_review.create_review_issue(
            self.connection, self.case.id_case, document_id, code="MISSING_REQUIRED_FIELD",
            field_name="manual_check", message="Revisar contrato firmado")
        resolved = document_review.list_open_issues(self.connection, self.case.id_case)[0]
        document_review.resolve_issue(self.connection, resolved.id_issue, value="2026-01-01", reason="Fecha verificada")
        self.connection.execute("DROP VIEW period_readings")
        self.connection.execute("DROP TABLE counter_reset_targets")
        self.connection.execute("DROP TABLE reading_periods")
        for column in ("source_context", "confirmed_by", "confirmed_at"):
            self.connection.execute(f"ALTER TABLE source_documents DROP COLUMN {column}")
        self.connection.execute("ALTER TABLE review_issues DROP COLUMN origin")
        self.connection.execute("DELETE FROM schema_migrations WHERE version>=7")
        self.connection.commit()
        db_migrations.migrate(self.connection)
        case_ingestion.reanalyze_case_documents(self.connection, self.case.id_case,
                                               analyser=lambda _: SourceAnalysis.reading())
        self.assertEqual([manual.id_issue], [i.id_issue for i in document_review.list_open_issues(self.connection, self.case.id_case)])
        self.assertEqual(1, self.connection.execute("SELECT COUNT(*) FROM manual_corrections").fetchone()[0])

    def test_migration_preserves_explicit_manual_issues_created_after_provenance_was_added(self):
        document = case_ingestion.add_document_to_case(self.connection, self.case.id_case,
            source_path=self.source_path, archive_root=self.archive_root,
            document_kind="invoice", candidates={}, required_fields=()).document
        manual = document_review.create_review_issue(self.connection, self.case.id_case, document.id_document,
            code="MISSING_REQUIRED_FIELD", field_name="importe_total",
            message="Falta el campo requerido: importe_total")
        self.connection.execute("UPDATE schema_migrations SET applied_at='2026-01-01 00:00:00' WHERE version=7")
        self.connection.execute("UPDATE review_issues SET created_at='2026-01-02 00:00:00' WHERE id_issue=?", (manual.id_issue,))
        self.connection.execute("DELETE FROM schema_migrations WHERE version=8")
        self.connection.commit()
        db_migrations.migrate(self.connection)
        case_ingestion.reanalyze_case_documents(self.connection, self.case.id_case,
                                               analyser=lambda _: SourceAnalysis.reading())
        self.assertEqual([manual.id_issue], [issue.id_issue for issue in document_review.list_open_issues(self.connection, self.case.id_case)])

    def test_shared_boundary_reading_is_available_to_both_periods_and_consumers(self):
        from openpyxl import Workbook
        self.connection.execute("INSERT INTO propietarios (id_comunidad,codigo_vivienda,nombre_propietario) VALUES (?,'A','Vecino')", (self.community_id,))
        header = "vivienda;tipo;fecha_ant;val_ant;fecha_act;val_act\n"
        self.reading_file.write_text(header + "A;ACS;2026-01-01;100;2026-01-31;120\n", encoding="utf-8")
        first = self.ingest(source_analysis.analyse_tabular(self.reading_file), self.reading_file)
        self.confirm(first)
        second_case = expedient_service.create_case(self.connection, self.community_id,
            name="Febrero", start_date=date(2026, 1, 31), end_date=date(2026, 2, 28))
        self.reading_file.write_text(header + "A;ACS;2026-01-31;120;2026-02-28;140\n", encoding="utf-8")
        second = self.ingest(source_analysis.analyse_tabular(self.reading_file), self.reading_file, second_case)
        self.confirm(second, second_case)
        owner_id = self.connection.execute("SELECT id_propietario FROM propietarios").fetchone()[0]
        concept = ConceptRule("acs_variable", "consumption", "unused", None, True)
        for case in (self.case, second_case):
            context = excel_export_service._input_case_context(self.connection, case.id_case)
            weights = case_distribution._consumption_weights(self.connection, context, concept, [owner_id])
            self.assertEqual(20, weights[owner_id])
            profile = SimpleNamespace(active_modules=("ACS",), concepts=(), onboarding_configuration={})
            excel_export_service._validate_normalized_inputs(self.connection, context, profile)
            for key in ("acs_variable_actual", "acs_fixed_actual"):
                self.connection.execute("INSERT INTO period_parameters(id_comunidad,id_periodo,parameter_key,numeric_value) VALUES (?,?,?,0)",
                    (self.community_id, context["id_periodo"], key))
            workbook = Workbook()
            workbook.active.title = "ACS"
            excel_export_service._write_meter_readings(self.connection, workbook, context, "ACS",
                {"sheet": "ACS", "start_row": 1, "end_row": 2,
                 "input_columns": {"initial": "A", "final": "B"}, "derived_columns": {"consumption": "C"}},
                materialize_consumption=True)
            self.assertEqual((100, 120, 20) if case == self.case else (120, 140, 20),
                tuple(workbook.active.cell(1, column).value for column in range(1, 4)))
            workbook.close()
        for case, expected_history in (
            (self.case, [("Febrero", 20.0)]),
            (second_case, [(self.case.name, 20.0)]),
        ):
            context = excel_export_service._input_case_context(self.connection, case.id_case)
            graph = case_letter_service._consumption_graphs(
                self.connection, context, [{"id_propietario": owner_id}],
            )[owner_id]
            with self.subTest(case=case.name, graph="current"):
                self.assertEqual(20.0, graph["owner_consumption"])
                self.assertEqual([20.0], graph["neighbor_consumptions"])
                self.assertEqual("m³", graph["unit"])
            with self.subTest(case=case.name, graph="historical"):
                self.assertEqual(expected_history, graph["history"])
        self.assertEqual(3, self.connection.execute("SELECT COUNT(*) FROM lecturas_vecino").fetchone()[0])
        original_period = excel_export_service._input_case_context(
            self.connection, self.case.id_case,
        )["id_periodo"]
        self.assertEqual(original_period, self.connection.execute(
            "SELECT id_periodo FROM lecturas_vecino WHERE fecha_lectura='2026-01-31'",
        ).fetchone()[0])

    def test_counter_reset_approval_targets_source_interval_inside_case(self):
        self.connection.execute("UPDATE regularization_cases SET fecha_fin='2026-02-28' WHERE id_case=?", (self.case.id_case,))
        document = self.add_confirmed_reading(final=5)
        case_ingestion.apply_confirmed_source(self.connection, self.case.id_case, document.id_document)
        issue = document_review.list_open_issues(self.connection, self.case.id_case)[0]
        self.connection.execute("INSERT INTO lecturas_vecino (id_propietario,id_periodo,tipo,fecha_lectura,valor_acumulado) SELECT id_propietario,id_periodo,tipo,'2026-02-28',150 FROM lecturas_vecino LIMIT 1")
        document_review.approve_counter_reset_estimate(self.connection, issue.id_issue,
            consumption="15", reason="Informe de sustitución", approved_by="Jose")
        rows = self.connection.execute("SELECT fecha_lectura,valor_acumulado FROM lecturas_vecino ORDER BY fecha_lectura").fetchall()
        self.assertEqual([("2026-01-01", 100), ("2026-01-31", 115), ("2026-02-28", 150)], [tuple(r) for r in rows])
        self.assertEqual("5", self.connection.execute("SELECT original_value FROM manual_corrections").fetchone()[0])

    def test_zero_final_reading_keeps_raw_observation_and_carries_previous_value(self):
        document = self.add_confirmed_reading(final=0)

        case_ingestion.apply_confirmed_source(self.connection, self.case.id_case, document.id_document)

        observations = self.connection.execute(
            """SELECT fecha_lectura,observed_value,status
                 FROM reading_observations ORDER BY fecha_lectura"""
        ).fetchall()
        effective = self.connection.execute(
            """SELECT valor_acumulado,estado,metodo_estimacion
                 FROM lecturas_vecino WHERE fecha_lectura='2026-01-31'"""
        ).fetchone()
        self.assertEqual(
            [("2026-01-01", 100, "observed"), ("2026-01-31", 0, "carried_forward")],
            [tuple(row) for row in observations],
        )
        self.assertEqual((100, "estimado", "carry_forward_zero"), tuple(effective))
        self.assertFalse(document_review.list_open_issues(self.connection, self.case.id_case))

    def test_zero_reading_without_prior_value_stays_open_for_review(self):
        self.connection.execute(
            """INSERT INTO propietarios (id_comunidad,codigo_vivienda,nombre_propietario)
               VALUES (?,'B','Vecino B')""",
            (self.community_id,),
        )
        self.connection.commit()
        document = case_ingestion.add_document_to_case(
            self.connection,
            self.case.id_case,
            source_path=self.reading_file,
            archive_root=self.archive_root,
            document_kind="reading",
            candidates={"vecinos": json.dumps([{
                "vivienda": "B", "tipo": "ACS", "fecha_ant": "2026-01-01",
                "val_ant": 0, "fecha_act": "2026-01-31", "val_act": 0,
            }])},
            required_fields=(),
        ).document

        case_ingestion.apply_confirmed_source(self.connection, self.case.id_case, document.id_document)

        self.assertEqual("review_required", self.connection.execute(
            "SELECT status FROM reading_observations WHERE observed_value=0 ORDER BY id_observation LIMIT 1"
        ).fetchone()[0])
        self.assertTrue(document_review.list_open_issues(self.connection, self.case.id_case))

    def test_confirming_initial_zeroes_for_one_source_applies_them_in_bulk(self):
        self.connection.execute(
            """INSERT INTO propietarios (id_comunidad,codigo_vivienda,nombre_propietario)
               VALUES (?,'B','Vecino B')""",
            (self.community_id,),
        )
        self.connection.commit()
        document = case_ingestion.add_document_to_case(
            self.connection,
            self.case.id_case,
            source_path=self.reading_file,
            archive_root=self.archive_root,
            document_kind="reading",
            candidates={"vecinos": json.dumps([{
                "vivienda": "B", "tipo": "ACS", "fecha_ant": "2026-01-01",
                "val_ant": 0, "fecha_act": "2026-01-31", "val_act": 0,
            }])},
            required_fields=(),
        ).document
        case_ingestion.apply_confirmed_source(self.connection, self.case.id_case, document.id_document)

        approve = getattr(document_review, "confirm_initial_zero_readings_for_source", None)
        self.assertTrue(callable(approve), "La fuente debe poder confirmar sus ceros iniciales en bloque")
        resolved = approve(
            self.connection, case_id=self.case.id_case, document_id=document.id_document,
            reason="Primer informe disponible", approved_by="Jose",
        )

        self.assertEqual(1, resolved)
        self.assertFalse(document_review.list_open_issues(self.connection, self.case.id_case))
        readings = self.connection.execute(
            """SELECT fecha_lectura,valor_acumulado,estado,metodo_estimacion
                 FROM lecturas_vecino ORDER BY fecha_lectura"""
        ).fetchall()
        self.assertEqual(
            [("2026-01-01", 0, "real", "confirmed_initial_zero"),
             ("2026-01-31", 0, "estimado", "carry_forward_zero")],
            [tuple(row) for row in readings],
        )
        self.assertEqual(
            [("observed",), ("carried_forward",)],
            [tuple(row) for row in self.connection.execute(
                "SELECT status FROM reading_observations ORDER BY fecha_lectura"
            )],
        )

    def test_confirming_zeroes_preserves_an_approved_counter_reset_reading(self):
        self.connection.execute(
            """UPDATE regularization_cases
               SET fecha_inicio='2025-07-01',fecha_fin='2026-07-31' WHERE id_case=?""",
            (self.case.id_case,),
        )
        self.connection.execute(
            """INSERT INTO propietarios (id_comunidad,codigo_vivienda,nombre_propietario)
               VALUES (?,'B','Vecino B')""",
            (self.community_id,),
        )
        owner_id = self.connection.execute(
            "SELECT id_propietario FROM propietarios WHERE id_comunidad=? AND codigo_vivienda='B'",
            (self.community_id,),
        ).fetchone()[0]
        period_id = expedient_service.link_case_to_period(self.connection, self.case.id_case)
        self.connection.execute(
            """INSERT INTO lecturas_vecino
               (id_propietario,id_periodo,tipo,fecha_lectura,valor_acumulado,estado,fuente)
               VALUES (?,?,'ACS','2025-07-01',9,'real','histórico')""",
            (owner_id, period_id),
        )
        self.connection.execute(
            """INSERT INTO lecturas_vecino
               (id_propietario,id_periodo,tipo,fecha_lectura,valor_acumulado,estado,fuente)
               VALUES (?,?,'ACS','2026-01-31',10,'real','lectura posterior')""",
            (owner_id, period_id),
        )
        self.connection.execute(
            """INSERT INTO lecturas_vecino
               (id_propietario,id_periodo,tipo,fecha_lectura,valor_acumulado,estado,
                metodo_estimacion,fuente,approved_by,approved_at)
               VALUES (?,?,'ACS','2026-07-31',9,'estimado','counter_reset_carry_forward',
                       'informe anterior','Jose',datetime('now'))""",
            (owner_id, period_id),
        )
        self.connection.commit()
        document = case_ingestion.add_document_to_case(
            self.connection,
            self.case.id_case,
            source_path=self.reading_file,
            archive_root=self.archive_root,
            document_kind="reading",
            candidates={"vecinos": json.dumps([{
                "vivienda": "B", "tipo": "ACS", "fecha_ant": "2025-07-01",
                "val_ant": 9, "fecha_act": "2026-07-31", "val_act": 0,
            }])},
            required_fields=(),
        ).document
        case_ingestion.apply_confirmed_source(self.connection, self.case.id_case, document.id_document)
        self.assertFalse(document_review.list_open_issues(self.connection, self.case.id_case))
        document_review.create_review_issue(
            self.connection, self.case.id_case, document.id_document,
            code="READING_ZERO_REVIEW", field_name="reading.B.ACS",
            message="Estado pendiente creado antes de aceptar reinicios aprobados",
        )

        resolved = document_review.confirm_initial_zero_readings_for_source(
            self.connection, case_id=self.case.id_case, document_id=document.id_document,
            reason="La fuente posterior contiene 0, pero ya existe una lectura aprobada", approved_by="Jose",
        )

        self.assertEqual(1, resolved)
        self.assertFalse(document_review.list_open_issues(self.connection, self.case.id_case))
        self.assertEqual(
            [
                ("2025-07-01", 9, "real"),
                ("2026-01-31", 10, "real"),
                ("2026-07-31", 9, "estimado"),
            ],
            [tuple(row) for row in self.connection.execute(
                "SELECT fecha_lectura,valor_acumulado,estado FROM lecturas_vecino ORDER BY fecha_lectura"
            )],
        )
        self.assertEqual(
            "counter_reset_carry_forward",
            self.connection.execute(
                "SELECT metodo_estimacion FROM lecturas_vecino WHERE fecha_lectura='2026-07-31'"
            ).fetchone()[0],
        )
        self.assertEqual(
            [("observed",), ("carried_forward",)],
            [tuple(row) for row in self.connection.execute(
                "SELECT status FROM reading_observations ORDER BY fecha_lectura"
            )],
        )

    def test_resolving_reading_conflict_updates_source_and_canonical_reading(self):
        self.connection.execute(
            """INSERT INTO propietarios (id_comunidad,codigo_vivienda,nombre_propietario)
               VALUES (?,'B','Vecino B')""",
            (self.community_id,),
        )
        owner_id = self.connection.execute(
            "SELECT id_propietario FROM propietarios WHERE id_comunidad=? AND codigo_vivienda='B'",
            (self.community_id,),
        ).fetchone()[0]
        period_id = expedient_service.link_case_to_period(self.connection, self.case.id_case)
        self.connection.executemany(
            """INSERT INTO lecturas_vecino
               (id_propietario,id_periodo,tipo,fecha_lectura,valor_acumulado,estado,fuente)
               VALUES (?,?,'ACS',? ,?,'real','histórico')""",
            [
                (owner_id, period_id, '2026-01-01', 100),
                (owner_id, period_id, '2026-01-31', 110),
            ],
        )
        self.connection.commit()
        document = case_ingestion.add_document_to_case(
            self.connection, self.case.id_case, source_path=self.reading_file,
            archive_root=self.archive_root, document_kind="reading",
            candidates={"vecinos": json.dumps([{
                "vivienda": "B", "tipo": "ACS", "fecha_ant": "2026-01-01",
                "val_ant": 100, "fecha_act": "2026-01-31", "val_act": 120,
            }])}, required_fields=(),
        ).document
        case_ingestion.apply_confirmed_source(self.connection, self.case.id_case, document.id_document)
        issue = document_review.list_open_issues(self.connection, self.case.id_case)[0]

        document_review.resolve_issue(
            self.connection, issue.id_issue, value="123", reason="Lectura comprobada en el informe",
        )

        self.assertFalse(document_review.list_open_issues(self.connection, self.case.id_case))
        self.assertEqual(
            (123.0, "real"),
            tuple(self.connection.execute(
                """SELECT valor_acumulado,estado FROM lecturas_vecino
                   WHERE id_propietario=? AND tipo='ACS' AND fecha_lectura='2026-01-31'""",
                (owner_id,),
            ).fetchone()),
        )
        rows = json.loads(self.connection.execute(
            "SELECT value FROM extraction_candidates WHERE id_document=? AND field_name='vecinos'",
            (document.id_document,),
        ).fetchone()[0])
        self.assertEqual(123.0, rows[0]["val_act"])

    def test_two_resets_in_one_source_have_independent_approval_targets(self):
        self.connection.execute("UPDATE regularization_cases SET fecha_fin='2026-02-28' WHERE id_case=?", (self.case.id_case,))
        self.connection.execute("INSERT INTO propietarios (id_comunidad,codigo_vivienda,nombre_propietario) VALUES (?,'A','Vecino')", (self.community_id,))
        rows = [
            {"vivienda": "A", "tipo": "ACS", "fecha_ant": "2026-01-01", "val_ant": 100,
             "fecha_act": "2026-01-31", "val_act": 5},
            {"vivienda": "A", "tipo": "ACS", "fecha_ant": "2026-02-01", "val_ant": 200,
             "fecha_act": "2026-02-28", "val_act": 8},
        ]
        document = self.ingest(SourceAnalysis.reading({"vecinos": json.dumps(rows)}))
        self.confirm(document)
        issues = document_review.list_open_issues(self.connection, self.case.id_case)
        self.assertEqual(2, len(issues))
        for issue in issues:
            document_review.approve_counter_reset_estimate(self.connection, issue.id_issue,
                consumption="15", reason="Cambio de contador", approved_by="Jose")
        self.assertEqual([100, 115, 200, 215], [row[0] for row in self.connection.execute(
            "SELECT valor_acumulado FROM lecturas_vecino ORDER BY fecha_lectura")])
        self.assertEqual(2, self.connection.execute("SELECT COUNT(*) FROM manual_corrections").fetchone()[0])
        self.assertFalse(document_review.case_has_unapplied_sources(self.connection, self.case.id_case))

    def test_invalid_reset_target_cannot_fall_back_to_unrelated_case_boundaries(self):
        document = self.add_confirmed_reading(final=5)
        case_ingestion.apply_confirmed_source(self.connection, self.case.id_case, document.id_document)
        issue = document_review.list_open_issues(self.connection, self.case.id_case)[0]
        owner = self.connection.execute("INSERT INTO propietarios (id_comunidad,codigo_vivienda,nombre_propietario) VALUES (?,'B','Otro vecino')", (self.community_id,)).lastrowid
        other = self.connection.execute("INSERT INTO lecturas_vecino (id_propietario,id_periodo,tipo,fecha_lectura,valor_acumulado) SELECT ?,id_periodo,'ACS','2026-01-31',8 FROM lecturas_vecino LIMIT 1", (owner,)).lastrowid
        self.connection.execute("UPDATE counter_reset_targets SET final_reading_id=? WHERE id_issue=?", (other, issue.id_issue))
        with self.assertRaises(LookupError):
            document_review.approve_counter_reset_estimate(self.connection, issue.id_issue,
                consumption="15", reason="Informe", approved_by="Jose")
        self.assertEqual(0, self.connection.execute("SELECT COUNT(*) FROM manual_corrections").fetchone()[0])

    def test_reading_correction_cannot_silently_keep_the_old_canonical_values(self):
        self.connection.execute("INSERT INTO propietarios (id_comunidad,codigo_vivienda,nombre_propietario) VALUES (?,'A','Vecino')", (self.community_id,))
        rows = [{"vivienda": "A", "tipo": "ACS", "fecha_ant": "2026-01-01", "val_ant": 100,
                 "fecha_act": "2026-01-31", "val_act": 120}]
        document = self.ingest(SourceAnalysis.reading({"vecinos": json.dumps(rows)}))
        self.confirm(document)
        issue = document_review.create_review_issue(self.connection, self.case.id_case, document.id_document,
            code="SOURCE_VALUE_REVIEW", field_name="vecinos", message="Comprobar lectura final")
        rows[0]["val_act"] = 125
        document_review.resolve_issue(self.connection, issue.id_issue, value=json.dumps(rows), reason="Original")
        final = self.connection.execute("SELECT valor_acumulado FROM lecturas_vecino WHERE fecha_lectura='2026-01-31'").fetchone()[0]
        self.assertTrue(final == 125 or document_review.list_open_issues(self.connection, self.case.id_case),
                        "A reviewed correction must apply or raise an explicit canonical conflict")

    def test_duplicate_reading_headers_require_review_instead_of_picking_one(self):
        self.reading_file.write_text("vivienda;tipo;fecha_ant;val_ant;fecha_act;val_act;lectura final\nA;ACS;2026-01-01;100;2026-01-31;120;130\n", encoding="utf-8")
        analysis = source_analysis.analyse_tabular(self.reading_file)
        self.assertEqual("unknown", analysis.kind)
        self.assertIn("ambig", analysis.review_message.lower())

    def test_xlsx_readings_with_date_cells_are_applied(self):
        from openpyxl import Workbook
        self.connection.execute("INSERT INTO propietarios (id_comunidad,codigo_vivienda,nombre_propietario) VALUES (?,'A','Vecino')", (self.community_id,))
        path = self.reading_file.with_suffix(".xlsx")
        workbook = Workbook()
        workbook.active.title = "Lecturas"
        workbook.active.append(["Vivienda", "Servicio", "Fecha anterior", "Lectura anterior", "Fecha actual", "Lectura actual"])
        workbook.active.append(["A", "ACS", date(2026, 1, 1), 100, date(2026, 1, 31), 120])
        workbook.save(path)
        workbook.close()
        analysis = source_analysis.analyse_tabular(path)
        self.assertEqual("Lecturas", analysis.locator.sheet)
        self.confirm(self.ingest(analysis, path))
        self.assertEqual(2, self.connection.execute("SELECT COUNT(*) FROM lecturas_vecino").fetchone()[0])

    def test_legacy_xls_rows_are_extracted_and_request_only_missing_service(self):
        import xlwt
        self.connection.execute("INSERT INTO propietarios (id_comunidad,codigo_vivienda,nombre_propietario) VALUES (?,'A','Vecino')", (self.community_id,))
        path = self.reading_file.with_suffix(".xls")
        workbook = xlwt.Workbook()
        sheet = workbook.add_sheet("Lecturas")
        for row_index, row in enumerate([["Cod.", "Propiedad", "12/25", "1/26"], [1, "A", 100, 120]]):
            for column, value in enumerate(row):
                sheet.write(row_index, column, value)
        workbook.save(str(path))
        analysis = source_analysis.analyse_tabular(path)
        self.assertEqual("reading", analysis.kind)
        self.assertIn("vecinos", analysis.candidates)
        document = self.ingest(analysis, path)
        issues = document_review.list_open_issues(self.connection, self.case.id_case)
        self.assertEqual({"tipo"}, {issue.field_name for issue in issues})
        self.assertEqual("2025-12-01", analysis.candidates["fecha_inicio"])
        self.assertEqual("2026-01-31", analysis.candidates["fecha_fin"])
        values = {"tipo": "ACS"}
        for issue in issues:
            document_review.resolve_issue(self.connection, issue.id_issue, value=values[issue.field_name], reason="Cabecera verificada")
        self.confirm(document)
        self.assertEqual([100, 120], [row[0] for row in self.connection.execute(
            "SELECT valor_acumulado FROM lecturas_vecino ORDER BY fecha_lectura")])

    def test_unknown_document_context_is_visible_without_candidate_rows(self):
        self.ingest(SourceAnalysis.unknown(locator=SourceLocator(page=2, fragment="Ambiguous fragment")))
        issue = document_review.list_open_issues(self.connection, self.case.id_case)[0]
        app = SimpleNamespace(ruta_bd_expedientes=self.database_path)
        with patch.object(expedient_ui.messagebox, "showinfo") as shown:
            expedient_ui.show_issue_context(app, issue)
        self.assertIn("Página 2", shown.call_args.args[1])
        self.assertIn("Ambiguous fragment", shown.call_args.args[1])

    def test_database_reset_rejects_noop_and_incomplete_schema_initializers(self):
        self.connection.close()
        for initializer in (lambda path: None, lambda path: sqlite3.connect(path).execute("CREATE TABLE marker(value TEXT)").connection.close()):
            with self.subTest(initializer=initializer):
                with self.assertRaisesRegex(database_reset.DatabaseResetError, "esquema"):
                    database_reset.reset_database(self.database_path,
                        backup_root=self.database_path.parent / "backups", initialise=initializer)
                with sqlite3.connect(self.database_path) as check:
                    self.assertEqual(1, check.execute("SELECT COUNT(*) FROM comunidades").fetchone()[0])


if __name__ == "__main__":
    unittest.main()
