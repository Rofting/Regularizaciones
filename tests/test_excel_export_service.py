import json
import hashlib
import sqlite3
import sys
import tempfile
import unittest
from copy import copy
from contextlib import redirect_stdout
from datetime import date
from io import StringIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import PatternFill


PROJECT_ROOT = Path(__file__).resolve().parents[1]
CORE_DIR = PROJECT_ROOT / "core"
if str(CORE_DIR) not in sys.path:
    sys.path.insert(0, str(CORE_DIR))

import excel_generator
import gestor_bd
from excel_export_service import (
    ExportBlockedError, _case_context, _input_hash, _profile_for_community,
    _is_winter, _season_boundary, _validate_normalized_inputs,
    generate_official_excel,
)
from excel_validation import WorkbookValidationError, validate_workbook, workbook_fingerprint
from office_recalculation import (
    DeterministicRecalculator,
    LibreOfficeRecalculator,
    RecalculationError,
)


PROFILE_DATA = {
    "key": "658_acs_v1",
    "version": "1",
    "community_code": "658",
    "template_relative_path": "plantillas/comunidades/658/658_acs_v1.xlsx",
    "active_modules": ["GAS", "ELECTRICIDAD", "AGUA", "OTROS_GASTOS", "ACS"],
    "required_sheets": [
        "DATOS", "GAS", "ELECTRICIDAD", "AGUA", "OTROS GASTOS",
        "LECTURAS ACS M3", "ANALISIS",
    ],
    "required_formula_cells": [["ANALISIS", "H23"]],
    "concepts": [
        {
            "key": "acs_fixed", "allocation_method": "equal",
            "actual_source": "period_parameters.acs_fixed_actual",
            "billed_source": "period_parameters.acs_fixed_billed", "required": True,
        },
        {
            "key": "acs_variable", "allocation_method": "consumption",
            "actual_source": "period_parameters.acs_variable_actual",
            "billed_source": "period_parameters.acs_variable_billed", "required": True,
        },
        {
            "key": "extraordinary_expense", "allocation_method": "coefficient",
            "actual_source": "period_parameters.extraordinary_expense_actual",
            "billed_source": None, "required": False,
        },
    ],
    "workbook_layout": {
        "metadata_cells": {
            "community_name": ["DATOS", "A1"],
            "period_label": ["DATOS", "A3"],
            "owner_count": ["DATOS", "D4"],
        },
        "tables": {
            "GAS": {
                "sheet": "GAS", "start_row": 10, "end_row": 30,
                "input_columns": {
                    "invoice_date": "B", "end_date": "D", "consumption": "J",
                    "fixed": "K", "variable": "L", "provider": "O",
                },
                "derived_columns": {"total": "M"},
            },
            "ELECTRICIDAD": {
                "sheet": "ELECTRICIDAD", "start_row": 10, "end_row": 21,
                "input_columns": {
                    "invoice_date": "B", "start_date": "C", "end_date": "D",
                    "consumption": "E", "fixed": "F", "variable": "G",
                    "provider": "J",
                },
                "derived_columns": {"total": "H"},
            },
            "AGUA": {
                "sheet": "AGUA", "start_row": 10, "end_row": 21,
                "input_columns": {
                    "invoice_date": "B", "end_date": "D", "consumption": "I",
                    "variable": "T", "fixed": "U",
                },
                "derived_columns": {"total": "R"},
            },
            "OTROS_GASTOS": {
                "sheet": "OTROS GASTOS", "start_row": 13, "end_row": 23,
                "columns": {"date": "D", "description": "E", "amount": "F"},
            },
            "ACS": {
                "sheet": "LECTURAS ACS M3", "start_row": 8, "end_row": 25,
                "input_columns": {
                    "charge_date": "B", "final_date": "C", "final": "D",
                    "initial_date": "E", "initial": "F", "variable_fee": "H",
                    "fixed_fee": "I",
                },
                "derived_columns": {
                    "consumption": "G", "total": "J", "variable_unit": "L",
                    "fixed_unit": "M",
                },
            },
        },
        "parameter_cells": {
            "acs_fixed_actual": ["ANALISIS", "H64"],
            "acs_variable_actual": ["ANALISIS", "J64"],
            "acs_fixed_billed": ["ANALISIS", "H66"],
            "acs_variable_billed": ["ANALISIS", "J66"],
        },
        "total_checks": {
            "invoice_total:GAS": ["GAS", "M10:M30"],
            "invoice_component:GAS:fixed": ["GAS", "K10:K30"],
            "invoice_component:GAS:variable": ["GAS", "L10:L30"],
            "invoice_total:ELECTRICIDAD": ["ELECTRICIDAD", "H10:H21"],
            "invoice_component:ELECTRICIDAD:fixed": ["ELECTRICIDAD", "F10:F21"],
            "invoice_component:ELECTRICIDAD:variable": ["ELECTRICIDAD", "G10:G21"],
            "invoice_total:AGUA": ["AGUA", "R10:R21"],
            "invoice_component:AGUA:fixed": ["AGUA", "U10:U21"],
            "invoice_component:AGUA:variable": ["AGUA", "T10:T21"],
            "parameter:extraordinary_expense_actual": ["OTROS GASTOS", "F13:F23"],
            "reading_total:ACS": ["LECTURAS ACS M3", "G8:G25"],
            "parameter:acs_fixed_actual": ["ANALISIS", "H64"],
            "parameter:acs_variable_actual": ["ANALISIS", "J64"],
            "parameter:acs_fixed_billed": ["ANALISIS", "H66"],
            "parameter:acs_variable_billed": ["ANALISIS", "J66"],
        },
    },
}


def _make_template(path: Path) -> Path:
    workbook = Workbook()
    workbook.active.title = "DATOS"
    for sheet_name in PROFILE_DATA["required_sheets"][1:]:
        workbook.create_sheet(sheet_name)
    for sheet in workbook.worksheets:
        sheet.print_area = "A1:X40"
        sheet.freeze_panes = "B8"
        sheet.row_dimensions[2].height = 27
        sheet.column_dimensions["A"].width = 19
    workbook["GAS"].merge_cells("A2:C2")
    workbook["GAS"]["A2"] = "Diseño saneado"
    workbook["ANALISIS"]["H23"] = "=SUM(H1:H22)"
    workbook["ANALISIS"]["B2"] = "NO TOCAR"
    workbook["ANALISIS"]["B2"].fill = PatternFill("solid", fgColor="D9EAF7")
    workbook["GAS"]["M10"] = "=K10+L10"
    workbook["GAS"]["M11"] = "=K11+L11"
    workbook["GAS"]["M31"] = "=SUM(M10:M30)"
    workbook["ELECTRICIDAD"]["H10"] = "=F10+G10"
    workbook["ELECTRICIDAD"]["H23"] = "=SUM(H10:H21)"
    workbook["AGUA"]["R10"] = "=T10+U10"
    workbook["AGUA"]["R23"] = "=SUM(R10:R21)"
    workbook["LECTURAS ACS M3"]["G8"] = "=D8-F8"
    workbook["LECTURAS ACS M3"]["J8"] = "=SUM(H8:I8)"
    workbook["LECTURAS ACS M3"]["L8"] = "=H8/G8"
    workbook["LECTURAS ACS M3"]["J9"] = "=SUM(H9:I9)"
    workbook["LECTURAS ACS M3"]["M9"] = "=I9/'DATOS'!$D$4"
    workbook["LECTURAS ACS M3"]["G15"] = 999
    workbook["LECTURAS ACS M3"]["J15"] = "=SUM(H15:I15)"
    workbook["LECTURAS ACS M3"]["L15"] = 999
    workbook["LECTURAS ACS M3"]["M15"] = 999
    for column in ("G", "H", "I", "J"):
        workbook["LECTURAS ACS M3"][f"{column}27"] = f"=SUM({column}8:{column}25)"
    for sheet_name, columns, first, last in (
        ("GAS", ("B", "D"), 10, 30),
        ("ELECTRICIDAD", ("B", "C", "D"), 10, 21),
        ("AGUA", ("B", "D"), 10, 21),
        ("OTROS GASTOS", ("D",), 13, 23),
        ("LECTURAS ACS M3", ("B", "C", "E"), 8, 25),
    ):
        for column in columns:
            for row in range(first, last + 1):
                workbook[sheet_name][f"{column}{row}"].number_format = "DD/MM/YYYY"
    chart = BarChart()
    chart.add_data(Reference(workbook["GAS"], min_col=11, min_row=10, max_row=11))
    workbook["GAS"].add_chart(chart, "P2")
    workbook["GAS"]["P5"] = "CELDA AJENA"
    workbook.save(path)
    with ZipFile(path, "a", ZIP_DEFLATED) as archive:
        archive.writestr("customXml/item1.xml", b"<audit>preservar</audit>")
    return path


class ExcelExportServiceTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.root = Path(self.directory.name)
        self.project_root = self.root / "proyecto"
        config_dir = self.project_root / "config" / "excel_profiles"
        template_dir = self.project_root / "plantillas" / "comunidades" / "658"
        config_dir.mkdir(parents=True)
        template_dir.mkdir(parents=True)
        (config_dir / "658_acs_v1.json").write_text(
            json.dumps(PROFILE_DATA), encoding="utf-8"
        )
        self.template = _make_template(template_dir / "658_acs_v1.xlsx")
        self.output_root = self.root / "salidas"
        self.database_path = self.root / "gestion.db"
        with redirect_stdout(StringIO()):
            gestor_bd.crear_bd(str(self.database_path))
        self.connection = sqlite3.connect(self.database_path)
        self.connection.row_factory = sqlite3.Row
        self.connection.execute("PRAGMA foreign_keys = ON")
        self.community_id = gestor_bd.obtener_o_crear_comunidad(
            self.connection, "658", "Comunidad sintética 658"
        )
        case_cursor = self.connection.execute(
            """INSERT INTO regularization_cases
               (id_comunidad,nombre,fecha_inicio,fecha_fin,estado)
               VALUES (?,?,?,?,?)""",
            (self.community_id, "2025-2026", "2025-09-01", "2026-08-31",
             "ready_for_calculation"),
        )
        self.case_id = int(case_cursor.lastrowid)
        period_cursor = self.connection.execute(
            """INSERT INTO periodos
               (id_comunidad,nombre,fecha_inicio,fecha_fin,estado)
               VALUES (?,?,?,?,?)""",
            (self.community_id, "2025-2026", "2025-09-01", "2026-08-31", "abierto"),
        )
        self.period_id = int(period_cursor.lastrowid)
        self.connection.execute(
            "UPDATE regularization_cases SET id_periodo=? WHERE id_case=?",
            (self.period_id, self.case_id),
        )
        self._insert_normalized_inputs()
        self.connection.commit()

    def tearDown(self):
        self.connection.close()
        self.directory.cleanup()

    @property
    def official_output(self) -> Path:
        return self.output_root / "Excels_Maestros" / "Comunidad_658.xlsx"

    def _insert_normalized_inputs(self):
        invoices = [
            ("GAS", "Gas sintético", "G-1", "2025-10-05", "2025-09-01",
             "2025-09-30", 30, 110.0, "kWh", 20.0, 80.0, 100.0),
            ("ELECTRICIDAD", "Luz sintética", "E-1", "2025-10-08", "2025-09-01",
             "2025-09-30", 30, 250.0, "kWh", 25.0, 75.0, 100.0),
            ("AGUA", "Agua sintética", "A-1", "2025-10-09", "2025-09-01",
             "2025-09-30", 30, 15.0, "m3", 20.0, 40.0, 60.0),
        ]
        for row in invoices:
            cursor = self.connection.execute(
                """INSERT INTO facturas
                   (id_comunidad,id_periodo,tipo_suministro,proveedor,
                    cups_o_referencia,num_factura,fecha_factura,fecha_inicio,
                    fecha_fin,dias_facturados,consumo_total,unidad_consumo,
                    termino_fijo,termino_variable,importe_total)
                   VALUES (?,?,?,?,'SINTETICO',?,?,?,?,?,?,?,?,?,?)""",
                (self.community_id, self.period_id, *row),
            )
            for key, amount in (("fixed", row[-3]), ("variable", row[-2]), ("total", row[-1])):
                self.connection.execute(
                    """INSERT INTO invoice_components
                       (id_factura,component_key,amount,unit,source_sheet,source_cell)
                       VALUES (?,?,?,?,?,?)""",
                    (cursor.lastrowid, key, amount, "EUR", row[0], "A1"),
                )
        parameters = {
            "acs_fixed_actual": 78.0,
            "acs_variable_actual": 182.0,
            "acs_fixed_billed": 112.0,
            "acs_variable_billed": 168.0,
            "extraordinary_expense_actual": 120.0,
        }
        for key, amount in parameters.items():
            self.connection.execute(
                """INSERT INTO period_parameters
                   (id_comunidad,id_periodo,parameter_key,numeric_value,unit,
                    source_sheet,source_cell) VALUES (?,?,?,?,?,?,?)""",
                (self.community_id, self.period_id, key, amount, "EUR", "ANALISIS", "A1"),
            )
        for code, name, coefficient, first, last in (
            ("P1-A", "Persona Uno", 50.0, 100.0, 125.0),
            ("P1-B", "Persona Dos", 50.0, 200.0, 215.0),
        ):
            owner = self.connection.execute(
                """INSERT INTO propietarios
                   (id_comunidad,codigo_vivienda,nombre_propietario,coeficiente,activo)
                   VALUES (?,?,?,?,1)""",
                (self.community_id, code, name, coefficient),
            )
            self.connection.executemany(
                """INSERT INTO lecturas_vecino
                   (id_propietario,id_periodo,tipo,fecha_lectura,valor_acumulado,
                    estado,fuente) VALUES (?,?, 'ACS', ?, ?, 'real', 'sintetico')""",
                [
                    (owner.lastrowid, self.period_id, "2025-09-01", first),
                    (owner.lastrowid, self.period_id, "2026-08-31", last),
                ],
            )

    def _expected_totals(self):
        return {
            "invoice_total:GAS": 10_000,
            "invoice_component:GAS:fixed": 2_000,
            "invoice_component:GAS:variable": 8_000,
            "invoice_total:ELECTRICIDAD": 10_000,
            "invoice_component:ELECTRICIDAD:fixed": 2_500,
            "invoice_component:ELECTRICIDAD:variable": 7_500,
            "invoice_total:AGUA": 6_000,
            "invoice_component:AGUA:fixed": 2_000,
            "invoice_component:AGUA:variable": 4_000,
            "parameter:extraordinary_expense_actual": 12_000,
            "reading_total:ACS": 4_000,
            "parameter:acs_fixed_actual": 7_800,
            "parameter:acs_variable_actual": 18_200,
            "parameter:acs_fixed_billed": 11_200,
            "parameter:acs_variable_billed": 16_800,
        }

    def test_export_writes_typed_data_preserves_template_and_reports_progress(self):
        stages = []
        result = generate_official_excel(
            self.connection,
            id_case=self.case_id,
            project_root=self.project_root,
            output_root=self.output_root,
            progress=lambda stage, payload: stages.append(stage),
            recalculator=DeterministicRecalculator(),
        )

        self.assertEqual(self.official_output, result.output_path)
        self.assertIsNone(result.backup_path)
        self.assertEqual("validated", result.stage)
        self.assertEqual(
            "validated",
            self.connection.execute(
                "SELECT status FROM excel_export_runs WHERE id_export_run=?",
                (result.id_export_run,),
            ).fetchone()[0],
        )
        self.assertEqual(
            ["validate_case", "prepare_template", "write_gas",
             "write_electricidad", "write_agua", "write_otros_gastos",
             "write_acs", "recalculate", "reconcile", "parameters", "publish"],
            stages,
        )
        workbook = load_workbook(result.output_path, data_only=False)
        self.assertIsInstance(workbook["GAS"]["B10"].value, date)
        self.assertEqual(20, workbook["GAS"]["K10"].value)
        self.assertEqual(80, workbook["GAS"]["L10"].value)
        self.assertEqual("=K10+L10", workbook["GAS"]["M10"].value)
        self.assertEqual("=SUM(M10:M30)", workbook["GAS"]["M31"].value)
        self.assertIsNone(workbook["GAS"]["B32"].value)
        self.assertEqual("=SUM(H1:H22)", workbook["ANALISIS"]["H23"].value)
        self.assertIsInstance(workbook["LECTURAS ACS M3"]["B8"].value, date)
        self.assertIsNone(workbook["LECTURAS ACS M3"]["A8"].value)
        self.assertEqual("=D8-F8", workbook["LECTURAS ACS M3"]["G8"].value)
        self.assertEqual(182, workbook["LECTURAS ACS M3"]["H8"].value)
        self.assertEqual(78, workbook["LECTURAS ACS M3"]["I9"].value)
        self.assertEqual("=SUM(H8:I8)", workbook["LECTURAS ACS M3"]["J8"].value)
        self.assertEqual("=H8/G8", workbook["LECTURAS ACS M3"]["L8"].value)
        self.assertEqual("=I9/'DATOS'!$D$4", workbook["LECTURAS ACS M3"]["M9"].value)
        self.assertEqual("=SUM(G8:G25)", workbook["LECTURAS ACS M3"]["G27"].value)
        for column in ("G", "J", "L", "M"):
            self.assertIsNone(workbook["LECTURAS ACS M3"][f"{column}15"].value)
        self.assertEqual("NO TOCAR", workbook["ANALISIS"]["B2"].value)
        self.assertEqual("00D9EAF7", workbook["ANALISIS"]["B2"].fill.fgColor.rgb)
        self.assertEqual("'ANALISIS'!$A$1:$X$40", str(workbook["ANALISIS"].print_area))
        self.assertEqual("A2:C2", str(next(iter(workbook["GAS"].merged_cells.ranges))))
        self.assertEqual("B8", workbook["GAS"].freeze_panes)
        self.assertEqual(1, len(workbook["GAS"]._charts))
        workbook.close()
        with ZipFile(result.output_path) as archive:
            self.assertEqual(b"<audit>preservar</audit>", archive.read("customXml/item1.xml"))
            with ZipFile(self.template) as template_archive:
                self.assertEqual(
                    template_archive.read("xl/styles.xml"), archive.read("xl/styles.xml")
                )

    def test_open_issue_blocks_export_without_replacing_predecessor(self):
        self.official_output.parent.mkdir(parents=True)
        self.official_output.write_bytes(b"SALIDA ANTERIOR")
        document = self.connection.execute(
            """INSERT INTO source_documents
               (id_case,original_name,archived_path,sha256,document_kind,status)
               VALUES (?,?,?,?,?,'under_review')""",
            (self.case_id, "dato.xlsx", "archivo/dato.xlsx", "a" * 64, "excel_master"),
        )
        self.connection.execute(
            """INSERT INTO review_issues
               (id_case,id_document,code,field_name,message,status)
               VALUES (?,?,?,?,?,'open')""",
            (self.case_id, document.lastrowid, "MISSING_DATA", "AGUA.total", "Falta dato"),
        )
        self.connection.commit()

        with self.assertRaisesRegex(ExportBlockedError, "incidencia"):
            generate_official_excel(
                self.connection, id_case=self.case_id,
                project_root=self.project_root, output_root=self.output_root,
                recalculator=DeterministicRecalculator(),
            )

        self.assertEqual(b"SALIDA ANTERIOR", self.official_output.read_bytes())

    def test_mismatched_invoice_components_block_official_output(self):
        self.connection.execute(
            """UPDATE invoice_components SET amount=10
               WHERE component_key='fixed' AND id_factura=(
                   SELECT id_factura FROM facturas WHERE tipo_suministro='AGUA'
               )"""
        )
        self.connection.commit()

        with self.assertRaisesRegex(ExportBlockedError, "componentes"):
            generate_official_excel(
                self.connection, id_case=self.case_id,
                project_root=self.project_root, output_root=self.output_root,
                recalculator=DeterministicRecalculator(),
            )

        self.assertFalse(self.official_output.exists())

    def test_invoice_without_optional_fixed_variable_breakdown_is_exportable(self):
        invoice_id = self.connection.execute(
            "SELECT id_factura FROM facturas WHERE tipo_suministro='GAS'"
        ).fetchone()[0]
        self.connection.execute(
            "DELETE FROM invoice_components WHERE id_factura=? AND component_key IN ('fixed','variable')",
            (invoice_id,),
        )
        self.connection.execute(
            "UPDATE facturas SET termino_fijo=NULL,termino_variable=NULL WHERE id_factura=?",
            (invoice_id,),
        )
        self.connection.commit()

        case = _case_context(self.connection, self.case_id)
        profile = _profile_for_community(
            self.connection, self.project_root, "658", self.community_id,
        )

        _validate_normalized_inputs(self.connection, case, profile)

    def test_formula_error_fails_auditable_run_without_replacing_predecessor(self):
        self.official_output.parent.mkdir(parents=True)
        self.official_output.write_bytes(b"SALIDA VALIDADA PREVIA")

        def corrupt_formula(path: Path):
            workbook = load_workbook(path)
            workbook["ANALISIS"]["H23"] = "#VALUE!"
            workbook.save(path)

        with self.assertRaisesRegex(WorkbookValidationError, "fórmula"):
            generate_official_excel(
                self.connection, id_case=self.case_id,
                project_root=self.project_root, output_root=self.output_root,
                recalculator=DeterministicRecalculator(after_recalculate=corrupt_formula),
            )

        self.assertEqual(b"SALIDA VALIDADA PREVIA", self.official_output.read_bytes())
        run = self.connection.execute(
            """SELECT status,error_message,diagnostic_path FROM excel_export_runs
               ORDER BY id_export_run DESC LIMIT 1"""
        ).fetchone()
        self.assertEqual("failed", run["status"])
        self.assertIn("fórmula", run["error_message"])
        self.assertTrue(Path(run["diagnostic_path"]).is_file())

    def test_second_success_backs_up_exact_preceding_workbook(self):
        first = generate_official_excel(
            self.connection, id_case=self.case_id,
            project_root=self.project_root, output_root=self.output_root,
            recalculator=DeterministicRecalculator(),
        )
        preceding = first.output_path.read_bytes()
        self.connection.execute(
            "UPDATE period_parameters SET numeric_value=79 WHERE parameter_key='acs_fixed_actual'"
        )
        self.connection.commit()

        second = generate_official_excel(
            self.connection, id_case=self.case_id,
            project_root=self.project_root, output_root=self.output_root,
            recalculator=DeterministicRecalculator(),
        )

        self.assertIsNotNone(second.backup_path)
        self.assertEqual(preceding, second.backup_path.read_bytes())
        self.assertNotEqual(preceding, second.output_path.read_bytes())

    def test_optional_extraordinary_expense_does_not_block_export(self):
        self.connection.execute(
            "DELETE FROM period_parameters WHERE parameter_key='extraordinary_expense_actual'"
        )
        self.connection.commit()

        result = generate_official_excel(
            self.connection,
            id_case=self.case_id,
            project_root=self.project_root,
            output_root=self.output_root,
            recalculator=DeterministicRecalculator(),
        )

        self.assertEqual("validated", result.stage)
        workbook = load_workbook(result.output_path, data_only=False)
        self.assertIsNone(workbook["OTROS GASTOS"]["F13"].value)
        workbook.close()

    def test_validator_rejects_wrong_declared_total(self):
        result = generate_official_excel(
            self.connection, id_case=self.case_id,
            project_root=self.project_root, output_root=self.output_root,
            recalculator=DeterministicRecalculator(),
        )
        workbook = load_workbook(result.output_path)
        workbook["ANALISIS"]["H64"] = 999
        workbook.save(result.output_path)

        from excel_profiles import load_profile
        with self.assertRaisesRegex(WorkbookValidationError, "acs_fixed_actual"):
            validate_workbook(
                result.output_path,
                load_profile("658_acs_v1", self.project_root),
                self._expected_totals(),
            )

    def test_input_hash_covers_eligible_expenses_and_active_owner_identity(self):
        from excel_profiles import load_profile

        case = _case_context(self.connection, self.case_id)
        profile = load_profile("658_acs_v1", self.project_root)
        initial = _input_hash(self.connection, case, profile)
        self.connection.execute(
            """INSERT INTO gastos_extra(id_comunidad,fecha,descripcion,importe_total,activo)
               VALUES (?,?,?,?,1)""",
            (self.community_id, "2026-01-10", "Gasto saneado", 12.5),
        )
        self.connection.commit()
        self.assertNotEqual(initial, _input_hash(self.connection, case, profile))

        after_expense = _input_hash(self.connection, case, profile)
        self.connection.execute(
            """UPDATE propietarios SET nombre_propietario='Identidad modificada',activo=0
               WHERE id_comunidad=? AND codigo_vivienda='P1-A'""",
            (self.community_id,),
        )
        self.connection.commit()
        self.assertNotEqual(after_expense, _input_hash(self.connection, case, profile))

    def test_registered_active_profile_selects_its_exact_version(self):
        v2 = dict(PROFILE_DATA)
        v2["key"] = "658_acs_v2"
        v2["version"] = "2"
        (self.project_root / "config" / "excel_profiles" / "658_acs_v2.json").write_text(
            json.dumps(v2), encoding="utf-8"
        )
        self.connection.execute(
            """INSERT INTO excel_template_profiles
               (id_comunidad,profile_key,profile_version,template_relative_path,
                template_sha256,profile_sha256,status)
               VALUES (?,?,?,?,?,?,'active')""",
            (
                self.community_id, "658_acs_v1", "1",
                PROFILE_DATA["template_relative_path"], "a" * 64,
                hashlib.sha256(
                    (self.project_root / "config" / "excel_profiles" / "658_acs_v1.json").read_bytes()
                ).hexdigest(),
            ),
        )
        self.connection.commit()
        self.assertEqual(
            "658_acs_v1",
            _profile_for_community(self.connection, self.project_root, "658", self.community_id).key,
        )
        self.connection.execute(
            """INSERT INTO excel_template_profiles
               (id_comunidad,profile_key,profile_version,template_relative_path,
                template_sha256,profile_sha256,status)
               VALUES (?,?,?,?,?,?,'active')""",
            (self.community_id, "658_acs_v2", "2",
             PROFILE_DATA["template_relative_path"], "c" * 64, "d" * 64),
        )
        self.connection.commit()
        with self.assertRaisesRegex(ExportBlockedError, "varios perfiles activos"):
            _profile_for_community(self.connection, self.project_root, "658", self.community_id)

    def test_active_profile_rejects_same_version_with_changed_bytes(self):
        self.connection.execute(
            """INSERT INTO excel_template_profiles
               (id_comunidad,profile_key,profile_version,template_relative_path,
                template_sha256,profile_sha256,status)
               VALUES (?,?,?,?,?,?,'active')""",
            (
                self.community_id, "658_acs_v1", "1",
                PROFILE_DATA["template_relative_path"], "a" * 64,
                hashlib.sha256(
                    (self.project_root / "config" / "excel_profiles" / "658_acs_v1.json").read_bytes()
                ).hexdigest(),
            ),
        )
        profile_path = self.project_root / "config" / "excel_profiles" / "658_acs_v1.json"
        profile_path.write_bytes(profile_path.read_bytes() + b"\n")
        self.connection.commit()

        with self.assertRaisesRegex(ExportBlockedError, "huella"):
            _profile_for_community(self.connection, self.project_root, "658", self.community_id)

    def test_validator_rejects_layout_or_chart_design_change(self):
        result = generate_official_excel(
            self.connection, id_case=self.case_id,
            project_root=self.project_root, output_root=self.output_root,
            recalculator=DeterministicRecalculator(),
        )
        expected = workbook_fingerprint(self.template, _profile_for_community(
            self.connection, self.project_root, "658", self.community_id
        ))
        workbook = load_workbook(result.output_path)
        workbook["GAS"].unmerge_cells("A2:C2")
        workbook.save(result.output_path)
        workbook.close()
        from excel_profiles import load_profile
        with self.assertRaisesRegex(WorkbookValidationError, "combinadas"):
            validate_workbook(
                result.output_path, load_profile("658_acs_v1", self.project_root),
                self._expected_totals(), expected_fingerprint=expected,
            )

    def test_validator_rejects_style_definition_change_without_changing_style_ids(self):
        result = generate_official_excel(
            self.connection, id_case=self.case_id,
            project_root=self.project_root, output_root=self.output_root,
            recalculator=DeterministicRecalculator(),
        )
        expected = workbook_fingerprint(self.template, _profile_for_community(
            self.connection, self.project_root, "658", self.community_id
        ))
        with ZipFile(result.output_path, "r") as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        parts["xl/styles.xml"] = parts["xl/styles.xml"].replace(
            b"00D9EAF7", b"00FFFFFF", 1
        )
        altered = result.output_path.with_name("estilo-alterado.xlsx")
        with ZipFile(altered, "w", ZIP_DEFLATED) as archive:
            for name, content in parts.items():
                archive.writestr(name, content)
        altered.replace(result.output_path)
        from excel_profiles import load_profile
        with self.assertRaisesRegex(WorkbookValidationError, "OOXML de diseño"):
            validate_workbook(
                result.output_path, load_profile("658_acs_v1", self.project_root),
                self._expected_totals(), expected_fingerprint=expected,
            )

    def test_validator_rejects_style_change_inside_mutable_input_cell(self):
        result = generate_official_excel(
            self.connection, id_case=self.case_id,
            project_root=self.project_root, output_root=self.output_root,
            recalculator=DeterministicRecalculator(),
        )
        expected = workbook_fingerprint(self.template, _profile_for_community(
            self.connection, self.project_root, "658", self.community_id
        ))
        workbook = load_workbook(result.output_path)
        workbook["LECTURAS ACS M3"]["B8"]._style = copy(workbook["ANALISIS"]["B2"]._style)
        workbook.save(result.output_path)
        workbook.close()
        from excel_profiles import load_profile
        with self.assertRaisesRegex(WorkbookValidationError, "estilo de una entrada"):
            validate_workbook(
                result.output_path, load_profile("658_acs_v1", self.project_root),
                self._expected_totals(), expected_fingerprint=expected,
            )

    def test_validator_accepts_libreoffice_normalization_of_implicit_default_style(self):
        def normalize_implicit_formula_style(path):
            workbook = load_workbook(path)
            cell = workbook["ANALISIS"]["H23"]
            normalized_font = copy(cell.font)
            normalized_font.name = "Arial"
            cell.font = normalized_font
            cell.number_format = "#,##0.00"
            workbook.save(path)
            workbook.close()

        result = generate_official_excel(
            self.connection, id_case=self.case_id,
            project_root=self.project_root, output_root=self.output_root,
            recalculator=DeterministicRecalculator(normalize_implicit_formula_style),
        )

        self.assertEqual("validated", result.stage)
        workbook = load_workbook(result.output_path)
        self.assertEqual("Arial", workbook["ANALISIS"]["H23"].font.name)
        self.assertEqual("#,##0.00", workbook["ANALISIS"]["H23"].number_format)
        workbook.close()

    def test_missing_libreoffice_has_an_understandable_error(self):
        recalculator = LibreOfficeRecalculator(
            executable=self.root / "programa-inexistente" / "soffice.exe"
        )
        with self.assertRaisesRegex(RecalculationError, "LibreOffice"):
            recalculator.recalculate(self.template, self.root / "recalculo")

    def test_legacy_entry_point_accepts_an_eligible_case(self):
        self.connection.close()
        result = excel_generator.regenerar_excel_comunidad(
            "658", ruta_bd=str(self.database_path),
            ruta_excels=str(self.output_root / "Excels_Maestros"),
            id_case=self.case_id, project_root=self.project_root,
            recalculator=DeterministicRecalculator(),
        )
        self.connection = sqlite3.connect(self.database_path)
        self.connection.row_factory = sqlite3.Row

        self.assertTrue(result["ok"])
        self.assertEqual(str(self.official_output), result["archivo"])
        self.assertEqual([], result["errores"])


class HeatingSeasonTest(unittest.TestCase):
    """La temporada decide si un consumo es de ACS o de calefacción.

    Las fechas son las del estudio real de la 658: cada factura cayó de un lado
    y el reparto de la fila de sumas depende de ello, así que la regla se fija
    aquí con esas mismas fechas.
    """

    GAS_INICIO, GAS_FIN = (12, 1), (5, 31)
    LUZ_INICIO, LUZ_FIN = (11, 1), (4, 30)

    def test_gas_invoices_follow_december_to_may_season(self):
        verano = [date(2025, 8, 27), date(2025, 9, 22), date(2025, 10, 1),
                  date(2025, 10, 27), date(2025, 11, 28), date(2026, 6, 24),
                  date(2026, 7, 23)]
        invierno = [date(2025, 12, 30), date(2026, 1, 22), date(2026, 2, 25),
                    date(2026, 3, 18), date(2026, 3, 30), date(2026, 4, 23),
                    date(2026, 4, 25), date(2026, 5, 29)]
        for dia in verano:
            self.assertFalse(_is_winter(dia, self.GAS_INICIO, self.GAS_FIN), dia)
        for dia in invierno:
            self.assertTrue(_is_winter(dia, self.GAS_INICIO, self.GAS_FIN), dia)

    def test_electricity_uses_its_own_november_to_april_season(self):
        # El mismo noviembre es verano para el gas e invierno para la luz.
        noviembre = date(2025, 11, 30)
        self.assertTrue(_is_winter(noviembre, self.LUZ_INICIO, self.LUZ_FIN))
        self.assertFalse(_is_winter(noviembre, self.GAS_INICIO, self.GAS_FIN))
        mayo = date(2026, 5, 31)
        self.assertFalse(_is_winter(mayo, self.LUZ_INICIO, self.LUZ_FIN))
        self.assertTrue(_is_winter(mayo, self.GAS_INICIO, self.GAS_FIN))

    def test_season_without_end_of_year_wrap_and_missing_date(self):
        self.assertTrue(_is_winter(date(2026, 3, 15), (1, 1), (6, 30)))
        self.assertFalse(_is_winter(date(2026, 8, 15), (1, 1), (6, 30)))
        self.assertFalse(_is_winter(None, self.GAS_INICIO, self.GAS_FIN))

    def test_boundaries_are_read_and_validated(self):
        self.assertEqual((12, 1), _season_boundary("12-01", "05-31"))
        self.assertEqual((5, 31), _season_boundary(None, "05-31"))
        for invalido in ("13-01", "12-45", "diciembre"):
            with self.assertRaises(Exception):
                _season_boundary(invalido, "05-31")


if __name__ == "__main__":
    unittest.main()
