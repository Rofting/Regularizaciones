import csv
import sqlite3
import tempfile
from contextlib import contextmanager
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import ContextManager, Iterator

from openpyxl import Workbook


def cents(value: str | int | float | Decimal) -> int:
    """Convierte un importe a céntimos con redondeo comercial explícito."""
    amount = Decimal(str(value)).quantize(Decimal("0.01"), ROUND_HALF_UP)
    return int(amount * 100)


@contextmanager
def temporary_database() -> Iterator[tuple[sqlite3.Connection, Path]]:
    """Entrega una conexión SQLite temporal y elimina el archivo al salir."""
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "gestion-test.db"
        connection = sqlite3.connect(path)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        try:
            yield connection, path
        finally:
            connection.close()


def make_reference_workbook(
    path: Path,
    *,
    period: str,
    actual_cents: int,
    billed_cents: int,
) -> Path:
    """Crea una referencia económica sintética sin datos personales."""
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "DATOS"
    data_sheet["A1"] = "COMUNIDAD DE PRUEBA"
    data_sheet["A3"] = period
    data_sheet["A4"] = "Nº de VIVIENDAS:"
    data_sheet["D4"] = 2

    readings_sheet = workbook.create_sheet("LECTURAS ACS M3")
    readings_sheet["B8"] = date(2024, 7, 1)
    readings_sheet["C8"] = date(2025, 7, 1)

    analysis_sheet = workbook.create_sheet("ANALISIS")
    analysis_sheet["F64"] = "IMPORTE COSTE REAL"
    analysis_sheet["G64"] = actual_cents / 100
    analysis_sheet["F66"] = "IMPORTE COBRADO"
    analysis_sheet["G66"] = billed_cents / 100
    analysis_sheet["F67"] = "DIFERENCIA"
    analysis_sheet["G67"] = (actual_cents - billed_cents) / 100

    workbook.save(path)
    return path


def make_owner_csv(path: Path, rows: list[dict[str, str]]) -> Path:
    """Crea un listado sintético con las cabeceras del origen real."""
    fieldnames = ["Codigo", "Nombre", "Fdenominacion", "Coeficiente", "Email"]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, delimiter=";")
        writer.writeheader()
        writer.writerows(rows)
    return path
