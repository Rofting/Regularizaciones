"""Aceptación de la mejora 12: recorrido completo con LibreOffice real.

Comunidad nueva creada desde el modelo común, dos propietarios, facturas y
lecturas confirmadas, cuotas cobradas, Excel recalculado por LibreOffice,
reparto que cuadra en céntimos y una carta por propietario. Se omite si
LibreOffice no está instalado; CI lo instala en Linux y Windows.
"""

import json
import shutil
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from datetime import date
from io import StringIO
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
CORE_DIR = PROJECT_ROOT / "core"
if str(CORE_DIR) not in sys.path:
    sys.path.insert(0, str(CORE_DIR))

import case_ingestion
import case_workflow_actions as workflow
import cuotas_servicio
import document_review
import expedient_service
import gestor_bd
import office_settings


def _soffice_available() -> bool:
    """Misma búsqueda que la aplicación (PATH y carpetas de programa en Windows)."""
    from office_recalculation import LibreOfficeRecalculator
    try:
        LibreOfficeRecalculator()._find_executable()
    except Exception:
        return False
    return True


@unittest.skipUnless(_soffice_available(), "LibreOffice no está instalado")
class LibreOfficeJourneyTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.root = Path(self.directory.name).resolve() / "proyecto"
        for name in ("config", "plantillas"):
            shutil.copytree(PROJECT_ROOT / name, self.root / name)
        # Proyecto de un despacho nuevo: sin perfiles ni plantillas de comunidades.
        shutil.rmtree(self.root / "config" / "excel_profiles", ignore_errors=True)
        shutil.rmtree(self.root / "plantillas" / "comunidades", ignore_errors=True)
        self.database = self.root / "data" / "gestion.db"
        with redirect_stdout(StringIO()):
            gestor_bd.crear_bd(str(self.database))

    def tearDown(self):
        self.directory.cleanup()

    def _prepare_case(self):
        connection = gestor_bd.conectar(str(self.database))
        try:
            office_settings.save_office_settings(
                connection, office_settings.OfficeSettings(name="Despacho Ensayo", city="Valencia"))
            community = gestor_bd.obtener_o_crear_comunidad(connection, "901", "CP Ensayo")
            case = expedient_service.create_case(
                connection, community, name="2026", start_date=date(2026, 1, 1), end_date=date(2026, 12, 31))
            for code in ("A-1", "B-2"):
                connection.execute(
                    """INSERT INTO propietarios(id_comunidad,codigo_vivienda,nombre_propietario,
                       coeficiente,activo,tipo_unidad) VALUES (?,?,?,50,1,'vivienda')""",
                    (community, code, f"Vecino {code}"))
            connection.commit()
            sources = self.root / "fuentes"
            sources.mkdir()
            documents = []
            for supply, amount, consumption in (("GAS", "500.00", "1000"), ("ELECTRICIDAD", "200.00", "300"),
                                                ("AGUA", "100.00", "40")):
                path = sources / f"{supply}.pdf"
                path.write_bytes(b"%PDF sintetico " + supply.encode())
                documents.append(case_ingestion.add_document_to_case(
                    connection, case.id_case, source_path=path,
                    archive_root=self.root / "data" / "expedientes", document_kind="invoice",
                    candidates={
                        "tipo_suministro": supply, "num_factura": f"F-{supply}",
                        "fecha_factura": "2026-12-31", "fecha_inicio": "2026-01-01",
                        "fecha_fin": "2026-12-31", "importe_total": amount,
                        "termino_fijo": str(float(amount) / 2), "termino_variable": str(float(amount) / 2),
                        ("consumo_m3" if supply == "AGUA" else "consumo_kwh"): consumption,
                        "cups": f"ES0000000000000000{supply[:2]}",
                    }, required_fields=()).document)
            readings = sources / "lecturas.csv"
            readings.write_text("sintetico", encoding="utf-8")
            documents.append(case_ingestion.add_document_to_case(
                connection, case.id_case, source_path=readings,
                archive_root=self.root / "data" / "expedientes", document_kind="reading",
                candidates={"vecinos": json.dumps([
                    {"vivienda": "A-1", "tipo": "ACS", "fecha_ant": "2026-01-01", "val_ant": 10,
                     "fecha_act": "2026-12-31", "val_act": 40},
                    {"vivienda": "B-2", "tipo": "ACS", "fecha_ant": "2026-01-01", "val_ant": 5,
                     "fecha_act": "2026-12-31", "val_act": 15},
                ])}, required_fields=()).document)
            for document in documents:
                case_ingestion.confirm_source_candidates(
                    connection, case.id_case, document.id_document, confirmed_by="ensayo")
            period = expedient_service.get_case(connection, case.id_case).period_id
            cuotas_servicio.generar_cuotas_mensuales(
                connection, community_id=community, period_id=period, servicio="ACS",
                tramos=[("2026-01", "50")])
            cuotas_servicio.registrar_cuota(
                connection, community_id=community, period_id=period, servicio="ACS",
                concepto="variable", fecha="2026-12-31", importe="300")
            connection.commit()
            self.assertEqual("ready_for_calculation",
                             document_review.validate_case_ready(connection, case.id_case).status)
            return community, case.id_case
        finally:
            connection.close()

    def test_new_community_from_the_common_model_reaches_letters(self):
        community, case_id = self._prepare_case()
        common = dict(id_case=case_id, active_community_id=community, project_root=self.root)

        workflow.run_generate_excel(self.database, output_root=self.root / "salidas", **common)
        distribution = workflow.run_calculate_distribution(self.database, **common)
        concepts = workflow.available_case_letter_concepts(self.database, **common)
        letters = workflow.run_generate_letters(
            self.database, selected_concepts=tuple(key for key, _label in concepts), **common)

        from openpyxl import load_workbook
        book = load_workbook(next((self.root / "salidas" / "Excels_Maestros").glob("*.xlsx")))
        try:
            self.assertEqual("'ANALISIS'!$A$1:$L$88", book["ANALISIS"].print_area)
            self.assertEqual("ES0000000000000000GA", book["GAS"]["C5"].value)
        finally:
            book.close()

        connection = gestor_bd.conectar(str(self.database))
        try:
            results = connection.execute(
                """SELECT r.concept_key,p.codigo_vivienda,r.actual_cents,r.difference_cents
                   FROM owner_concept_results r JOIN propietarios p USING(id_propietario)
                   ORDER BY 1,2""").fetchall()
        finally:
            connection.close()
        for concept, total in distribution.concept_totals_cents.items():
            with self.subTest(concept=concept):
                self.assertEqual(total, sum(row[2] for row in results if row[0] == concept))
        variable = {row[1]: row[2] for row in results if row[0] == "acs_variable"}
        self.assertEqual(3 * variable["B-2"], variable["A-1"])  # 30 m³ frente a 10 m³

        self.assertEqual((2, ()), (letters.generated_count, letters.failures))
        from docx import Document
        files = sorted(Path(letters.output_path).glob("*.docx"))
        self.assertEqual(2, len(files))
        for path in files:
            document = Document(str(path))
            self.assertEqual(1, len(document.inline_shapes))
            text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            self.assertIn("€", text)


if __name__ == "__main__":
    unittest.main()
