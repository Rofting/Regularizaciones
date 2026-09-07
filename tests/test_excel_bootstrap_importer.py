import csv
import hashlib
import sqlite3
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from datetime import date
from io import StringIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
CORE_DIR = PROJECT_ROOT / "core"
if str(CORE_DIR) not in sys.path:
    sys.path.insert(0, str(CORE_DIR))

import expedient_service
import excel_generator
import excel_profiles
import gestor_bd
import document_review
import excel_bootstrap_importer
from excel_bootstrap_importer import import_companion_sources, import_master_excel
from excel_export_service import _case_context, _template_registration


def _make_master(path: Path, *, missing_water_fixed: bool = False) -> Path:
    workbook = Workbook()
    data = workbook.active
    data.title = "DATOS"
    data["A1"] = "COMUNIDAD SINTETICA 658"
    data["A3"] = "PERIODO: EJERCICIO 01/09/2025 - 31/08/2026"
    data["A4"] = "Nº de VIVIENDAS:"
    data["D4"] = 2

    gas = workbook.create_sheet("GAS")
    for cell, value in {
        "B8": "Fecha fra.", "C8": "Días", "D8": "Fecha ini",
        "F8": "Fecha fin", "H8": "m³", "I8": "kWh", "J8": "€fijo",
        "K8": "€var", "L8": "TOTAL", "N8": "Nota",
    }.items():
        gas[cell] = value
    gas["A9"] = "2025-2026"
    for cell, value in {
        "B10": date(2025, 10, 5), "C10": 30, "D10": date(2025, 9, 1),
        "F10": date(2025, 9, 30), "H10": 10, "I10": 110,
        "J10": 20, "K10": 80, "L10": 100, "N10": "SUMINISTRO SINTETICO",
    }.items():
        gas[cell] = value
    gas["B11"] = "Suma 2025-2026"

    electricity = workbook.create_sheet("ELECTRICIDAD")
    for cell, value in {
        "B8": "Fecha fra.", "C8": "Fecha ini", "D8": "Fecha fin",
        "E8": "kWh", "F8": "€fijo", "G8": "€var", "H8": "TOTAL",
        "J8": "Nota",
    }.items():
        electricity[cell] = value
    electricity["A9"] = "2025-2026"
    for cell, value in {
        "B10": date(2025, 10, 8), "C10": date(2025, 9, 1),
        "D10": date(2025, 9, 30), "E10": 250, "F10": 25,
        "G10": 75, "H10": 100, "J10": "ELECTRICA SINTETICA",
    }.items():
        electricity[cell] = value
    electricity["B11"] = "Suma 2025-2026"

    water = workbook.create_sheet("AGUA")
    for cell, value in {
        "B12": "Fecha fra.", "C12": "Periodo", "D12": "Fecha ini",
        "F12": "Fecha fin", "H12": "m³", "T12": "TOTAL IVA",
        "V12": "C.Variable", "W12": "C.Fija",
    }.items():
        water[cell] = value
    water["A13"] = "2025-2026"
    for cell, value in {
        "B14": date(2025, 10, 9), "C14": "09/2025",
        "D14": date(2025, 9, 1), "F14": date(2025, 9, 30),
        "H14": 15, "T14": 60, "V14": 40,
        "W14": None if missing_water_fixed else 20,
    }.items():
        water[cell] = value
    water["B15"] = "Suma 2025-2026"

    other = workbook.create_sheet("OTROS GASTOS")
    other["D11"] = "EJERCICIO 2025-2026"
    other["D12"] = "Fecha"
    other["E12"] = "Comentario"
    other["F12"] = "Debe"
    other["D13"] = date(2025, 11, 1)
    other["E13"] = "Gasto extraordinario sintético"
    other["F13"] = 120
    other["E14"] = "TOTAL €......"
    other["F14"] = 120

    readings = workbook.create_sheet("LECTURAS ACS M3")
    readings["A6"] = "Ejercicio"
    readings["B6"] = "Fecha lect."
    readings["C6"] = "Fecha ant."
    readings["D6"] = "Lec ant."
    readings["E6"] = "Fecha act."
    readings["F6"] = "Lec act."
    readings["G6"] = "m³/€"
    readings["A7"] = "2025-2026"
    readings["C7"] = date(2025, 9, 1)
    readings["D7"] = 100
    readings["E7"] = date(2026, 8, 31)
    readings["F7"] = 140
    readings["G7"] = 40

    analysis = workbook.create_sheet("ANALISIS")
    analysis["G54"] = "RESULTADO DEL ANALISIS Y DEL REPARTO DE COSTES"
    analysis["G56"] = "GASTO"
    analysis["H56"] = "CUOTA FIJA"
    analysis["J56"] = "CUOTA VARIABLE"
    analysis["F64"] = "TOTAL €."
    analysis["G64"] = 260
    analysis["H64"] = 78
    analysis["J64"] = 182
    analysis["F66"] = "IMPORTE COBRADO"
    analysis["G66"] = 280
    analysis["H66"] = 112
    analysis["J66"] = 168
    analysis["F67"] = "DIFERENCIA"
    analysis["G67"] = 20
    analysis["H67"] = 34
    analysis["J67"] = -14
    analysis["E70"] = "Cuota fija vivienda/mes (€/mes)"
    analysis["F70"] = 8
    analysis["G70"] = 7.5
    analysis["E71"] = "Cuota variable m3 consumido (€/m3)"
    analysis["F71"] = 12
    analysis["G71"] = 13
    analysis["H23"] = "=SUM(H1:H22)"

    workbook.save(path)
    return path


def _make_profile_layout_master(path: Path) -> Path:
    """Libro saneado que usa únicamente las coordenadas del perfil declarativo.

    Las cabeceras deliberadamente no son alias del importador heredado: si la
    prueba pasa, la importación ha seguido el ``workbook_layout`` y no ha
    vuelto a depender de una etiqueta concreta.
    """
    workbook_path = _make_master(path)
    workbook = load_workbook(workbook_path)

    gas = workbook["GAS"]
    for cell, value in {
        "B8": "Fecha documento", "D8": "Cierre suministro",
        "G8": "Inicio suministro", "J8": "Consumo declarado",
        "K8": "Cargo estable", "L8": "Cargo uso",
        "M8": "Importe TOTAL con IVA", "O8": "Emisor",
        "B10": date(2025, 10, 5), "D10": date(2025, 9, 30),
        "G10": date(2025, 9, 1), "J10": 110,
        "K10": 20, "L10": 80, "M10": 100, "O10": "Gas sintético",
    }.items():
        gas[cell] = value

    electricity = workbook["ELECTRICIDAD"]
    electricity["H8"] = "Importe TOTAL con IVA"

    water = workbook["AGUA"]
    for cell, value in {
        "B12": "Fecha documento", "D12": "Cierre suministro",
        "F12": "Inicio suministro", "I12": "Consumo declarado",
        "R12": "Importe TOTAL con IVA", "T12": "Cargo uso",
        "U12": "Cargo estable", "B14": date(2025, 10, 9),
        "D14": date(2025, 9, 30), "F14": date(2025, 9, 1),
        "I14": 15, "R14": 60, "T14": 40, "U14": 20,
    }.items():
        water[cell] = value

    # Fuerza el uso de parameter_cells; el análisis por rótulos no debe ser
    # condición para conservar los importes económicos declarados en perfil.
    workbook["ANALISIS"]["G54"] = "BLOQUE ECONOMICO NO ETIQUETADO"
    workbook["ANALISIS"]["H64"] = 0
    workbook.save(workbook_path)
    return workbook_path


def _append_gas_invoice(source: Path, target: Path) -> Path:
    workbook = load_workbook(source)
    gas = workbook["GAS"]
    gas["B11"] = date(2025, 11, 5)
    gas["C11"] = 31
    gas["D11"] = date(2025, 10, 1)
    gas["F11"] = date(2025, 10, 31)
    gas["H11"] = 12
    gas["I11"] = 132
    gas["J11"] = 22
    gas["K11"] = 88
    gas["L11"] = 110
    gas["N11"] = "SUMINISTRO SINTETICO"
    gas["B12"] = "Suma 2025-2026"
    workbook.save(target)
    return target


def _make_owners(
    path: Path,
    *,
    first_coefficient: str = "50,00",
    second_coefficient: str = "50,00",
) -> Path:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["Codigo", "Nombre", "Fdenominacion", "Coeficiente", "Email"],
            delimiter=";",
        )
        writer.writeheader()
        writer.writerows([
            {
                "Codigo": "1", "Nombre": "Persona Inventada Uno",
                "Fdenominacion": "P1-A", "Coeficiente": first_coefficient,
                "Email": "uno@example.invalid",
            },
            {
                "Codigo": "2", "Nombre": "Persona Inventada Dos",
                "Fdenominacion": "P1-B", "Coeficiente": second_coefficient,
                "Email": "dos@example.invalid",
            },
        ])
    return path


def _make_readings(
    path: Path,
    *,
    include_second: bool = True,
    initial_header: str = "09/2025",
    final_header: str = "08/2026",
) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "LECTURAS"
    headers = ["Cod.", "Propiedad", initial_header, final_header]
    sheet.append(headers)
    sheet.append(["1", "P1-A", 100, 150])
    if include_second:
        sheet.append(["2", "P1-B", 200, 5])
    workbook.save(path)
    return path


class ExcelBootstrapImporterTest(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.root = Path(self.directory.name)
        self.database_path = self.root / "gestion.db"
        with redirect_stdout(StringIO()):
            gestor_bd.crear_bd(str(self.database_path))
        self.connection = sqlite3.connect(self.database_path)
        self.connection.row_factory = sqlite3.Row
        self.connection.execute("PRAGMA foreign_keys = ON")
        self.community_id = gestor_bd.obtener_o_crear_comunidad(
            self.connection, "658", "Comunidad sintética 658"
        )
        self.connection.commit()
        self.case = expedient_service.create_case(
            self.connection,
            self.community_id,
            name="2025-2026",
            start_date=date(2025, 9, 1),
            end_date=date(2026, 8, 31),
        )
        self.profile = excel_profiles.load_profile("658_acs_v1", PROJECT_ROOT)

    def tearDown(self):
        self.connection.close()
        self.directory.cleanup()

    def _onboarding_profile_and_template(self) -> tuple[excel_profiles.ExcelProfile, Path]:
        template = self.root / "plantilla-onboarding.xlsx"
        layout = excel_generator.create_canonical_community_template(
            template,
            community_code="658",
            community_name="Comunidad sintética 658",
            active_modules=("ACS",),
        )
        reading_sha256 = "a" * 64
        payload = {
            "key": "658_onboarding_v1",
            "version": "1",
            "community_code": "658",
            "template_relative_path": "plantillas/comunidades/658/658_onboarding_v1.xlsx",
            "active_modules": ["ACS"],
            "required_sheets": ["DATOS", "LECTURAS ACS M3", "ANALISIS"],
            "required_formula_cells": [],
            "concepts": [
                {
                    "key": "acs_fixed",
                    "allocation_method": "equal",
                    "actual_source": "period_parameters.acs_fixed_actual",
                    "billed_source": "period_parameters.acs_fixed_billed",
                    "required": True,
                },
                {
                    "key": "acs_variable",
                    "allocation_method": "consumption",
                    "actual_source": "period_parameters.acs_variable_actual",
                    "billed_source": "period_parameters.acs_variable_billed",
                    "required": True,
                },
            ],
            "onboarding_configuration": {
                "schema_version": 2,
                "service_decision": "ACS",
                "active_modules": ["ACS"],
                "reading_column": "Lectura ACS",
                "reading_bindings": [{
                    "module": "ACS",
                    "column": "Lectura ACS",
                    "meter": "Contador ACS",
                    "date": "2026-08-31",
                    "value": "150",
                    "source_sha256s": [reading_sha256],
                }],
                "invoice_decisions": [],
                "not_applicable_modules": [],
                "sources": [{
                    "kind": "meter_reading_excel",
                    "name": "lecturas.xlsx",
                    "sha256": reading_sha256,
                }],
            },
            "workbook_layout": layout,
        }
        return excel_profiles.validate_profile_payload(payload, self.root), template

    def test_master_import_is_idempotent_archived_and_traces_cell_values(self):
        workbook = _make_master(self.root / "modelo_sintetico.xlsx")

        first = import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=workbook, profile=self.profile, actor="Prueba",
        )
        second = import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=workbook, profile=self.profile, actor="Prueba",
        )

        self.assertEqual(first.id_batch, second.id_batch)
        self.assertEqual(3, first.imported_invoice_count)
        self.assertEqual(0, second.imported_invoice_count)
        self.assertEqual(
            3, self.connection.execute("SELECT COUNT(*) FROM facturas").fetchone()[0]
        )
        self.assertGreater(
            self.connection.execute(
                "SELECT COUNT(*) FROM source_values WHERE id_batch=?",
                (first.id_batch,),
            ).fetchone()[0],
            8,
        )
        traced = self.connection.execute(
            """SELECT sheet_name,cell_address FROM source_values
               WHERE id_batch=? AND entity_type='invoice'
               ORDER BY id_source_value LIMIT 1""",
            (first.id_batch,),
        ).fetchone()
        self.assertIsNotNone(traced[0])
        self.assertRegex(traced[1], r"^[A-Z]+\d+$")
        document = self.connection.execute(
            "SELECT archived_path,sha256 FROM source_documents"
        ).fetchone()
        self.assertTrue(Path(document["archived_path"]).is_file())
        self.assertNotEqual(workbook.resolve(), Path(document["archived_path"]).resolve())
        self.assertEqual(64, len(document["sha256"]))

    def test_later_empty_onboarding_master_reports_missing_required_field(self):
        profile, template = self._onboarding_profile_and_template()
        initial = import_master_excel(
            self.connection,
            id_case=self.case.id_case,
            workbook_path=template,
            profile=profile,
            actor="Prueba",
        )
        later_master = self.root / "maestro-posterior-vacio.xlsx"
        workbook = load_workbook(template)
        workbook["DATOS"]["A3"] = (
            "PERIODO: EJERCICIO 01/09/2025 - 31/08/2026"
        )
        workbook.save(later_master)
        workbook.close()

        later = import_master_excel(
            self.connection,
            id_case=self.case.id_case,
            workbook_path=later_master,
            profile=profile,
            actor="Prueba",
        )

        self.assertEqual(0, initial.open_issue_count)
        self.assertGreater(later.open_issue_count, 0)
        self.assertTrue(any(
            row["code"] == "MISSING_REQUIRED_FIELD"
            for row in self.connection.execute(
                "SELECT code FROM review_issues WHERE id_case=? AND status='open'",
                (self.case.id_case,),
            )
        ))

    def test_profile_layout_imports_descriptive_headers_endpoints_and_parameters(self):
        workbook = _make_profile_layout_master(self.root / "modelo_por_perfil.xlsx")

        result = import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=workbook, profile=self.profile, actor="Prueba",
        )

        self.assertEqual(3, result.imported_invoice_count)
        endpoints = {
            row["tipo_suministro"]: (row["fecha_inicio"], row["fecha_fin"])
            for row in self.connection.execute(
                "SELECT tipo_suministro,fecha_inicio,fecha_fin FROM facturas"
            )
        }
        self.assertEqual(("2025-09-01", "2025-09-30"), endpoints["GAS"])
        self.assertEqual(("2025-09-01", "2025-09-30"), endpoints["AGUA"])
        parameter_values = {
            row["parameter_key"]: row["numeric_value"]
            for row in self.connection.execute(
                "SELECT parameter_key,numeric_value FROM period_parameters"
            )
        }
        self.assertEqual(0, parameter_values["acs_fixed_actual"])
        self.assertEqual(168, parameter_values["acs_variable_billed"])
        sources = self.connection.execute(
            """SELECT COUNT(*) FROM source_values
               WHERE entity_type='period_parameter'
                 AND field_name LIKE 'parameter:%'"""
        ).fetchone()[0]
        self.assertEqual(8, sources)
        self.assertEqual(0, self.connection.execute(
            """SELECT COUNT(*) FROM review_issues
               WHERE code='UNRECOGNIZED_HEADER' AND field_name='ANALISIS.reference'"""
        ).fetchone()[0])

    def test_bootstrap_installs_verified_private_template_and_registers_it(self):
        project = self.root / "proyecto_portable"
        profile_dir = project / "config" / "excel_profiles"
        profile_dir.mkdir(parents=True)
        (profile_dir / "658_acs_v1.json").write_text(
            (PROJECT_ROOT / "config" / "excel_profiles" / "658_acs_v1.json").read_text(
                encoding="utf-8"
            ),
            encoding="utf-8",
        )
        profile = excel_profiles.load_profile("658_acs_v1", project)
        workbook = _make_master(self.root / "modelo_instalable.xlsx")
        original_source = workbook.read_bytes()
        source_hash = hashlib.sha256(original_source).hexdigest()

        first = import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=workbook, profile=profile, actor="Prueba",
            project_root=project,
        )
        template = project / profile.template_relative_path
        template_bytes = template.read_bytes()
        second = import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=workbook, profile=profile, actor="Prueba",
            project_root=project,
        )

        self.assertEqual(template, first.installed_template_path)
        self.assertEqual(source_hash, first.installed_template_sha256)
        self.assertEqual(template_bytes, template.read_bytes())
        self.assertEqual(original_source, workbook.read_bytes())
        self.assertEqual(template, second.installed_template_path)
        registration = self.connection.execute(
            """SELECT template_relative_path,template_sha256 FROM excel_template_profiles
               WHERE id_comunidad=? AND status='active'""",
            (self.community_id,),
        ).fetchone()
        self.assertEqual(profile.template_relative_path, registration["template_relative_path"])
        self.assertEqual(source_hash, registration["template_sha256"])
        document_review.validate_case_ready(self.connection, self.case.id_case)
        case_context = _case_context(self.connection, self.case.id_case)
        _profile_id, exported_template, exported_hash = _template_registration(
            self.connection, case_context, profile, project
        )
        self.assertEqual(template, exported_template)
        self.assertEqual(source_hash, exported_hash)
        document = self.connection.execute(
            "SELECT archived_path FROM source_documents WHERE document_kind='excel_master_bootstrap'"
        ).fetchone()
        self.assertNotEqual(template.resolve(), Path(document["archived_path"]).resolve())

    def test_invalid_bootstrap_does_not_install_template_before_valid_retry(self):
        project = self.root / "proyecto_reintento"
        profile_dir = project / "config" / "excel_profiles"
        profile_dir.mkdir(parents=True)
        (profile_dir / "658_acs_v1.json").write_text(
            (PROJECT_ROOT / "config" / "excel_profiles" / "658_acs_v1.json").read_text(
                encoding="utf-8"
            ),
            encoding="utf-8",
        )
        profile = excel_profiles.load_profile("658_acs_v1", project)
        invalid = self.root / "modelo_invalido.xlsx"
        invalid.write_bytes(b"no-es-un-libro-xlsx")

        with self.assertRaises(Exception):
            import_master_excel(
                self.connection, id_case=self.case.id_case,
                workbook_path=invalid, profile=profile, actor="Prueba",
                project_root=project,
            )

        template = project / profile.template_relative_path
        self.assertFalse(template.exists())
        self.assertEqual(
            0,
            self.connection.execute(
                "SELECT COUNT(*) FROM excel_template_profiles WHERE id_comunidad=?",
                (self.community_id,),
            ).fetchone()[0],
        )
        archived_invalid = self.connection.execute(
            """SELECT archived_path FROM source_documents
               WHERE original_name='modelo_invalido.xlsx'"""
        ).fetchone()
        self.assertTrue(Path(archived_invalid["archived_path"]).is_file())

        valid = _make_master(self.root / "modelo_valido.xlsx")
        result = import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=valid, profile=profile, actor="Prueba",
            project_root=project,
        )

        self.assertEqual(template, result.installed_template_path)
        self.assertTrue(template.is_file())
        self.assertEqual(
            1,
            self.connection.execute(
                "SELECT COUNT(*) FROM excel_template_profiles WHERE id_comunidad=?",
                (self.community_id,),
            ).fetchone()[0],
        )

    def test_different_master_hash_creates_template_conflict_without_overwriting(self):
        project = self.root / "proyecto_conflicto"
        profile_dir = project / "config" / "excel_profiles"
        profile_dir.mkdir(parents=True)
        (profile_dir / "658_acs_v1.json").write_text(
            (PROJECT_ROOT / "config" / "excel_profiles" / "658_acs_v1.json").read_text(
                encoding="utf-8"
            ), encoding="utf-8"
        )
        profile = excel_profiles.load_profile("658_acs_v1", project)
        initial = _make_master(self.root / "modelo_original.xlsx")
        import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=initial, profile=profile, actor="Prueba", project_root=project,
        )
        template = project / profile.template_relative_path
        original_template = template.read_bytes()
        changed = _append_gas_invoice(initial, self.root / "modelo_distinto.xlsx")

        with self.assertRaisesRegex(excel_bootstrap_importer.TemplateInstallationConflictError, "nueva versión"):
            import_master_excel(
                self.connection, id_case=self.case.id_case,
                workbook_path=changed, profile=profile, actor="Prueba", project_root=project,
            )

        self.assertEqual(original_template, template.read_bytes())
        issue = self.connection.execute(
            """SELECT code,message FROM review_issues
               WHERE id_case=? AND code='TEMPLATE_VERSION_CONFLICT' AND status='open'""",
            (self.case.id_case,),
        ).fetchone()
        self.assertIsNotNone(issue)
        self.assertIn("nueva versión", issue["message"])

    def test_template_installation_rejects_a_path_outside_the_project(self):
        from dataclasses import replace

        workbook = _make_master(self.root / "modelo_ruta_segura.xlsx")
        invalid_profile = replace(self.profile, template_relative_path="../fuera.xlsx")
        with self.assertRaisesRegex(excel_bootstrap_importer.TemplateInstallationError, "fuera"):
            import_master_excel(
                self.connection, id_case=self.case.id_case,
                workbook_path=workbook, profile=invalid_profile, actor="Prueba",
                project_root=self.root / "proyecto_ruta_segura",
            )
        self.assertFalse((self.root / "fuera.xlsx").exists())

    def test_private_community_templates_are_ignored_by_git(self):
        ignored = (PROJECT_ROOT / ".gitignore").read_text(encoding="utf-8")
        self.assertIn("plantillas/comunidades/", ignored)

    def test_missing_required_water_component_creates_issue_and_blocks_case(self):
        workbook = _make_master(
            self.root / "modelo_incompleto.xlsx", missing_water_fixed=True
        )

        result = import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=workbook, profile=self.profile, actor="Prueba",
        )

        self.assertGreater(result.open_issue_count, 0)
        issue = self.connection.execute(
            """SELECT code,field_name FROM review_issues
               WHERE status='open' ORDER BY id_issue LIMIT 1"""
        ).fetchone()
        self.assertEqual("MISSING_REQUIRED_FIELD", issue["code"])
        self.assertIn("fixed", issue["field_name"])
        self.assertEqual(
            "under_review",
            expedient_service.get_case(self.connection, self.case.id_case).status,
        )

    def test_changed_workbook_adds_only_the_new_invoice(self):
        first_path = _make_master(self.root / "modelo_v1.xlsx")
        second_path = _append_gas_invoice(first_path, self.root / "modelo_v2.xlsx")
        first = import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=first_path, profile=self.profile, actor="Prueba",
        )

        second = import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=second_path, profile=self.profile, actor="Prueba",
        )

        self.assertNotEqual(first.id_batch, second.id_batch)
        self.assertEqual(1, second.imported_invoice_count)
        self.assertEqual(
            4, self.connection.execute("SELECT COUNT(*) FROM facturas").fetchone()[0]
        )
        self.assertEqual(
            2, self.connection.execute("SELECT COUNT(*) FROM import_batches").fetchone()[0]
        )

    def test_gas_uses_cubic_metres_when_kwh_is_empty(self):
        workbook_path = _make_master(self.root / "gas_en_m3.xlsx")
        workbook = load_workbook(workbook_path)
        workbook["GAS"]["I10"] = None
        workbook.save(workbook_path)

        import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=workbook_path, profile=self.profile, actor="Prueba",
        )

        gas = self.connection.execute(
            """SELECT consumo_total,unidad_consumo FROM facturas
               WHERE tipo_suministro='GAS'"""
        ).fetchone()
        self.assertEqual(10, gas["consumo_total"])
        self.assertEqual("m3", gas["unidad_consumo"])

    def test_detects_unrecognized_header_incompatible_date_and_total_mismatch(self):
        workbook_path = _make_master(self.root / "modelo_inconsistente.xlsx")
        workbook = load_workbook(workbook_path)
        workbook["GAS"]["L8"] = "Importe irreconocible"
        workbook["ELECTRICIDAD"]["B10"] = date(2027, 1, 1)
        workbook["AGUA"]["W14"] = 10
        workbook.save(workbook_path)

        import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=workbook_path, profile=self.profile, actor="Prueba",
        )

        codes = {
            row[0] for row in self.connection.execute(
                "SELECT code FROM review_issues WHERE status='open'"
            )
        }
        self.assertTrue(
            {"UNRECOGNIZED_HEADER", "INVOICE_OUTSIDE_PERIOD", "TOTAL_MISMATCH"}.issubset(codes)
        )

    def test_idempotent_retry_preserves_resolved_issue_and_manual_correction(self):
        workbook = _make_master(
            self.root / "modelo_revisado.xlsx", missing_water_fixed=True
        )
        first = import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=workbook, profile=self.profile, actor="Prueba",
        )
        issue_id = self.connection.execute(
            """SELECT id_issue FROM review_issues
               WHERE code='MISSING_REQUIRED_FIELD' AND status='open'"""
        ).fetchone()[0]
        document_review.resolve_issue(
            self.connection, issue_id, value="20", reason="Dato confirmado",
            resolved_by="Prueba",
        )

        second = import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=workbook, profile=self.profile, actor="Prueba",
        )

        self.assertEqual(first.id_batch, second.id_batch)
        self.assertEqual(1, self.connection.execute(
            "SELECT COUNT(*) FROM manual_corrections"
        ).fetchone()[0])
        self.assertEqual(0, self.connection.execute(
            "SELECT COUNT(*) FROM review_issues WHERE status='open'"
        ).fetchone()[0])

    def test_companions_preserve_owner_and_mark_counter_reset_for_review(self):
        self.connection.execute(
            """INSERT INTO propietarios
               (id_comunidad,codigo_vivienda,nombre_propietario,coeficiente,email)
               VALUES (?,?,?,?,?)""",
            (self.community_id, "P1-A", "Nombre confirmado", 51, "confirmado@example.invalid"),
        )
        self.connection.commit()
        owners = _make_owners(self.root / "propietarios_sinteticos.csv")
        readings = _make_readings(self.root / "lecturas_sinteticas.xlsx")

        first = import_companion_sources(
            self.connection, id_case=self.case.id_case,
            owner_list_path=owners, readings_path=readings,
            profile=self.profile, actor="Prueba",
        )
        second = import_companion_sources(
            self.connection, id_case=self.case.id_case,
            owner_list_path=owners, readings_path=readings,
            profile=self.profile, actor="Prueba",
        )

        self.assertEqual(1, first.imported_owner_count)
        self.assertEqual(4, first.imported_reading_count)
        self.assertEqual(0, second.imported_owner_count)
        self.assertEqual(0, second.imported_reading_count)
        owner = self.connection.execute(
            """SELECT nombre_propietario,email FROM propietarios
               WHERE id_comunidad=? AND codigo_vivienda='P1-A'""",
            (self.community_id,),
        ).fetchone()
        self.assertEqual("Nombre confirmado", owner["nombre_propietario"])
        reset = self.connection.execute(
            """SELECT estado,valor_acumulado FROM lecturas_vecino lv
               JOIN propietarios p ON p.id_propietario=lv.id_propietario
               WHERE p.codigo_vivienda='P1-B' AND lv.fecha_lectura='2026-08-31'"""
        ).fetchone()
        self.assertEqual("contador_averiado", reset["estado"])
        self.assertEqual(5, reset["valor_acumulado"])
        self.assertEqual(
            1, self.connection.execute(
                "SELECT COUNT(*) FROM review_issues WHERE code='COUNTER_RESET' AND status='open'"
            ).fetchone()[0]
        )
        self.assertEqual(2, self.connection.execute(
            "SELECT COUNT(*) FROM source_documents"
        ).fetchone()[0])
        statuses = {
            row["document_kind"]: row["status"]
            for row in self.connection.execute(
                "SELECT document_kind,status FROM source_documents"
            )
        }
        self.assertEqual("validated", statuses["owner_list"])
        self.assertEqual("under_review", statuses["meter_readings"])

    def test_missing_active_owner_reading_range_creates_review_issue(self):
        owners = _make_owners(self.root / "propietarios_sinteticos.csv")
        readings = _make_readings(
            self.root / "lecturas_incompletas.xlsx", include_second=False
        )

        result = import_companion_sources(
            self.connection, id_case=self.case.id_case,
            owner_list_path=owners, readings_path=readings,
            profile=self.profile, actor="Prueba",
        )

        self.assertGreater(result.open_issue_count, 0)
        issue = self.connection.execute(
            """SELECT code,field_name FROM review_issues
               WHERE code='MISSING_READING_RANGE' AND status='open'"""
        ).fetchone()
        self.assertIsNotNone(issue)
        self.assertIn("P1-B", issue["field_name"])

    def test_reading_headers_from_another_period_are_not_persisted(self):
        owners = _make_owners(self.root / "propietarios_periodo.csv")
        readings = _make_readings(
            self.root / "lecturas_2020.xlsx",
            initial_header="09/2020",
            final_header="08/2021",
        )

        result = import_companion_sources(
            self.connection, id_case=self.case.id_case,
            owner_list_path=owners, readings_path=readings,
            profile=self.profile, actor="Prueba",
        )

        self.assertEqual(0, result.imported_reading_count)
        self.assertEqual(0, self.connection.execute(
            "SELECT COUNT(*) FROM lecturas_vecino"
        ).fetchone()[0])
        self.assertEqual(1, self.connection.execute(
            """SELECT COUNT(*) FROM review_issues
               WHERE code='INCOMPATIBLE_DATE' AND status='open'"""
        ).fetchone()[0])

    def test_previous_month_header_is_accepted_only_as_initial_period_reading(self):
        owners = _make_owners(self.root / "propietarios_previos.csv")
        readings = _make_readings(
            self.root / "lecturas_mes_anterior.xlsx", initial_header="08/2025"
        )

        result = import_companion_sources(
            self.connection, id_case=self.case.id_case,
            owner_list_path=owners, readings_path=readings,
            profile=self.profile, actor="Prueba",
        )

        self.assertEqual(4, result.imported_reading_count)
        self.assertEqual(0, self.connection.execute(
            """SELECT COUNT(*) FROM review_issues
               WHERE code IN ('INCOMPATIBLE_DATE','MISSING_READING_RANGE')"""
        ).fetchone()[0])
        initial_sources = self.connection.execute(
            """SELECT cell_address FROM source_values
               WHERE entity_type='reading' AND field_name='initial'
               ORDER BY id_source_value"""
        ).fetchall()
        self.assertEqual(["C2", "C3"], [row["cell_address"] for row in initial_sources])

    def test_missing_invoice_values_are_not_replaced_with_zero(self):
        workbook_path = _make_master(self.root / "factura_incompleta.xlsx")
        workbook = load_workbook(workbook_path)
        workbook["GAS"]["J10"] = None
        workbook["AGUA"]["W14"] = 0
        workbook.save(workbook_path)

        result = import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=workbook_path, profile=self.profile, actor="Prueba",
        )

        self.assertEqual(2, result.imported_invoice_count)
        self.assertIsNone(self.connection.execute(
            "SELECT id_factura FROM facturas WHERE tipo_suministro='GAS'"
        ).fetchone())
        water_fixed = self.connection.execute(
            """SELECT amount FROM invoice_components c
               JOIN facturas f ON f.id_factura=c.id_factura
               WHERE f.tipo_suministro='AGUA' AND c.component_key='fixed'"""
        ).fetchone()
        self.assertEqual(0, water_fixed["amount"])
        self.assertGreater(self.connection.execute(
            """SELECT COUNT(*) FROM review_issues
               WHERE code='MISSING_REQUIRED_FIELD' AND field_name LIKE 'GAS.%fixed'"""
        ).fetchone()[0], 0)

    def test_missing_coefficient_does_not_create_owner_but_explicit_zero_does(self):
        owners = _make_owners(
            self.root / "coeficientes.csv",
            first_coefficient="",
            second_coefficient="0",
        )
        readings = _make_readings(self.root / "lecturas_coeficientes.xlsx")

        import_companion_sources(
            self.connection, id_case=self.case.id_case,
            owner_list_path=owners, readings_path=readings,
            profile=self.profile, actor="Prueba",
        )

        self.assertIsNone(self.connection.execute(
            """SELECT id_propietario FROM propietarios
               WHERE id_comunidad=? AND codigo_vivienda='P1-A'""",
            (self.community_id,),
        ).fetchone())
        explicit_zero = self.connection.execute(
            """SELECT coeficiente FROM propietarios
               WHERE id_comunidad=? AND codigo_vivienda='P1-B'""",
            (self.community_id,),
        ).fetchone()
        self.assertEqual(0, explicit_zero["coeficiente"])

    def test_same_owner_batch_is_linked_to_each_case_and_period(self):
        owners = _make_owners(self.root / "propietarios_compartidos.csv")
        readings_first = _make_readings(self.root / "lecturas_2025.xlsx")
        first = import_companion_sources(
            self.connection, id_case=self.case.id_case,
            owner_list_path=owners, readings_path=readings_first,
            profile=self.profile, actor="Prueba",
        )
        second_case = expedient_service.create_case(
            self.connection,
            self.community_id,
            name="2026-2027",
            start_date=date(2026, 9, 1),
            end_date=date(2027, 8, 31),
        )
        readings_second = _make_readings(
            self.root / "lecturas_2026.xlsx",
            initial_header="09/2026",
            final_header="08/2027",
        )

        second = import_companion_sources(
            self.connection, id_case=second_case.id_case,
            owner_list_path=owners, readings_path=readings_second,
            profile=self.profile, actor="Prueba",
        )

        self.assertEqual(first.id_batches[0], second.id_batches[0])
        links = self.connection.execute(
            """SELECT id_case,id_periodo FROM case_import_batches
               WHERE source_kind='owner_list' ORDER BY id_case"""
        ).fetchall()
        self.assertEqual(
            [(self.case.id_case, first.id_periodo), (second_case.id_case, second.id_periodo)],
            [(row["id_case"], row["id_periodo"]) for row in links],
        )

    def test_master_is_parsed_from_archived_copy_after_original_mutates(self):
        workbook_path = _make_master(self.root / "original_mutable.xlsx")
        real_register = excel_bootstrap_importer.register_source_document

        def archive_then_mutate(*args, **kwargs):
            result = real_register(*args, **kwargs)
            changed = load_workbook(workbook_path)
            changed["GAS"]["L10"] = 999
            changed.save(workbook_path)
            return result

        with patch.object(
            excel_bootstrap_importer,
            "register_source_document",
            side_effect=archive_then_mutate,
        ):
            import_master_excel(
                self.connection, id_case=self.case.id_case,
                workbook_path=workbook_path, profile=self.profile, actor="Prueba",
            )

        gas_total = self.connection.execute(
            "SELECT importe_total FROM facturas WHERE tipo_suministro='GAS'"
        ).fetchone()[0]
        self.assertEqual(100, gas_total)

    def test_tampered_archived_copy_is_rejected_before_parsing(self):
        workbook_path = _make_master(self.root / "origen_verificable.xlsx")
        real_register = excel_bootstrap_importer.register_source_document

        def archive_then_tamper(*args, **kwargs):
            result = real_register(*args, **kwargs)
            with Path(result[0].archived_path).open("ab") as archived:
                archived.write(b"alterado")
            return result

        with patch.object(
            excel_bootstrap_importer,
            "register_source_document",
            side_effect=archive_then_tamper,
        ):
            with self.assertRaisesRegex(RuntimeError, "hash"):
                import_master_excel(
                    self.connection, id_case=self.case.id_case,
                    workbook_path=workbook_path, profile=self.profile, actor="Prueba",
                )

        self.assertEqual(0, self.connection.execute(
            "SELECT COUNT(*) FROM facturas"
        ).fetchone()[0])

    def test_invoice_row_with_unreadable_date_creates_issue(self):
        workbook_path = _make_master(self.root / "fecha_factura_ilegible.xlsx")
        workbook = load_workbook(workbook_path)
        workbook["GAS"]["B10"] = "fecha imposible"
        workbook.save(workbook_path)

        import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=workbook_path, profile=self.profile, actor="Prueba",
        )

        self.assertIsNone(self.connection.execute(
            "SELECT id_factura FROM facturas WHERE tipo_suministro='GAS'"
        ).fetchone())
        issue = self.connection.execute(
            """SELECT code,field_name FROM review_issues
               WHERE id_case=? AND field_name='GAS.B10.invoice_date'""",
            (self.case.id_case,),
        ).fetchone()
        self.assertEqual("INCOMPATIBLE_DATE", issue["code"])

    def test_invoice_date_outside_case_period_uses_dedicated_issue_code(self):
        workbook_path = _make_master(self.root / "fecha_factura_fuera_periodo.xlsx")
        workbook = load_workbook(workbook_path)
        workbook["GAS"]["B10"] = date(2027, 1, 1)
        workbook.save(workbook_path)

        import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=workbook_path, profile=self.profile, actor="Prueba",
        )

        issue = self.connection.execute(
            """SELECT code FROM review_issues
               WHERE id_case=? AND field_name='GAS.B10.invoice_date'""",
            (self.case.id_case,),
        ).fetchone()
        self.assertEqual("INVOICE_OUTSIDE_PERIOD", issue["code"])

    def test_invoice_row_requires_valid_supply_endpoints(self):
        workbook_path = _make_master(self.root / "extremos_incompletos.xlsx")
        workbook = load_workbook(workbook_path)
        workbook["GAS"]["D10"] = None
        workbook["GAS"]["F10"] = "fin ilegible"
        workbook.save(workbook_path)

        import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=workbook_path, profile=self.profile, actor="Prueba",
        )

        self.assertIsNone(self.connection.execute(
            "SELECT id_factura FROM facturas WHERE tipo_suministro='GAS'"
        ).fetchone())
        issues = {
            row["field_name"]: row["code"]
            for row in self.connection.execute(
                """SELECT field_name,code FROM review_issues
                   WHERE id_case=? AND field_name LIKE 'GAS.%date'""",
                (self.case.id_case,),
            )
        }
        self.assertEqual("MISSING_REQUIRED_FIELD", issues["GAS.D10.start_date"])
        self.assertEqual("INCOMPATIBLE_DATE", issues["GAS.F10.end_date"])

    def test_reused_master_hash_validates_new_case_and_returns_current_period(self):
        workbook = _make_master(self.root / "maestro_reutilizado.xlsx")
        first = import_master_excel(
            self.connection, id_case=self.case.id_case,
            workbook_path=workbook, profile=self.profile, actor="Prueba",
        )
        second_case = expedient_service.create_case(
            self.connection,
            self.community_id,
            name="2026-2027",
            start_date=date(2026, 9, 1),
            end_date=date(2027, 8, 31),
        )

        second = import_master_excel(
            self.connection, id_case=second_case.id_case,
            workbook_path=workbook, profile=self.profile, actor="Prueba",
        )

        self.assertEqual(first.id_batch, second.id_batch)
        self.assertNotEqual(first.id_periodo, second.id_periodo)
        self.assertEqual(
            expedient_service.get_case(
                self.connection, second_case.id_case
            ).period_id,
            second.id_periodo,
        )
        self.assertGreater(self.connection.execute(
            """SELECT COUNT(*) FROM review_issues
               WHERE id_case=? AND code='INCOMPATIBLE_DATE' AND status='open'""",
            (second_case.id_case,),
        ).fetchone()[0], 0)
        second_document = self.connection.execute(
            """SELECT status FROM source_documents
               WHERE id_case=? AND document_kind='excel_master_bootstrap'""",
            (second_case.id_case,),
        ).fetchone()
        self.assertEqual("under_review", second_document["status"])

    def test_reused_companion_hash_validates_new_period_and_document_states(self):
        owners = _make_owners(self.root / "propietarios_reutilizados.csv")
        readings = _make_readings(self.root / "lecturas_reutilizadas.xlsx")
        import_companion_sources(
            self.connection, id_case=self.case.id_case,
            owner_list_path=owners, readings_path=readings,
            profile=self.profile, actor="Prueba",
        )
        second_case = expedient_service.create_case(
            self.connection,
            self.community_id,
            name="2026-2027",
            start_date=date(2026, 9, 1),
            end_date=date(2027, 8, 31),
        )

        second = import_companion_sources(
            self.connection, id_case=second_case.id_case,
            owner_list_path=owners, readings_path=readings,
            profile=self.profile, actor="Prueba",
        )

        self.assertEqual(
            expedient_service.get_case(
                self.connection, second_case.id_case
            ).period_id,
            second.id_periodo,
        )
        statuses = {
            row["document_kind"]: row["status"]
            for row in self.connection.execute(
                """SELECT document_kind,status FROM source_documents
                   WHERE id_case=?""",
                (second_case.id_case,),
            )
        }
        self.assertEqual("validated", statuses["owner_list"])
        self.assertEqual("under_review", statuses["meter_readings"])
        self.assertGreater(self.connection.execute(
            """SELECT COUNT(*) FROM review_issues
               WHERE id_case=? AND code='INCOMPATIBLE_DATE' AND status='open'""",
            (second_case.id_case,),
        ).fetchone()[0], 0)
        self.assertEqual(0, self.connection.execute(
            "SELECT COUNT(*) FROM lecturas_vecino WHERE id_periodo=?",
            (second.id_periodo,),
        ).fetchone()[0])


if __name__ == "__main__":
    unittest.main()
